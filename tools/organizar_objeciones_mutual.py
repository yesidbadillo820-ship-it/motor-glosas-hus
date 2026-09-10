"""organizar_objeciones_mutual.py — Objeciones de MUTUAL SER → Excel OBJECIONES (DGH).

Toma el Excel de glosas que entrega **MUTUAL SER** (hoja `CONSOLIDADO`, 7
columnas) y lo convierte al formato de trabajo de 16 columnas (hoja
`OBJECIONES`) que se importa en Dinámica Gerencial, el mismo de FAMISANAR,
SANITAS, SAVIA y el Dispensario.

CÓMO VIENE EL ARCHIVO DE MUTUAL. Los encabezados sí están bien rotulados, así
que las columnas se resuelven por nombre (tolerando tildes y mayúsculas):

    Número de factura    → HUS0000544271
    SERVICIO             → el CÓDIGO del servicio (903883, FMQ0817, 10A004…)
    Cantidad facturada   → cantidad
    Valor glosado        → viene como texto con signo de pesos ("$ 304.500")
    Concepto de glosa    → el texto del concepto ("Recargos no pactados - TARIFAS")
    Código de glosa      → TA0201, TA2901, FA1305… (ya limpio)
    Observacion          → el motivo, y adentro el NOMBRE del servicio

OJO CON LA COLUMNA «SERVICIO». Pese al nombre, no trae la descripción sino el
**código**. El nombre del servicio viaja escondido en la observación:

    "La tecnología 903883 - GLUCOSA SEMIAUTOMATIZADA [GLUCOMETRIA] no se
     encuentra dentro del contrato número 20352."
    "La tarifa facturada 31800 no coincide con la tarifa 26712.0 definida en
     el contrato para la tecnologia 931001 - TERAPIA FISICA INTEGRAL."

`nombre_del_servicio()` lo saca de ahí para que el cruce contra el DGH tenga
con qué desempatar. Cuando la observación no lo trae ("El servicio facturado
389002 no se encuentra habilitado"), el cruce se apoya sólo en el código y el
valor, y si no alcanza el renglón va a REVISAR sin adivinar nada.

DOBLE GLOSA: EL MISMO SERVICIO BAJO DOS CONCEPTOS. MUTUAL lista un servicio
una vez por cada concepto de glosa que le aplica, pero en su total lo cuenta
**una sola vez, por el mayor de los valores**:

    FMQ0463  x1  $244.400  TA0201  "Consultas… - TARIFAS"
    FMQ0463  x1  $244.400  TA0601  "Dispositivos médicos - TARIFAS"
    389002   x1  $ 54.594  TA2901  "Recargos no pactados - TARIFAS"
    389002   x1  $181.900  FA1305  "…no se encuentra habilitado"

Sumar todas las filas infla la glosa: en el lote del 7 de septiembre daba
$26.636.056 cuando MUTUAL reportó **$24.462.346** — $2.173.710 de más. Por eso
`fusionar_dobles_glosas()` deja un renglón por (factura, servicio, cantidad)
con el valor mayor, y anota los demás códigos en CRDOBSERV: no se pierde
ninguna glosa, sólo se deja de contar dos veces la misma. Es el mismo trato
que el bot de EMSSANAR le da a sus dobles glosas.

Si dentro de un grupo se repite el MISMO código, eso ya no es doble glosa sino
una repetición legítima (dos renglones de verdad, como los del Dispensario) y
el bot NO fusiona: los deja y avisa.

REGLAS FIJAS DEL FORMATO (ver CLAUDE.md, valen para todas las entidades):
    CTNCENCOS  → siempre vacía
    CROTIPOBJ  → por factura: 0 = ADMINISTRATIVA, 1 = MEDICA, 2 = MIXTA
    SLNSERPRO  → prohibido inventar: sin cruce confiable, celda vacía
    Ningún renglón se descarta por no haber cruzado: va igual, con la celda
    vacía para completarlo a mano. Lo único que se junta es la doble glosa
    de arriba, y precisamente para que el total cuadre con el de MUTUAL.

USO:
    py tools\\organizar_objeciones_mutual.py ^
        --entrada       "D:\\...\\MUTUAL_7_SEPTIEMBRE.xlsx" ^
        --servicios-dgh "D:\\...\\SERVICIOS_FACTURADOS_DGH.xlsx" ^
        --salida        "D:\\...\\OBJECIONES_MUTUAL_07092026.xlsx" --consolidado ^
        --reporte-cruce "D:\\...\\CRUCE_MUTUAL_07092026.xlsx" ^
        --fecha 2026-09-07

INSTALACIÓN (una vez):  py -m pip install openpyxl
"""

from __future__ import annotations

import argparse
import datetime as _dt
import logging
import re
import sys
from collections import defaultdict
from pathlib import Path

# Motor del cruce y lector de pesos, comunes a todos los bots.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from _cruce_dgh import (  # noqa: E402
    escribir_reporte_cruce,
    factura_larga,
    leer_servicios_dgh,
    norm_header,
    resolver_servicio,
    traza,
    verificar_reglas,
)
from _dinero import a_entero  # noqa: E402

logger = logging.getLogger("organizar_mutual")


# ─── Formato de salida: layout de 16 columnas (hoja OBJECIONES) ──────────────

COLUMNAS_OBJECIONES: tuple[str, ...] = (
    "CDCONSEC",
    "CDFECDOC",
    "CRNCXC",
    "CROFECOBJ",
    "CROREFERE",
    "CROOBSERV",
    "CROCLAOBJ",
    "CRNCLAOBJ",
    "GENUSUARIO4",
    "CRNCONOBJ",
    "SLNSERPRO",
    "IDRIPS",
    "CTNCENCOS",
    "CROVALOBJ",
    "CRDOBSERV",
    "CROTIPOBJ",
)

# number_format por columna, igual que en los archivos reales.
FORMATOS: dict[str, str] = {
    "CDCONSEC": "@",
    "CDFECDOC": "mm-dd-yy",
    "CRNCXC": "@",
    "CROFECOBJ": "mm-dd-yy",
    "CROREFERE": "@",
    "CROOBSERV": "@",
    "CROCLAOBJ": "General",
    "CRNCLAOBJ": "@",
    "GENUSUARIO4": "@",
    "CRNCONOBJ": "@",
    "SLNSERPRO": "@",
    "IDRIPS": "@",
    "CTNCENCOS": "@",
    "CROVALOBJ": '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-',
    "CRDOBSERV": "@",
    "CROTIPOBJ": "0",
}

CROCLAOBJ_CONST = 0
GENUSUARIO4_CONST = "999"

# Encabezados que reconoce el lector (normalizados: sin tildes, en mayúscula).
ALIAS_COLUMNAS: dict[str, set[str]] = {
    "factura": {"NUMERO DE FACTURA", "NUMERO FACTURA", "FACTURA"},
    "cod_servicio": {"SERVICIO", "CODIGO SERVICIO", "CODIGO DEL SERVICIO"},
    "cantidad": {"CANTIDAD FACTURADA", "CANTIDAD"},
    "valor": {"VALOR GLOSADO", "VALOR GLOSA", "VALOR OBJETADO"},
    "concepto": {"CONCEPTO DE GLOSA", "CONCEPTO GLOSA", "CONCEPTO"},
    "codigo": {"CODIGO DE GLOSA", "CODIGO GLOSA", "COD GLOSA"},
    "observacion": {"OBSERVACION", "OBSERVACIONES"},
}

# Sin estas cuatro el archivo no sirve: se para antes de entregar algo malo.
COLUMNAS_OBLIGATORIAS = ("factura", "cod_servicio", "valor", "codigo")

_RE_CODIGO_GLOSA = re.compile(r"^([A-Z]{2}\d{2})\s*(\d{2})$")

GRUPO_CLINICO = "CL"


def _texto(v: object) -> str:
    return " ".join(str(v if v is not None else "").split())


def codigo_glosa(valor: object) -> str:
    """Deja el código de glosa limpio: 'TA02 01' → 'TA0201'; 'TA0201' tal cual."""
    t = _texto(valor).upper()
    m = _RE_CODIGO_GLOSA.match(t)
    if m:
        return m.group(1) + m.group(2)
    return t.replace(" ", "")


# El nombre del servicio va detrás de "<código> - " dentro de la observación.
_RE_NOMBRE = re.compile(r"-\s*(.+)$", re.DOTALL)

# Colas que MUTUAL le pega al nombre y que no son parte del nombre.
_RE_COLA = re.compile(
    r"\s+(?:no se encuentra|no esta|no está|no corresponde|no cumple|"
    r"se encuentra|esta incluid|está incluid)\b.*$",
    re.IGNORECASE | re.DOTALL,
)


def nombre_del_servicio(observacion: object, cod_servicio: str = "") -> str:
    """Saca el nombre del servicio del texto de la observación de MUTUAL.

    'La tecnología 903883 - GLUCOSA SEMIAUTOMATIZADA [GLUCOMETRIA] no se
    encuentra dentro del contrato número 20352.' → 'GLUCOSA SEMIAUTOMATIZADA
    [GLUCOMETRIA]'.

    Se corta a partir del código del servicio para no confundirse con los
    guiones que traiga el propio texto ('Recargos no pactados - TARIFAS'). Si
    la observación no nombra el servicio, se devuelve cadena vacía: el cruce
    se apoyará en el código y el valor, y si no alcanza no se inventa nada.
    """
    t = _texto(observacion)
    if not t:
        return ""
    cod = _texto(cod_servicio).upper()
    if cod:
        pos = t.upper().rfind(cod)
        if pos < 0:
            return ""
        t = t[pos + len(cod) :]
    m = _RE_NOMBRE.match(t.strip()) if cod else None
    if m is None:
        return ""
    nombre = _RE_COLA.sub("", m.group(1)).strip(" .;,")
    return nombre


def crotipobj_factura(grupos: set[str]) -> int:
    """Tipo de objeción de la factura según sus grupos de concepto (2 letras).

    Los tres valores que reconoce DGH: **0 = ADMINISTRATIVA** (sólo TA/FA/SO/
    AU/CO…), **1 = MEDICA** (sólo CL) y **2 = MIXTA** (CL junto con cualquier
    administrativa). Se decide por FACTURA, no por renglón.
    """
    tiene_cl = GRUPO_CLINICO in grupos
    tiene_admin = any(g and g != GRUPO_CLINICO for g in grupos)
    if tiene_cl and tiene_admin:
        return 2
    if tiene_cl:
        return 1
    return 0


def construir_crdobserv(
    codigo: str,
    concepto: str,
    observacion: str,
    valor: int,
    fusionadas: list[tuple[str, int]] | None = None,
) -> str:
    """CRDOBSERV = ``<código> <concepto>: <observación>$<valor>``.

    Si el renglón absorbió otras glosas del mismo servicio (doble glosa), se
    anotan al final con su valor, para que no se pierda de vista bajo qué
    otros conceptos lo objetó MUTUAL.
    """
    partes = [p for p in (concepto, observacion) if p]
    texto = ": ".join(partes) if len(partes) == 2 else (partes[0] if partes else "")
    prefijo = f"{codigo} " if codigo else ""
    cola = ""
    if fusionadas:
        detalle = ", ".join(f"{c} ${v}" for c, v in fusionadas)
        cola = f" (también glosado como {detalle})"
    return f"{prefijo}{texto}${valor}{cola}"


# ─── Lectura del Excel de MUTUAL ─────────────────────────────────────────────


def mapear_columnas(encabezados: tuple) -> dict[str, int]:
    """{campo: posición} a partir de la fila de encabezados.

    Si falta alguna de las columnas indispensables se detiene con un mensaje
    que dice cuál, en vez de entregar un archivo silenciosamente incompleto.
    """
    normalizados = [norm_header(h) for h in encabezados]
    indices: dict[str, int] = {}
    for campo, nombres in ALIAS_COLUMNAS.items():
        pos = next((i for i, h in enumerate(normalizados) if h in nombres), None)
        if pos is not None:
            indices[campo] = pos

    faltan = [c for c in COLUMNAS_OBLIGATORIAS if c not in indices]
    if faltan:
        raise ValueError(
            "El archivo de MUTUAL no trae las columnas "
            f"{', '.join(faltan)}. Encabezados leídos: {[h for h in normalizados if h]}. "
            "Se esperaba el consolidado de MUTUAL SER (Número de factura, SERVICIO, "
            "Cantidad facturada, Valor glosado, Concepto de glosa, Código de glosa, "
            "Observacion)."
        )
    return indices


def leer_mutual(ruta: Path, hoja: str | None = None) -> list[dict]:
    """Lee el Excel de MUTUAL SER y devuelve una lista de objeciones."""
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(ruta), data_only=True, read_only=True)
    try:
        if hoja:
            objetivo = _texto(hoja).upper()
            ws = next((x for x in wb.worksheets if _texto(x.title).upper() == objetivo), None)
            if ws is None:
                raise ValueError(f"La hoja '{hoja}' no existe en {ruta.name}.")
        else:
            ws = next(
                (x for x in wb.worksheets if _texto(x.title).upper() == "CONSOLIDADO"), wb.active
            )
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            raise ValueError(f"La hoja '{ws.title}' de {ruta.name} está vacía.")
        idx = mapear_columnas(filas[0])

        objeciones: list[dict] = []
        for n, fila in enumerate(filas[1:], start=2):
            if not any(fila):
                continue

            def dato(campo, f=fila):
                pos = idx.get(campo)
                return f[pos] if pos is not None and pos < len(f) else None

            factura = _texto(dato("factura"))
            if not factura:
                continue
            cod_servicio = _texto(dato("cod_servicio"))
            observacion = _texto(dato("observacion"))
            objeciones.append(
                {
                    "fila_excel": n,
                    "cxc": factura_larga(factura),
                    "codigo": codigo_glosa(dato("codigo")),
                    "cod_servicio": cod_servicio,
                    "servicio": nombre_del_servicio(observacion, cod_servicio),
                    "cantidad": a_entero(dato("cantidad")),
                    "concepto": _texto(dato("concepto")),
                    "observacion": observacion,
                    "valor": a_entero(dato("valor")),
                }
            )
        return objeciones
    finally:
        wb.close()


# ─── Doble glosa: el mismo servicio bajo dos conceptos ───────────────────────


def fusionar_dobles_glosas(objeciones: list[dict], avisar=None) -> list[dict]:
    """Deja un renglón por (factura, servicio, cantidad), con el valor MAYOR.

    MUTUAL objeta un mismo servicio una vez por cada concepto que le aplica
    —el mismo dispositivo sale como «Consultas… - TARIFAS» (TA0201) y como
    «Dispositivos médicos - TARIFAS» (TA0601)— pero en su total lo cuenta una
    sola vez. Sumar las dos filas infla la glosa; el archivo que se sube al
    DGH quedaría reclamando más de lo que la entidad realmente objetó.

    El renglón que sobrevive es el de mayor valor, y los otros códigos quedan
    anotados en `fusionadas` para que CRDOBSERV los muestre. Nada se pierde.

    Cuando dentro de un grupo se repite el MISMO código de glosa, no es doble
    glosa sino una repetición legítima (dos renglones de verdad de la misma
    cuenta): ahí el bot NO fusiona y avisa, porque juntarlos sí sería borrar
    una objeción.
    """
    grupos: dict[tuple, list[dict]] = defaultdict(list)
    for o in objeciones:
        grupos[(o["cxc"], o["cod_servicio"], o["cantidad"])].append(o)

    conservar: set[int] = set()
    for clave, items in grupos.items():
        if len(items) == 1:
            conservar.add(id(items[0]))
            continue
        codigos = [o["codigo"] for o in items]
        if len(set(codigos)) != len(codigos):
            # Un código repetido: son renglones distintos, no doble glosa.
            if avisar is not None:
                avisar(
                    f"  ⚠ {clave[1]} x{clave[2]} aparece {len(items)} veces con el "
                    f"código {codigos[0]} repetido: NO se fusiona (se cuentan todas). "
                    "Verificá el total contra el acta de MUTUAL."
                )
            conservar.update(id(o) for o in items)
            continue
        ganador = max(items, key=lambda o: o["valor"] or 0)
        ganador["fusionadas"] = [(o["codigo"], o["valor"] or 0) for o in items if o is not ganador]
        conservar.add(id(ganador))
        # Los códigos absorbidos siguen contando para el tipo de la factura.
        ganador["codigos_absorbidos"] = [o["codigo"] for o in items if o is not ganador]

    return [o for o in objeciones if id(o) in conservar]


# ─── Armado de los renglones ─────────────────────────────────────────────────


def construir_filas(
    objeciones: list[dict],
    fecha: _dt.datetime,
    servicios_dgh: dict | None = None,
    trazas: list[dict] | None = None,
) -> list[dict]:
    """Los renglones del formato OBJECIONES, uno por objeción."""
    consec: dict[str, int] = {}
    grupos: dict[str, set[str]] = defaultdict(set)
    for o in objeciones:
        # Los códigos que absorbió la fusión son glosas reales de la factura:
        # cuentan para decidir si es administrativa, médica o mixta.
        for cod in [o["codigo"], *o.get("codigos_absorbidos", [])]:
            grupos[o["cxc"]].add((cod or "")[:2].upper())

    filas: list[dict] = []
    for o in objeciones:
        if o["cxc"] not in consec:
            consec[o["cxc"]] = len(consec) + 1

        slnserpro = ""
        if servicios_dgh is not None:
            unitario = o["valor"] // o["cantidad"] if o["cantidad"] else 0
            cruce = resolver_servicio(
                servicios_dgh.get(o["cxc"], []),
                codigo=o["cod_servicio"],
                descripcion=o["servicio"],
                valor=o["valor"],
                valor_unitario=unitario,
                cantidad=o["cantidad"],
            )
            if cruce.linea is not None:
                slnserpro = cruce.linea.codigo
            if trazas is not None:
                trazas.append(
                    traza(
                        factura=o["cxc"],
                        codigo_objecion=o["codigo"],
                        valor=o["valor"],
                        cod_entidad=o["cod_servicio"],
                        desc_entidad=o["servicio"],
                        unitario_entidad=unitario,
                        observacion=o["observacion"],
                        cruce=cruce,
                    )
                )

        filas.append(
            {
                "CDCONSEC": str(consec[o["cxc"]]),
                "CDFECDOC": fecha,
                "CRNCXC": o["cxc"],
                "CROFECOBJ": fecha,
                "CROREFERE": None,
                "CROOBSERV": None,
                "CROCLAOBJ": CROCLAOBJ_CONST,
                "CRNCLAOBJ": None,
                "GENUSUARIO4": GENUSUARIO4_CONST,
                "CRNCONOBJ": o["codigo"] or None,
                "SLNSERPRO": slnserpro or None,
                "IDRIPS": None,
                # CTNCENCOS va SIEMPRE vacía (regla del área).
                "CTNCENCOS": None,
                "CROVALOBJ": o["valor"],
                "CRDOBSERV": construir_crdobserv(
                    o["codigo"],
                    o["concepto"],
                    o["observacion"],
                    o["valor"],
                    o.get("fusionadas"),
                ),
                "CROTIPOBJ": crotipobj_factura(grupos[o["cxc"]]),
            }
        )
    return filas


# ─── Escritura ───────────────────────────────────────────────────────────────


def escribir_objeciones(filas: list[dict], salida: Path) -> None:
    """Un .xlsx con hoja OBJECIONES y las 16 columnas del formato."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "OBJECIONES"
    fill_hdr = PatternFill("solid", fgColor="1F4E78")
    font_hdr = Font(bold=True, color="FFFFFF")
    for col, nombre in enumerate(COLUMNAS_OBJECIONES, start=1):
        c = ws.cell(row=1, column=col, value=nombre)
        c.fill = fill_hdr
        c.font = font_hdr

    for i, fila in enumerate(filas, start=2):
        for col, nombre in enumerate(COLUMNAS_OBJECIONES, start=1):
            celda = ws.cell(row=i, column=col, value=fila.get(nombre))
            celda.number_format = FORMATOS[nombre]

    ws.freeze_panes = "A2"
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(salida))


def escribir_por_factura(filas: list[dict], carpeta: Path, prefijo: str) -> list[Path]:
    """Un archivo por factura; cada uno reinicia su CDCONSEC en 1."""
    por_factura: dict[str, list[dict]] = defaultdict(list)
    for fila in filas:
        por_factura[fila["CRNCXC"]].append(fila)
    generados = []
    for cxc, grupo in sorted(por_factura.items()):
        destino = carpeta / f"{prefijo}_{cxc}.xlsx"
        escribir_objeciones([{**f, "CDCONSEC": "1"} for f in grupo], destino)
        generados.append(destino)
        logger.info("  %s: %d objeciones", destino.name, len(grupo))
    return generados


# ─── CLI ─────────────────────────────────────────────────────────────────────


def _parse_fecha(texto: str | None) -> _dt.datetime:
    if not texto:
        hoy = _dt.date.today()
        return _dt.datetime(hoy.year, hoy.month, hoy.day)
    try:
        return _dt.datetime.strptime(texto.strip(), "%Y-%m-%d")
    except ValueError:
        logger.error("--fecha inválida: %r (usá YYYY-MM-DD, p. ej. 2026-09-07)", texto)
        sys.exit(2)


PREFIJO_DEFAULT = "OBJECIONES_MUTUAL"


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--entrada", type=Path, required=True, help="Excel de glosas de MUTUAL SER")
    ap.add_argument("--hoja", default=None, help="Hoja del archivo (por defecto 'CONSOLIDADO')")
    ap.add_argument(
        "--salida",
        type=Path,
        required=True,
        help="Carpeta destino, o archivo .xlsx si se usa --consolidado",
    )
    ap.add_argument("--consolidado", action="store_true", help="Un solo Excel con todo")
    ap.add_argument(
        "--prefijo", default=PREFIJO_DEFAULT, help="Prefijo de los archivos por factura"
    )
    ap.add_argument(
        "--servicios-dgh",
        type=Path,
        default=None,
        help="Export de servicios facturados del DGH: con él SLNSERPRO queda con el "
        "código que Dinámica Gerencial reconoce",
    )
    ap.add_argument(
        "--reporte-cruce",
        type=Path,
        default=None,
        help="Excel de trabajo con el detalle del cruce (hojas CRUCE, REVISAR y RESUMEN). "
        "Requiere --servicios-dgh",
    )
    ap.add_argument("--fecha", default=None, help="Fecha de la objeción (YYYY-MM-DD)")
    ap.add_argument("--log", type=Path, default=None, help="Guarda un log adicional a archivo")
    args = ap.parse_args(argv)

    handlers: list[logging.Handler] = [logging.StreamHandler(sys.stdout)]
    if args.log is not None:
        args.log.parent.mkdir(parents=True, exist_ok=True)
        handlers.append(logging.FileHandler(args.log, encoding="utf-8"))
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s", handlers=handlers
    )

    if not args.entrada.is_file():
        logger.error("No existe el archivo de entrada: %s", args.entrada)
        return 1
    if args.reporte_cruce is not None and args.servicios_dgh is None:
        logger.error("--reporte-cruce necesita --servicios-dgh (es el detalle de ese cruce).")
        return 1

    servicios_dgh = None
    if args.servicios_dgh is not None:
        if not args.servicios_dgh.is_file():
            logger.error("No existe el export del DGH: %s", args.servicios_dgh)
            return 1
        logger.info("Leyendo servicios facturados del DGH: %s", args.servicios_dgh.name)
        try:
            servicios_dgh = leer_servicios_dgh(args.servicios_dgh, avisar=logger.warning)
        except ValueError as e:
            logger.error(str(e))
            return 1
        logger.info(
            "  %d renglones de servicio en %d factura(s).",
            sum(len(v) for v in servicios_dgh.values()),
            len(servicios_dgh),
        )

    logger.info("Leyendo glosas de MUTUAL SER: %s", args.entrada.name)
    try:
        objeciones = leer_mutual(args.entrada, args.hoja)
    except ValueError as e:
        logger.error(str(e))
        return 1
    if not objeciones:
        logger.error("No se encontró ninguna objeción en %s", args.entrada.name)
        return 1

    sin_codigo = [o for o in objeciones if not o["codigo"]]
    if sin_codigo:
        logger.warning(
            "  ⚠ %d objeciones sin código de glosa: CRNCONOBJ queda vacío (filas %s).",
            len(sin_codigo),
            ", ".join(str(o["fila_excel"]) for o in sin_codigo[:10]),
        )
    sin_nombre = [o for o in objeciones if not o["servicio"]]
    if sin_nombre:
        logger.info(
            "  %d objeciones sin nombre de servicio en la observación: el cruce se "
            "apoya sólo en el código y el valor (filas %s).",
            len(sin_nombre),
            ", ".join(str(o["fila_excel"]) for o in sin_nombre[:10]),
        )

    bruto = sum(o["valor"] or 0 for o in objeciones)
    antes = len(objeciones)
    objeciones = fusionar_dobles_glosas(objeciones, avisar=logger.warning)
    fusionados = antes - len(objeciones)
    if fusionados:
        neto = sum(o["valor"] or 0 for o in objeciones)
        logger.info(
            "  Doble glosa: %d renglón(es) repetían un servicio ya objetado bajo otro "
            "concepto. Se cuenta una sola vez, por el mayor valor.",
            fusionados,
        )
        logger.info(
            "    sumando todas las filas: $%s  →  glosa real: $%s  (%s menos).",
            f"{bruto:,}",
            f"{neto:,}",
            f"${bruto - neto:,}",
        )
        for o in objeciones:
            for cod, val in o.get("fusionadas", []):
                logger.info(
                    "      %s x%s: se cuenta %s $%s y se absorbe %s $%s",
                    o["cod_servicio"],
                    o["cantidad"],
                    o["codigo"],
                    f"{o['valor']:,}",
                    cod,
                    f"{val:,}",
                )

    trazas: list[dict] = []
    filas = construir_filas(objeciones, _parse_fecha(args.fecha), servicios_dgh, trazas)

    if servicios_dgh is not None:
        cuenta = {
            k: sum(1 for t in trazas if t["confianza"] == k)
            for k in ("ALTA", "MEDIA", "BAJA", "SIN CRUCE")
        }
        ubicados = cuenta["ALTA"] + cuenta["MEDIA"]
        logger.info(
            "  Cruce contra los servicios del DGH: %d de %d servicios ubicados con "
            "confianza alta/media (%.0f%%).",
            ubicados,
            len(trazas),
            ubicados / (len(trazas) or 1) * 100,
        )
        logger.info(
            "    ALTA=%d  MEDIA=%d  BAJA=%d  SIN CRUCE=%d  → revisar %d.",
            cuenta["ALTA"],
            cuenta["MEDIA"],
            cuenta["BAJA"],
            cuenta["SIN CRUCE"],
            cuenta["BAJA"] + cuenta["SIN CRUCE"],
        )

    verificar_reglas(
        [
            {
                "factura": f["CRNCXC"],
                "slnserpro": f["SLNSERPRO"],
                "ctncencos": f["CTNCENCOS"],
                "crotipobj": f["CROTIPOBJ"],
                "codigo_glosa": f["CRNCONOBJ"],
            }
            for f in filas
        ],
        servicios_dgh,
        avisar=logger.info,
    )

    facturas = sorted({o["cxc"] for o in objeciones})
    if args.consolidado:
        escribir_objeciones(filas, args.salida)
        logger.info(
            "\nConsolidado: %d objeciones de %d factura(s) → %s",
            len(filas),
            len(facturas),
            args.salida,
        )
    else:
        generados = escribir_por_factura(filas, args.salida, args.prefijo)
        logger.info("\n%d archivo(s) de MUTUAL en: %s", len(generados), args.salida)

    if args.reporte_cruce is not None:
        escribir_reporte_cruce(trazas, args.reporte_cruce, entidad="MUTUAL SER")
        pendientes = sum(1 for t in trazas if t["aviso"] or t["confianza"] in ("BAJA", "SIN CRUCE"))
        logger.info(
            "Detalle del cruce: %s (%d renglón(es) en la hoja REVISAR).",
            args.reporte_cruce,
            pendientes,
        )

    logger.info(
        "  Facturas: %d  |  Objeciones: %d  |  Valor glosado total: $%s",
        len(facturas),
        len(filas),
        f"{sum(o['valor'] or 0 for o in objeciones):,}",
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
