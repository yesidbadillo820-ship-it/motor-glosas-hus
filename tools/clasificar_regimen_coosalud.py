"""clasificar_regimen_coosalud.py — Clasifica facturas por regimen y cruza fechas RIPS vs factura.

Lee un Excel con la columna FACTURA (ej. "COOSALUD PARA BRAYAN.xlsx"), busca la
carpeta de soportes de cada factura en el share de facturacion electronica
(\\\\172.16.32.83\\factura_electronica_net22\\<AAAAMM>\\FACTURAS_SALUD\\<factura>\\) y:

 1. Identifica el REGIMEN (Subsidiado / Contributivo) leyendo el tipoUsuario del
    RIPS JSON (Res. 1036/2022 - FEV RIPS). Si la carpeta solo trae RIPS viejos
    en TXT (Res. 3374/2000) usa el archivo US; como ultimo recurso busca el dato
    en el XML de la factura (queda marcado "verificar manualmente").
 2. Copia la carpeta COMPLETA de la factura (XML, PDF, RIPS, CUV) a la carpeta
    maestra de su regimen:  <destino>\\Subsidiado\\  o  <destino>\\Contributivo\\.
    Lo que no se pueda clasificar queda en <destino>\\SIN_CLASIFICAR\\.
 3. Busca los SOPORTES DEL SERVICIO de cada factura en las rutas de radicacion
    (Y:\\ y \\\\Prime\\...; ver RUTAS_RADICACION_DEFECTO) y los copia en la
    subcarpeta SOPORTES_RADICACION\\ de la factura.
 4. Extrae la fecha de ATENCION y de EGRESO del RIPS (fechaEgreso de
    hospitalizacion/urgencias manda; si no hay, la ultima atencion) y la fecha
    de INGRESO y EGRESO de la factura. OJO aprendido con la factura real
    HUS349680: el EndDate del InvoicePeriod NO es el egreso clinico — es la
    fecha de facturacion (identica al IssueDate) — asi que el egreso de la
    factura se lee del PDF impreso (fv*.pdf: "Fec Ingreso ... Fec Egreso ..."),
    o del bloque Interoperabilidad si el emisor lo trae; del InvoicePeriod solo
    se toma el StartDate (ingreso). Sin evidencia de egreso, la celda queda
    vacia — no se inventa.
 5. Genera el Excel de auditoria con las columnas:
    Factura | Régimen | RIPS_Atencion | RIPS_Egreso | Factura_Ingreso |
    Factura_Egreso | Alerta_Diferencia (SI/NO) + columnas de apoyo.

USO tipico (PowerShell, desde C:\\temp-notas):

    py tools\\clasificar_regimen_coosalud.py `
        --excel "D:\\USUARIO CARTERA\\Downloads\\COOSALUD PARA BRAYAN.xlsx" `
        --destino "D:\\USUARIO CARTERA\\Documents\\COOSALUD\\PARA_BRAYAN" `
        --piloto 5

    * quitar --piloto para procesar todas las facturas.
    * --solo HUS472660 procesa SOLO esa factura (sin Excel): ideal para probar.
    * --lista facturas.txt procesa las facturas del TXT (una por linea, sin Excel).
    * --lote <nombre> arma las carpetas con el FORMATO DEL CARGUE de COOSALUD:
        <destino>\\<Regimen>\\<lote>\\RIPS\\HUS<n>.json  (+ CUV_HUS<n>.json)
        <destino>\\<Regimen>\\<lote>\\IMG\\HUS<n>\\<soportes del servicio>
      Sin --lote, cada factura queda en su propia carpeta con todo adentro.
    * --sin-copiar genera solo el informe (no copia carpetas ni busca soportes).
    * --sin-soportes procesa y copia la factura electronica pero NO busca en
      las rutas de radicacion (mas rapido).
    * --raiz-soportes "<ruta>" (repetible) reemplaza las rutas de radicacion.
    * --meses 202605,202606 limita la busqueda a esos meses del share.

DEPENDENCIAS:
    py -m pip install openpyxl pymupdf
    (pymupdf es para leer las fechas del PDF de la factura; sin el, esas
     fechas se omiten y se avisa en el log)
"""

from __future__ import annotations

import argparse
import html
import json
import logging
import os
import re
import shutil
import sys
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

logger = logging.getLogger("clasificar_regimen")

DEFAULT_SHARE = r"\\172.16.32.83\factura_electronica_net22"
RE_MES = re.compile(r"^\d{6}$")

# Rutas donde viven los soportes del servicio (PDF, FEV, HEV, etc.) que se
# radican con la factura. Se recorren TODAS y se copia lo que aparezca de cada
# factura (carpeta HUS<n> completa o archivos sueltos con el numero adentro).
RUTAS_RADICACION_DEFECTO: tuple[str, ...] = (
    r"Y:\6. JUNIO 2026 - SOPORTES RADICACION",
    "\\\\Prime\\servidor_radicación\\SERVIDOR RADICACION\\2. SINAC SC SAS - 2026",
    "\\\\Prime\\servidor_radicación\\SERVIDOR RADICACION\\1. SINAC SC SAS - 2025",
    r"\\Prime\radicacion_2026\8. AGOSTO 2026 - SOPORTES RADICACION",
    r"\\Prime\radicacion_2026\9.SEPTIEMBRE - SOPORTES RADICACION",
    r"\\Prime\radicacion_2026\2. FEBRERO 2026 - SOPORTES RADICACION CARPETA 2",
    r"\\Prime\radicacion_2026\3. MARZO 2026 - SOPORTES RADICACION",
    r"\\Prime\radicacion_2026\4. ABRIL 2026 - SOPORTES RADICACION",
    r"\\Prime\radicacion_2026\5. MAYO 2026 - SOPORTES RADICACION",
    r"\\Prime\radicacion_2026\6. JUNIO 2026 - SOPORTES RADICACION",
    r"\\Prime\radicacion_2026\7. JULIO 2026 - SOPORTES RADICACION",
    # Radicacion DIGITAL (portal): aqui viven los soportes electronicos por
    # EPS de los años viejos — Carpeta 2 trae 2024 y 2025 (p.ej.
    # ...\2025\01. ENERO\COOSALUD\...), y la carpeta del servidor de
    # radicacion trae RADICACIÓN 2023 / RADICACION 2024.
    "\\\\Prime\\radicacion_2026\\Radicacion Digital - Carpeta 2\\RADICACION\\RADICACION DIGITAL",
    "\\\\Prime\\servidor_radicación\\RADICACION DIGITAL",
)

# Numero de factura dentro de nombres de carpeta/archivo del share de
# radicacion (FEV_900006037_HUS349680.pdf, carpeta "HUS349680_PEND...", etc.).
_RE_NUM_FACTURA = re.compile(r"HUS\s*0*(\d{4,12})", re.IGNORECASE)

# tipoUsuario del RIPS JSON (Res. 2275/2023, tabla tipoUsuario)
TIPOS_CONTRIBUTIVO = {"01", "02", "03"}
TIPOS_SUBSIDIADO = {"04"}
DESCRIPCION_TIPO = {
    "01": "contributivo cotizante",
    "02": "contributivo beneficiario",
    "03": "contributivo adicional",
    "04": "subsidiado",
}

REGIMEN_SUB = "Subsidiado"
REGIMEN_CON = "Contributivo"
SIN_CLASIFICAR = "SIN_CLASIFICAR"


# ─────────────────────────────────────────────────────────────────────────────
# Utilidades comunes
# ─────────────────────────────────────────────────────────────────────────────


def norm_factura(valor: str) -> str:
    """HUS0000349680 → HUS349680 (tolerante a ceros, guiones y espacios)."""
    v = str(valor or "").strip().upper().replace(" ", "").replace("-", "")
    m = re.match(r"^([A-Z]*)0*(\d+)$", v)
    if m and m.group(2):
        return m.group(1) + m.group(2)
    return v


def clave_numerica(valor: str) -> str:
    """HUS0000349680 → '349680': la llave para cruzar contra nombres del share."""
    m = _RE_NUM_FACTURA.search(str(valor or "").upper())
    if m:
        return m.group(1).lstrip("0") or "0"
    digitos = re.sub(r"\D", "", str(valor or ""))
    return digitos.lstrip("0") or ("0" if digitos else "")


def _decodificar(crudo: bytes) -> str:
    for cod in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return crudo.decode(cod)
        except UnicodeDecodeError:
            continue
    return crudo.decode("utf-8", errors="ignore")


def parse_fecha(valor: object) -> date | None:
    """Acepta date/datetime, '2026-05-01', '2026-05-01 10:30', '01/05/2026'."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    if not texto:
        return None
    m = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", texto)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", texto)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


# ─────────────────────────────────────────────────────────────────────────────
# Indice del share:  <share>\AAAAMM\FACTURAS_SALUD\<factura>\
# ─────────────────────────────────────────────────────────────────────────────


def indexar_share(base: Path, meses: list[str] | None = None) -> dict[str, list[Path]]:
    """Mapa norm_factura → [carpetas encontradas] (puede haber mas de un mes)."""
    indice: dict[str, list[Path]] = {}
    if not base.is_dir():
        raise OSError(f"No se pudo acceder al share {base} (correr desde un PC del hospital).")
    dirs_mes = sorted(
        d
        for d in base.iterdir()
        if d.is_dir() and RE_MES.match(d.name) and (not meses or d.name in meses)
    )
    if not dirs_mes:
        raise OSError(f"El share {base} no tiene carpetas de mes AAAAMM (o --meses no coincide).")
    for mes in dirs_mes:
        contenedores = [
            c
            for c in mes.iterdir()
            if c.is_dir()
            and c.name.upper().startswith("FACTURAS")
            and c.name.upper() != "FACTURAS_NOTA"
        ]
        for cont in contenedores:
            try:
                hijos = list(cont.iterdir())
            except OSError:
                continue
            for carpeta in hijos:
                if not carpeta.is_dir():
                    continue
                clave = norm_factura(carpeta.name)
                if clave:
                    indice.setdefault(clave, []).append(carpeta)
        logger.info(f"  indexado {mes.name} ({len(indice)} facturas acumuladas)")
    return indice


# ─────────────────────────────────────────────────────────────────────────────
# RIPS — JSON nuevo (Res. 1036/2022) y TXT viejo (Res. 3374/2000)
# ─────────────────────────────────────────────────────────────────────────────

_CLAVES_EGRESO = ("fechaEgreso",)
_CLAVES_NO_ATENCION = {"fechaEgreso", "fechaNacimiento"}
_RE_TXT_RIPS = re.compile(r"^(US|AC|AP|AU|AH|AT|AM|AN|AF)\w*\.(txt|TXT)$")


@dataclass
class DatosRips:
    tipos_usuario: list[str] = field(default_factory=list)
    atencion: date | None = None
    egreso: date | None = None
    egreso_derivado: bool = False  # no habia fechaEgreso: se uso la ultima atencion
    fuente: str = ""
    obs: list[str] = field(default_factory=list)


def _es_rips_json(raw: object) -> bool:
    return isinstance(raw, dict) and isinstance(raw.get("usuarios"), list)


def leer_rips_carpeta(carpeta: Path) -> DatosRips:
    """Busca y consolida los RIPS de la carpeta (JSON nuevo o TXT viejo)."""
    datos = DatosRips()
    inicios: list[date] = []
    egresos: list[date] = []

    # --- RIPS JSON (formato FEV vigente) ---
    for ruta in sorted(carpeta.rglob("*.json")):
        try:
            raw = json.loads(_decodificar(ruta.read_bytes()))
        except (OSError, ValueError):
            continue
        if not _es_rips_json(raw):
            continue  # p.ej. ResultadosDoker_*.json (CUV), no es un RIPS
        datos.fuente = datos.fuente or f"JSON {ruta.name}"
        for u in raw.get("usuarios") or []:
            if not isinstance(u, dict):
                continue
            tipo = str(u.get("tipoUsuario") or "").strip().zfill(2)
            if tipo and tipo != "00":
                datos.tipos_usuario.append(tipo)
            servicios = u.get("servicios") or {}
            if not isinstance(servicios, dict):
                continue
            for lista in servicios.values():
                if not isinstance(lista, list):
                    continue
                for s in lista:
                    if not isinstance(s, dict):
                        continue
                    for clave, valor in s.items():
                        if not clave.lower().startswith("fecha"):
                            continue
                        f = parse_fecha(valor)
                        if f is None:
                            continue
                        if clave in _CLAVES_EGRESO:
                            egresos.append(f)
                        elif clave not in _CLAVES_NO_ATENCION:
                            inicios.append(f)

    # --- RIPS TXT viejos (solo si el JSON no aporto nada) ---
    if not datos.fuente:
        for ruta in sorted(carpeta.rglob("*.txt")):
            if not _RE_TXT_RIPS.match(ruta.name):
                continue
            try:
                texto = _decodificar(ruta.read_bytes())
            except OSError:
                continue
            tipo_archivo = ruta.name[:2].upper()
            datos.fuente = "TXT (RIPS Res. 3374/2000)"
            for linea in texto.splitlines():
                if not linea.strip():
                    continue
                if tipo_archivo == "US":
                    campos = linea.split(",")
                    if len(campos) >= 4:
                        t = campos[3].strip()
                        if t == "1":
                            datos.tipos_usuario.append("01")
                        elif t == "2":
                            datos.tipos_usuario.append("04")
                        elif t:
                            datos.tipos_usuario.append(f"txt:{t}")
                    continue
                fechas_linea = [
                    f
                    for f in (parse_fecha(tok) for tok in re.findall(r"[\d/-]{8,10}", linea))
                    if f is not None
                ]
                if not fechas_linea:
                    continue
                if tipo_archivo in ("AH", "AU"):
                    inicios.append(min(fechas_linea))
                    egresos.append(max(fechas_linea))
                else:
                    inicios.extend(fechas_linea)

    if inicios:
        datos.atencion = min(inicios)
    if egresos:
        datos.egreso = max(egresos)
    elif inicios:
        datos.egreso = max(inicios)
        datos.egreso_derivado = True
        datos.obs.append("sin fechaEgreso en RIPS: se uso la ultima fecha de atencion")
    return datos


def regimen_de_tipos(tipos: list[str]) -> tuple[str, str]:
    """(regimen, detalle). Con usuarios mixtos gana la mayoria y se anota."""
    if not tipos:
        return "", ""
    votos: Counter[str] = Counter()
    for t in tipos:
        if t in TIPOS_CONTRIBUTIVO:
            votos[REGIMEN_CON] += 1
        elif t in TIPOS_SUBSIDIADO:
            votos[REGIMEN_SUB] += 1
        else:
            votos[f"Otro ({t})"] += 1
    regimen, _ = votos.most_common(1)[0]
    unicos = sorted(set(tipos))
    detalle = ", ".join(f"{t} {DESCRIPCION_TIPO.get(t, '')}".strip() for t in unicos)
    if len(votos) > 1:
        detalle += " [MIXTO: " + ", ".join(f"{r}×{n}" for r, n in votos.most_common()) + "]"
    return regimen, detalle


# ─────────────────────────────────────────────────────────────────────────────
# XML de la factura electronica (DIAN, sector salud)
# ─────────────────────────────────────────────────────────────────────────────

_RE_PAR_SALUD = re.compile(
    r"<AdditionalInformation>\s*<Name>([^<]+)</Name>\s*<Value[^>]*>(.*?)</Value>",
    re.DOTALL | re.IGNORECASE,
)


@dataclass
class DatosFactura:
    ingreso: date | None = None
    egreso: date | None = None
    fuente: str = ""
    regimen_hint: str = ""
    archivo: str = ""
    obs: list[str] = field(default_factory=list)


def _texto_xml(ruta: Path) -> str:
    """Lee el XML y des-escapa el Invoice embebido (CDATA o entidades)."""
    try:
        contenido = _decodificar(ruta.read_bytes())
    except OSError:
        return ""
    contenido = contenido.replace("<![CDATA[", "").replace("]]>", "")
    if "&lt;" in contenido:
        contenido = html.unescape(contenido)
    return contenido


def _elegir_xml(carpeta: Path) -> Path | None:
    candidatos = sorted(p for p in carpeta.glob("*.xml") if p.is_file())
    if not candidatos:
        candidatos = sorted(p for p in carpeta.rglob("*.xml") if p.is_file())
    for p in candidatos:
        if p.name.lower().startswith("ad"):
            return p
    for p in candidatos:
        if "<Invoice" in _texto_xml(p):
            return p
    return candidatos[0] if candidatos else None


# ── Fechas clinicas desde el PDF impreso de la factura (fv*.pdf) ─────────────
#
# En las facturas del HUS el XML NO trae el egreso clinico (el EndDate del
# InvoicePeriod es la fecha de facturacion). El PDF impreso si lo trae:
#   "Fec Ingreso 07 feb. 2025 11:02 a. m."  /  "Fec Egreso 10 feb. 2025 12:15 p. m."

_MESES_ES = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dic": 12,
}  # fmt: skip
_RE_FEC_PDF = re.compile(
    r"Fec\.?\s*(Ingreso|Egreso)\s*:?\s*(\d{1,2})\s+([a-záéíóú]{3,5})\.?\s+(\d{4})",
    re.IGNORECASE,
)
_aviso_sin_pymupdf = False


def _texto_pdf(ruta: Path) -> str | None:
    """Texto de la primera pagina del PDF. None si no hay lector instalado."""
    global _aviso_sin_pymupdf
    try:
        import pymupdf  # type: ignore[import-not-found]
    except ImportError:
        try:
            import fitz as pymupdf  # type: ignore[import-not-found]  # PyMuPDF viejo
        except ImportError:
            if not _aviso_sin_pymupdf:
                _aviso_sin_pymupdf = True
                logger.warning(
                    "pymupdf no esta instalado: no se leeran las fechas del PDF de la "
                    "factura. Instalalo con: py -m pip install pymupdf"
                )
            return None
    try:
        with pymupdf.open(ruta) as doc:
            return doc[0].get_text() if doc.page_count else ""
    except Exception:
        return ""


def _fechas_desde_texto_pdf(texto: str) -> dict[str, date]:
    """{'ingreso': date, 'egreso': date} con lo que el PDF traiga legible."""
    fechas: dict[str, date] = {}
    for etiqueta, dia, mes_txt, anio in _RE_FEC_PDF.findall(texto or ""):
        mes = _MESES_ES.get(mes_txt.lower().rstrip("."))
        if mes is None:
            continue
        try:
            f = date(int(anio), mes, int(dia))
        except ValueError:
            continue
        clave = etiqueta.lower()
        if clave not in fechas:
            fechas[clave] = f
    return fechas


def _elegir_pdf(carpeta: Path) -> Path | None:
    candidatos = sorted(p for p in carpeta.glob("*.pdf") if p.is_file())
    for p in candidatos:
        if p.name.lower().startswith("fv"):
            return p
    return candidatos[0] if candidatos else None


def analizar_factura(carpeta: Path) -> DatosFactura:
    """Fechas de ingreso/egreso de la factura + pista de regimen.

    Prioridad de fechas: (1) el PDF impreso de la factura, (2) el bloque
    Interoperabilidad del XML, (3) SOLO el StartDate del InvoicePeriod como
    ingreso. El EndDate del InvoicePeriod NUNCA se usa como egreso: en el HUS
    es la fecha de facturacion (= IssueDate), no el egreso clinico."""
    datos = DatosFactura()

    # 1) PDF impreso de la factura.
    pdf = _elegir_pdf(carpeta)
    if pdf is not None:
        texto_pdf = _texto_pdf(pdf)
        if texto_pdf:
            fechas = _fechas_desde_texto_pdf(texto_pdf)
            if fechas:
                datos.ingreso = fechas.get("ingreso")
                datos.egreso = fechas.get("egreso")
                datos.fuente = "PDF factura"
                datos.archivo = pdf.name

    # 2) XML: Interoperabilidad + InvoicePeriod (solo ingreso) + pista regimen.
    ruta = _elegir_xml(carpeta)
    if ruta is None:
        if datos.fuente == "":
            datos.obs.append("sin XML en la carpeta")
        return datos
    datos.archivo = datos.archivo or ruta.name
    texto = _texto_xml(ruta)
    if texto:
        pares = {
            n.strip().upper(): re.sub(r"\s+", " ", v).strip()
            for n, v in _RE_PAR_SALUD.findall(texto)
        }
        uso_interop = False
        for nombre, valor in pares.items():
            if "FECHA" not in nombre:
                continue
            f = parse_fecha(valor)
            if f is None:
                continue
            if datos.ingreso is None and ("INGRESO" in nombre or "INICIO" in nombre):
                datos.ingreso = f
                uso_interop = True
            if datos.egreso is None and ("EGRESO" in nombre or "FIN" in nombre):
                datos.egreso = f
                uso_interop = True
        if uso_interop:
            datos.fuente = (datos.fuente + "+Interoperabilidad").lstrip("+")

        if datos.ingreso is None:
            m = re.search(
                r"<(?:\w+:)?InvoicePeriod>(.*?)</(?:\w+:)?InvoicePeriod>",
                texto,
                re.DOTALL | re.IGNORECASE,
            )
            if m:
                inicio = re.search(r"<(?:\w+:)?StartDate>([^<]+)<", m.group(1), re.IGNORECASE)
                if inicio:
                    datos.ingreso = parse_fecha(inicio.group(1))
                    if datos.ingreso is not None:
                        datos.fuente = (datos.fuente + "+InvoicePeriod (solo ingreso)").lstrip("+")

        for nombre, valor in pares.items():
            if "REGIMEN" in nombre or "TIPO_USUARIO" in nombre:
                datos.regimen_hint = valor
                break
        if not datos.regimen_hint:
            mayus = texto.upper()
            if "SUBSIDIADO" in mayus and "CONTRIBUTIVO" not in mayus:
                datos.regimen_hint = "SUBSIDIADO"
            elif "CONTRIBUTIVO" in mayus and "SUBSIDIADO" not in mayus:
                datos.regimen_hint = "CONTRIBUTIVO"

    if datos.egreso is None:
        datos.obs.append(
            "la factura no trae egreso clinico legible (el EndDate del XML es la "
            "fecha de facturacion, no se usa)"
        )
    return datos


def regimen_de_hint(hint: str) -> str:
    h = hint.strip().upper()
    if not h:
        return ""
    if "SUBSIDIADO" in h or norm_factura(h) in ("2", "04", "4"):
        return REGIMEN_SUB
    if "CONTRIBUTIVO" in h or norm_factura(h) in ("1", "01", "02", "03", "2", "3"):
        return REGIMEN_CON
    return ""


# ─────────────────────────────────────────────────────────────────────────────
# Soportes del servicio en las rutas de radicacion (Y:\ y \\Prime\...)
# ─────────────────────────────────────────────────────────────────────────────


# En los PC del hospital las unidades Y: (\\Prime\radicacion_2026) y
# X: (\\Prime\servidor_radicación) son mapeos con credenciales guardadas.
# Si una ruta UNC no abre, se prueba su equivalente por letra — y al
# reves — antes de darla por inaccesible.
_EQUIV_RADICACION: tuple[tuple[str, str], ...] = (
    (r"\\Prime\radicacion_2026", "Y:"),
    ("\\\\Prime\\servidor_radicación", "X:"),
)


def _canon_raiz(ruta: str) -> str:
    """Forma canonica para no escanear dos veces la misma carpeta
    (Y:\\3. MARZO... y \\\\Prime\\radicacion_2026\\3. MARZO... son la misma)."""
    low = str(ruta).lower().rstrip("\\/")
    for unc, letra in _EQUIV_RADICACION:
        if low.startswith(unc.lower()):
            low = letra.lower() + low[len(unc) :]
    return low


def resolver_raices(rutas: list[str]) -> list[Path]:
    """Devuelve las raices ACCESIBLES, sin duplicados (Y:/X: ≡ \\\\Prime\\...).

    Para cada ruta prueba la escrita y su equivalente mapeada; avisa las que
    no abren por ninguna de las dos."""
    accesibles: list[Path] = []
    vistas: set[str] = set()
    for original in rutas:
        clave = _canon_raiz(original)
        if clave in vistas:
            continue
        vistas.add(clave)
        candidatas = [str(original)]
        low = str(original).lower()
        for unc, letra in _EQUIV_RADICACION:
            if low.startswith(unc.lower()):
                candidatas.append(letra + str(original)[len(unc) :])
            elif low.startswith(letra.lower()):
                candidatas.append(unc + str(original)[len(letra) :])
        elegida: Path | None = None
        for cand in candidatas:
            p = Path(cand)
            if p.is_dir():
                elegida = p
                break
        if elegida is None:
            logger.warning(f"  ruta de soportes NO accesible (se omite): {original}")
        else:
            if str(elegida) != str(original):
                logger.info(f"  ruta {original} no abrio: se usa su equivalente {elegida}")
            accesibles.append(elegida)
    return accesibles


def indexar_radicacion(raices: list[Path], objetivos: set[str]) -> dict[str, list[Path]]:
    """Recorre cada raiz UNA vez y devuelve {clave_numerica: [hallazgos]}.

    Un hallazgo es una carpeta nombrada con la factura (se copia completa) o un
    archivo suelto con el numero en el nombre (FEV_..._HUS349680.pdf). Poda: no
    desciende a carpetas de OTRAS facturas (las mas numerosas del share) ni a
    las carpetas objetivo ya encontradas. Un hilo por raiz: el costo es
    latencia de red, no CPU."""
    hallados: dict[str, list[Path]] = {k: [] for k in objetivos}

    def _explorar(raiz: Path) -> tuple[Path, int, int]:
        vistos = 0
        hits = 0
        for root, dirs, files in os.walk(str(raiz), onerror=lambda _e: None):
            vistos += 1
            if vistos % 2000 == 0:
                logger.info(
                    f"    ... {vistos} carpetas revisadas en {raiz.name} ({hits} hallazgos)"
                )
            conservar = []
            for d in dirs:
                m = _RE_NUM_FACTURA.search(d)
                if m:
                    num = m.group(1).lstrip("0") or "0"
                    if num in objetivos:
                        hallados[num].append(Path(root) / d)
                        hits += 1
                    continue  # carpeta de factura (objetivo o no): no descender
                conservar.append(d)
            dirs[:] = conservar
            for fn in files:
                m = _RE_NUM_FACTURA.search(fn)
                if not m:
                    continue
                num = m.group(1).lstrip("0") or "0"
                if num in objetivos:
                    hallados[num].append(Path(root) / fn)
                    hits += 1
        return raiz, vistos, hits

    with ThreadPoolExecutor(max_workers=max(1, min(len(raices), 12))) as pool:
        for raiz, vistos, hits in pool.map(_explorar, raices):
            if hits == 0:
                logger.warning(
                    f"  soportes: {raiz} — {vistos} carpetas revisadas y NINGUN hallazgo "
                    f"(¿la ruta es la correcta para estas facturas?)"
                )
            else:
                logger.info(f"  soportes: {raiz} — {vistos} carpetas revisadas, {hits} hallazgos")
    return hallados


def armar_radicacion(
    factura: str,
    carpeta_fe: Path | None,
    soportes: list[Path],
    lote_dir: Path,
) -> tuple[int, list[str]]:
    """Arma la carpeta como la pide el cargue de COOSALUD:

        <lote>\\RIPS\\HUS<n>.json          (el RIPS de la factura)
        <lote>\\RIPS\\CUV_HUS<n>.json      (el resultado de validacion MinSalud)
        <lote>\\IMG\\HUS<n>\\*.pdf          (los soportes del servicio)

    Los RIPS van PLANOS en una sola carpeta (renombrados al nombre que espera
    el portal) y cada factura tiene su subcarpeta dentro de IMG. Devuelve
    (archivos copiados, observaciones)."""
    obs: list[str] = []
    copiados = 0
    fac = norm_factura(factura) or factura
    rips_dir = lote_dir / "RIPS"
    img_dir = lote_dir / "IMG" / fac

    # --- RIPS y CUV desde la carpeta de facturacion electronica ---
    if carpeta_fe is not None:
        for ruta in sorted(carpeta_fe.rglob("*.json")):
            try:
                raw = json.loads(_decodificar(ruta.read_bytes()))
            except (OSError, ValueError):
                continue
            if _es_rips_json(raw):
                destino_json = rips_dir / f"{fac}.json"
            elif "resultadosdoker" in ruta.name.lower() or "cuv" in ruta.name.lower():
                destino_json = rips_dir / f"CUV_{fac}.json"
            else:
                continue
            try:
                rips_dir.mkdir(parents=True, exist_ok=True)
                shutil.copy2(ruta, destino_json)
                copiados += 1
            except OSError as exc:
                obs.append(f"NO se pudo copiar {ruta.name}: {exc}")

    # --- Soportes del servicio a IMG\HUS<n>\ (aplanados: el portal no
    #     entiende subcarpetas dentro de la carpeta de la factura) ---
    for hit in soportes:
        archivos = sorted(p for p in hit.rglob("*") if p.is_file()) if hit.is_dir() else [hit]
        for a in archivos:
            try:
                img_dir.mkdir(parents=True, exist_ok=True)
                shutil.copy2(a, img_dir / a.name)
                copiados += 1
            except OSError as exc:
                obs.append(f"NO se pudo copiar soporte {a.name}: {exc}")
    return copiados, obs


def copiar_soportes(soportes: list[Path], destino_fac: Path) -> tuple[int, list[str]]:
    """Copia los hallazgos a <destino_fac>\\SOPORTES_RADICACION\\. Devuelve
    (archivos copiados, observaciones de fallos)."""
    obs: list[str] = []
    sop_dir = destino_fac / "SOPORTES_RADICACION"
    for hit in soportes:
        try:
            if hit.is_dir():
                shutil.copytree(hit, sop_dir / hit.name, dirs_exist_ok=True)
            else:
                sop_dir.mkdir(parents=True, exist_ok=True)
                shutil.copy2(hit, sop_dir / hit.name)
        except OSError as exc:
            obs.append(f"NO se pudo copiar soporte {hit.name}: {exc}")
    copiados = sum(1 for p in sop_dir.rglob("*") if p.is_file()) if sop_dir.is_dir() else 0
    return copiados, obs


# ─────────────────────────────────────────────────────────────────────────────
# Proceso por factura
# ─────────────────────────────────────────────────────────────────────────────


@dataclass
class Resultado:
    factura: str
    regimen: str = ""
    rips_atencion: date | None = None
    rips_egreso: date | None = None
    fact_ingreso: date | None = None
    fact_egreso: date | None = None
    alerta: str = "SIN DATOS"
    detalle_dif: str = ""
    tipos_usuario: str = ""
    fuente_regimen: str = ""
    fuente_fechas_xml: str = ""
    origen: str = ""
    destino: str = ""
    copiados: int = 0
    soportes_origen: str = ""
    soportes_copiados: int = 0
    obs: list[str] = field(default_factory=list)


def _comparar_fechas(r: Resultado) -> None:
    fechas = (r.rips_atencion, r.rips_egreso, r.fact_ingreso, r.fact_egreso)
    if any(f is None for f in fechas):
        r.alerta = "SIN DATOS"
        faltan = [
            n
            for n, f in zip(
                ("RIPS_Atencion", "RIPS_Egreso", "Factura_Ingreso", "Factura_Egreso"), fechas
            )
            if f is None
        ]
        r.detalle_dif = "faltan: " + ", ".join(faltan)
        return
    difs = []
    if r.rips_atencion != r.fact_ingreso:
        dias = (r.fact_ingreso - r.rips_atencion).days  # type: ignore[operator]
        difs.append(f"atencion/ingreso difieren {abs(dias)} dia(s)")
    if r.rips_egreso != r.fact_egreso:
        dias = (r.fact_egreso - r.rips_egreso).days  # type: ignore[operator]
        difs.append(f"egreso difiere {abs(dias)} dia(s)")
    r.alerta = "SI" if difs else "NO"
    r.detalle_dif = "; ".join(difs)


def procesar_factura(
    factura: str,
    carpetas: list[Path],
    destino: Path,
    copiar: bool,
    soportes: list[Path] | None = None,
    con_soportes: bool = False,
    lote: str = "",
) -> Resultado:
    r = Resultado(factura=factura)
    soportes = soportes or []
    if soportes:
        r.soportes_origen = "; ".join(str(h) for h in soportes[:5]) + (
            f" (+{len(soportes) - 5} mas)" if len(soportes) > 5 else ""
        )
    if not carpetas:
        r.regimen = "NO_ENCONTRADA"
        r.obs.append("la factura no aparece en el share de facturacion electronica")
        # Aunque no este la factura electronica, los soportes de radicacion
        # que si aparecieron se guardan para no perderlos.
        if copiar and soportes:
            if lote:
                lote_dir = destino / SIN_CLASIFICAR / lote
                r.soportes_copiados, obs_sop = armar_radicacion(factura, None, soportes, lote_dir)
                r.destino = str(lote_dir / "IMG" / (norm_factura(factura) or factura))
            else:
                destino_fac = destino / SIN_CLASIFICAR / (norm_factura(factura) or factura)
                r.soportes_copiados, obs_sop = copiar_soportes(soportes, destino_fac)
                r.destino = str(destino_fac)
            r.obs.extend(obs_sop)
        return r
    # Si esta en varios meses se usa la carpeta del mes mas reciente.
    carpeta = sorted(carpetas, key=lambda p: str(p))[-1]
    r.origen = str(carpeta)
    if len(carpetas) > 1:
        r.obs.append(f"aparece en {len(carpetas)} carpetas del share; se uso la mas reciente")

    rips = leer_rips_carpeta(carpeta)
    fe = analizar_factura(carpeta)

    r.rips_atencion, r.rips_egreso = rips.atencion, rips.egreso
    r.fact_ingreso, r.fact_egreso = fe.ingreso, fe.egreso
    r.tipos_usuario = ", ".join(sorted(set(rips.tipos_usuario)))
    r.fuente_fechas_xml = fe.fuente or ("sin fechas en " + fe.archivo if fe.archivo else "sin XML")
    r.obs.extend(rips.obs)
    r.obs.extend(fe.obs)

    regimen, detalle = regimen_de_tipos(rips.tipos_usuario)
    if regimen in (REGIMEN_SUB, REGIMEN_CON):
        r.regimen = regimen
        r.fuente_regimen = rips.fuente
        if detalle:
            r.obs.append("tipoUsuario: " + detalle)
    elif regimen:
        r.regimen = regimen
        r.fuente_regimen = rips.fuente
        r.obs.append("tipoUsuario fuera de contributivo/subsidiado: " + detalle)
    else:
        desde_xml = regimen_de_hint(fe.regimen_hint)
        if desde_xml:
            r.regimen = desde_xml
            r.fuente_regimen = "XML (verificar manualmente)"
            r.obs.append(f"regimen tomado del XML ('{fe.regimen_hint}'): sin RIPS legible")
        else:
            r.regimen = SIN_CLASIFICAR
            r.obs.append("sin RIPS legible ni pista de regimen en el XML")

    _comparar_fechas(r)

    if copiar:
        subcarpeta = r.regimen if r.regimen in (REGIMEN_SUB, REGIMEN_CON) else SIN_CLASIFICAR
        if lote:
            # Formato del cargue de COOSALUD: <regimen>\<lote>\RIPS + IMG\HUS<n>
            lote_dir = destino / subcarpeta / lote
            r.copiados, obs_rad = armar_radicacion(factura, carpeta, soportes, lote_dir)
            r.soportes_copiados = sum(
                1 for p in (lote_dir / "IMG" / carpeta.name).rglob("*") if p.is_file()
            )
            r.destino = str(lote_dir / "IMG" / carpeta.name)
            r.obs.extend(obs_rad)
        else:
            destino_fac = destino / subcarpeta / carpeta.name
            try:
                shutil.copytree(carpeta, destino_fac, dirs_exist_ok=True)
                r.destino = str(destino_fac)
                r.copiados = sum(1 for p in destino_fac.rglob("*") if p.is_file())
            except OSError as exc:
                r.obs.append(f"NO se pudo copiar: {exc}")
            if soportes:
                r.soportes_copiados, obs_sop = copiar_soportes(soportes, destino_fac)
                r.obs.extend(obs_sop)
                # El conteo de la carpeta de la factura ya incluye los soportes.
                r.copiados = sum(1 for p in destino_fac.rglob("*") if p.is_file())
        if not soportes and con_soportes:
            r.obs.append("sin soportes en las rutas de radicacion")
    return r


# ─────────────────────────────────────────────────────────────────────────────
# Entrada (Excel de facturas) y salida (Excel de auditoria)
# ─────────────────────────────────────────────────────────────────────────────


def leer_facturas_lista(ruta: Path) -> list[str]:
    """TXT con una factura por linea. Tolera el BOM UTF-16/UTF-8 que dejan
    `Out-File` y `>` de PowerShell, comillas envolventes, lineas vacias y
    comentarios (#). Deduplica conservando el orden."""
    raw = ruta.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        texto = raw.decode("utf-16")
    elif raw.startswith(b"\xef\xbb\xbf"):
        texto = raw.decode("utf-8-sig")
    else:
        try:
            texto = raw.decode("utf-8")
        except UnicodeDecodeError:
            texto = raw.decode("latin-1")
    facturas: list[str] = []
    vistas: set[str] = set()
    for linea in texto.splitlines():
        # El ﻿ suelto aparece cuando el TXT ya traia BOM y PowerShell le
        # agrego otro al reescribirlo: si no se quita, la factura no cruza.
        s = linea.replace("﻿", "").strip().strip('"').strip("'")
        if not s or s.startswith("#"):
            continue
        clave = norm_factura(s)
        if not clave or clave in vistas:
            continue
        vistas.add(clave)
        facturas.append(s)
    return facturas


def leer_facturas_excel(ruta: Path, hoja: str | None, columna: str) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        nombre = hoja or wb.sheetnames[0]
        if nombre not in wb.sheetnames:
            objetivo = nombre.strip().casefold()
            real = next((s for s in wb.sheetnames if s.strip().casefold() == objetivo), None)
            if real is None:
                raise ValueError(f"El Excel no tiene la hoja '{nombre}'. Hojas: {wb.sheetnames}")
            nombre = real
        ws = wb[nombre]
        filas = ws.iter_rows(values_only=True)
        encabezado = next(filas, None)
        if not encabezado:
            raise ValueError("La hoja esta vacia.")
        objetivo = columna.strip().casefold()
        idx = next(
            (i for i, c in enumerate(encabezado) if str(c or "").strip().casefold() == objetivo),
            None,
        )
        if idx is None:
            raise ValueError(f"No encontre la columna '{columna}'. Encabezados: {list(encabezado)}")
        facturas: list[str] = []
        vistas: set[str] = set()
        for fila in filas:
            valor = str(fila[idx] or "").strip() if idx < len(fila) else ""
            if not valor:
                continue
            clave = norm_factura(valor)
            if clave in vistas:
                continue
            vistas.add(clave)
            facturas.append(valor)
        return facturas
    finally:
        wb.close()


ENCABEZADOS = (
    "Factura",
    "Régimen",
    "RIPS_Atencion",
    "RIPS_Egreso",
    "Factura_Ingreso",
    "Factura_Egreso",
    "Alerta_Diferencia",
    "Detalle_Diferencia",
    "Tipo_Usuario_RIPS",
    "Fuente_Regimen",
    "Fuente_Fechas_Factura",
    "Archivos_Copiados",
    "Soportes_Radicacion",
    "Archivos_Soportes",
    "Carpeta_Origen",
    "Carpeta_Destino",
    "Observaciones",
)


def escribir_auditoria(resultados: list[Resultado], salida: Path) -> Path:
    """Guarda el Excel y devuelve la ruta real donde quedo.

    Si `salida` esta abierta en Excel (PermissionError de Windows), NO se
    pierde la corrida: se guarda con un sufijo de hora junto al original."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "AUDITORIA"
    ws.append(list(ENCABEZADOS))
    for celda in ws[1]:
        celda.font = Font(bold=True)
    rojo = PatternFill("solid", start_color="FFC7CE")
    verde = PatternFill("solid", start_color="C6EFCE")
    for r in resultados:
        ws.append(
            [
                r.factura,
                r.regimen,
                r.rips_atencion,
                r.rips_egreso,
                r.fact_ingreso,
                r.fact_egreso,
                r.alerta,
                r.detalle_dif,
                r.tipos_usuario,
                r.fuente_regimen,
                r.fuente_fechas_xml,
                r.copiados or "",
                r.soportes_origen,
                r.soportes_copiados or "",
                r.origen,
                r.destino,
                " | ".join(r.obs),
            ]
        )
        fila = ws.max_row
        for col in ("C", "D", "E", "F"):
            ws[f"{col}{fila}"].number_format = "DD/MM/YYYY"
        celda_alerta = ws[f"G{fila}"]
        if r.alerta == "SI":
            celda_alerta.fill = rojo
        elif r.alerta == "NO":
            celda_alerta.fill = verde
    anchos = (16, 14, 13, 13, 14, 14, 16, 34, 18, 22, 22, 16, 52, 16, 52, 52, 60)
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(ENCABEZADOS)).column_letter}{ws.max_row}"
    salida.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(salida)
        return salida
    except PermissionError:
        alterna = salida.with_name(f"{salida.stem}_{datetime.now():%H%M%S}{salida.suffix}")
        logger.warning(
            f"'{salida.name}' esta abierto en Excel y Windows lo bloquea: "
            f"guardo el informe como '{alterna.name}' para no perder la corrida."
        )
        wb.save(alterna)
        return alterna


# ─────────────────────────────────────────────────────────────────────────────
# main
# ─────────────────────────────────────────────────────────────────────────────


def main() -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--excel", type=Path, default=None, help="Excel con la columna FACTURA.")
    parser.add_argument(
        "--solo",
        type=str,
        default="",
        help="Procesar SOLO estas facturas por coma (p.ej. --solo HUS472660). "
        "Con este flag el Excel es opcional: sirve para probar una factura puntual.",
    )
    parser.add_argument(
        "--lista",
        type=Path,
        default=None,
        help="TXT con una factura por linea (alternativa al Excel; tolera el BOM "
        "de PowerShell, comillas y lineas vacias).",
    )
    parser.add_argument(
        "--lote",
        type=str,
        default="",
        nargs="?",
        const="auto",
        help="Arma las carpetas con el FORMATO DEL CARGUE de COOSALUD: "
        "<destino>\\<Regimen>\\<lote>\\RIPS\\HUS<n>.json + CUV_HUS<n>.json y "
        "<lote>\\IMG\\HUS<n>\\<soportes>. Se le puede dar el nombre del lote "
        "(p.ej. --lote 605505_20260908_135357); sin valor usa la fecha y hora.",
    )
    parser.add_argument("--hoja", type=str, default=None, help="Hoja a leer (default: la primera).")
    parser.add_argument("--columna", type=str, default="FACTURA", help="Columna de facturas.")
    parser.add_argument(
        "--share", type=Path, default=Path(DEFAULT_SHARE), help="Raiz del share de facturacion."
    )
    parser.add_argument(
        "--destino",
        type=Path,
        default=Path(r"D:\USUARIO CARTERA\Documents\COOSALUD\CLASIFICADO_REGIMEN"),
        help="Carpeta donde crear Subsidiado\\ y Contributivo\\.",
    )
    parser.add_argument(
        "--salida", type=Path, default=None, help="Excel de auditoria (default: en --destino)."
    )
    parser.add_argument(
        "--meses", type=str, default="", help="Limitar la busqueda a meses AAAAMM (por coma)."
    )
    parser.add_argument("--piloto", type=int, default=0, help="Procesar solo las primeras N.")
    parser.add_argument(
        "--sin-copiar", action="store_true", help="Solo el informe: no copia carpetas."
    )
    parser.add_argument(
        "--sin-soportes",
        action="store_true",
        help="No buscar soportes en las rutas de radicacion (mas rapido).",
    )
    parser.add_argument(
        "--raiz-soportes",
        action="append",
        default=None,
        metavar="RUTA",
        help="Ruta de radicacion donde buscar soportes (repetible; reemplaza las default).",
    )
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.stderr.write("ERROR: falta openpyxl. Instalalo con: py -m pip install openpyxl\n")
        return 2

    if args.solo.strip():
        facturas = [f.strip() for f in args.solo.split(",") if f.strip()]
        logger.info(f"SOLO estas facturas (sin leer Excel): {', '.join(facturas)}")
    elif args.lista is not None:
        if not args.lista.is_file():
            logger.error(f"No existe la lista: {args.lista}")
            return 1
        facturas = leer_facturas_lista(args.lista)
        if not facturas:
            logger.error(f"La lista esta vacia: {args.lista}")
            return 1
        logger.info(f"Lista {args.lista.name}: {len(facturas)} facturas (sin leer Excel)")
    else:
        if args.excel is None:
            logger.error("Falta --excel (o --lista <txt>, o --solo HUS<n> para pruebas).")
            return 1
        if not args.excel.is_file():
            logger.error(f"No existe el Excel: {args.excel}")
            return 1
        try:
            facturas = leer_facturas_excel(args.excel, args.hoja, args.columna)
        except ValueError as exc:
            logger.error(str(exc))
            return 1
    if args.piloto > 0:
        facturas = facturas[: args.piloto]
        logger.info(f"PILOTO: solo las primeras {len(facturas)} facturas")
    logger.info(f"Facturas a procesar: {len(facturas)}")

    meses = [m.strip() for m in args.meses.split(",") if m.strip()] or None
    logger.info(f"Indexando el share {args.share} ...")
    try:
        indice = indexar_share(args.share, meses)
    except OSError as exc:
        logger.error(str(exc))
        return 1
    logger.info(f"Indice listo: {len(indice)} carpetas de factura en el share.")

    # Soportes del servicio en las rutas de radicacion (solo si se va a copiar).
    soportes_idx: dict[str, list[Path]] = {}
    buscar_soportes = not args.sin_copiar and not args.sin_soportes
    if buscar_soportes:
        accesibles = resolver_raices(list(args.raiz_soportes or RUTAS_RADICACION_DEFECTO))
        if accesibles:
            logger.info(
                f"Buscando soportes de radicacion de {len(facturas)} facturas "
                f"en {len(accesibles)} rutas (esto puede tardar)..."
            )
            soportes_idx = indexar_radicacion(accesibles, {clave_numerica(f) for f in facturas})
        else:
            buscar_soportes = False
            logger.warning("Ninguna ruta de soportes accesible: se continua sin soportes.")

    lote = args.lote.strip()
    if lote == "auto":
        lote = datetime.now().strftime("%Y%m%d_%H%M%S")
    if lote:
        logger.info(
            f"Formato CARGUE COOSALUD: las carpetas quedan como "
            f"<Regimen>\\{lote}\\RIPS y <Regimen>\\{lote}\\IMG\\HUS<n>"
        )

    salida = args.salida or (args.destino / "AUDITORIA_FECHAS_REGIMEN.xlsx")
    resultados: list[Resultado] = []
    for i, fac in enumerate(facturas, 1):
        r = procesar_factura(
            fac,
            indice.get(norm_factura(fac), []),
            args.destino,
            not args.sin_copiar,
            soportes=soportes_idx.get(clave_numerica(fac), []),
            con_soportes=buscar_soportes,
            lote=lote,
        )
        resultados.append(r)
        rango_rips = (
            f"{r.rips_atencion:%d/%m}–{r.rips_egreso:%d/%m}"
            if r.rips_atencion and r.rips_egreso
            else "—"
        )
        rango_fe = (
            f"{r.fact_ingreso:%d/%m}–{r.fact_egreso:%d/%m}"
            if r.fact_ingreso and r.fact_egreso
            else "—"
        )
        sop = f" | sop {r.soportes_copiados}" if r.soportes_copiados else ""
        logger.info(
            f"[{i}/{len(facturas)}] {fac} → {r.regimen or '?'} | RIPS {rango_rips} | FE {rango_fe} | alerta {r.alerta}{sop}"
        )

    salida_real = escribir_auditoria(resultados, salida)

    conteo = Counter(r.regimen for r in resultados)
    alertas = sum(1 for r in resultados if r.alerta == "SI")
    logger.info("")
    logger.info("========== RESUMEN ==========")
    for regimen, n in conteo.most_common():
        logger.info(f"  {regimen}: {n}")
    logger.info(f"  Con diferencia de fechas (Alerta SI): {alertas}")
    if buscar_soportes:
        con_sop = sum(1 for r in resultados if r.soportes_copiados)
        logger.info(f"  Con soportes de radicacion hallados: {con_sop} de {len(resultados)}")
    logger.info(f"  Excel de auditoria: {salida_real}")
    if not args.sin_copiar:
        logger.info(f"  Carpetas copiadas bajo: {args.destino}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
