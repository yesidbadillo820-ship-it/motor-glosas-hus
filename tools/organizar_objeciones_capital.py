"""organizar_objeciones_capital.py — Objeciones de CAPITAL SALUD → Excel OBJECIONES (DGH).

Toma el Excel de glosas que entrega **CAPITAL SALUD** (6 columnas) y lo
convierte al formato de trabajo de 16 columnas (hoja `OBJECIONES`) que se
importa en Dinámica Gerencial, el mismo de FAMISANAR, SANITAS, MUTUAL y el
Dispensario.

CÓMO VIENE EL ARCHIVO DE CAPITAL SALUD:

    FACTURA           → HUS0000532847
    servicio          → el NOMBRE del servicio, con guiones en vez de espacios
                        ("INTERNACION-ADULTOS-COMPLEJIDAD-ALTA-HABITACION-MULTIPLE")
    Cantidad          → cantidad
    DescripcionGlosa  → el CÓDIGO y su texto, pegados ("SO4201 - Existe ausencia…")
    OBSERVACION       → lo que escribió el auditor de la EPS
    VALOR GLOSA       → valor objetado

DOS PARTICULARIDADES, Y DE DÓNDE SALE CADA DATO:

1. **No manda el código del servicio**, sólo el nombre — como SALUD TOTAL. El
   cruce contra el DGH se apoya entonces en el NOMBRE y el VALOR; sin cruce
   confiable la celda queda vacía y el renglón va a REVISAR. Nunca se inventa.

2. **No manda el código de glosa en columna propia**: va al principio de
   `DescripcionGlosa`, antes del guion. `codigo_y_concepto()` lo separa:

       "SO4201 - Existe ausencia total, parcial o inconsistencia de la lista
        de precios"  →  ("SO4201", "Existe ausencia total, parcial o …")

   Si esa celda no empieza por un código reconocible, `CRNCONOBJ` queda vacío
   y el bot lo avisa, en vez de inventar uno.

LOS GUIONES DEL NOMBRE. CAPITAL escribe los nombres con guion donde va el
espacio. Se cambian por espacios al leer, que es como los tiene el DGH
("INTERNACION ADULTOS COMPLEJIDAD ALTA HABITACION MULTIPLE") y como se leen
mejor en el archivo de salida.

RENGLONES REPETIDOS. CAPITAL repite el mismo servicio tantas veces como
renglones tenga la factura (cuatro hemogramas, tres terapias respiratorias…).
Se comprobó contra el export del DGH que son renglones de verdad, así que
**se conservan todos**: juntarlos borraría objeciones reales.

REGLAS FIJAS DEL FORMATO (ver CLAUDE.md, valen para todas las entidades):
    CTNCENCOS  → siempre vacía
    CROTIPOBJ  → por factura: 0 = ADMINISTRATIVA, 1 = MEDICA, 2 = MIXTA
    SLNSERPRO  → prohibido inventar: sin cruce confiable, celda vacía
    El archivo lleva el 100% de los renglones; los no cruzados van con la
    celda vacía para completarlos a mano.

USO:
    py tools\\organizar_objeciones_capital.py ^
        --entrada       "D:\\...\\CAPITAL_SALUD_9_SEPTIEMBRE.xlsx" ^
        --servicios-dgh "D:\\...\\SERVICIOS_FACTURADOS_DGH.xlsx" ^
        --salida        "D:\\...\\OBJECIONES_CAPITAL_09092026.xlsx" --consolidado ^
        --reporte-cruce "D:\\...\\CRUCE_CAPITAL_09092026.xlsx" ^
        --fecha 2026-09-09

INSTALACIÓN (una vez):  py -m pip install openpyxl
"""

from __future__ import annotations

import argparse
import datetime as _dt
import logging
import re
import sys
from collections import defaultdict
from pathlib import Path

# Motor del cruce y lector de pesos, comunes a todos los bots.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from _cruce_dgh import (  # noqa: E402
    escribir_reporte_cruce,
    factura_larga,
    leer_servicios_dgh,
    norm_header,
    resolver_servicio,
    traza,
    verificar_reglas,
)
from _dinero import a_entero  # noqa: E402

logger = logging.getLogger("organizar_capital")


# ─── Formato de salida: layout de 16 columnas (hoja OBJECIONES) ──────────────

COLUMNAS_OBJECIONES: tuple[str, ...] = (
    "CDCONSEC",
    "CDFECDOC",
    "CRNCXC",
    "CROFECOBJ",
    "CROREFERE",
    "CROOBSERV",
    "CROCLAOBJ",
    "CRNCLAOBJ",
    "GENUSUARIO4",
    "CRNCONOBJ",
    "SLNSERPRO",
    "IDRIPS",
    "CTNCENCOS",
    "CROVALOBJ",
    "CRDOBSERV",
    "CROTIPOBJ",
)

# number_format por columna, igual que en los archivos reales.
FORMATOS: dict[str, str] = {
    "CDCONSEC": "@",
    "CDFECDOC": "mm-dd-yy",
    "CRNCXC": "@",
    "CROFECOBJ": "mm-dd-yy",
    "CROREFERE": "@",
    "CROOBSERV": "@",
    "CROCLAOBJ": "General",
    "CRNCLAOBJ": "@",
    "GENUSUARIO4": "@",
    "CRNCONOBJ": "@",
    "SLNSERPRO": "@",
    "IDRIPS": "@",
    "CTNCENCOS": "@",
    "CROVALOBJ": '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-',
    "CRDOBSERV": "@",
    "CROTIPOBJ": "0",
}

CROCLAOBJ_CONST = 0
GENUSUARIO4_CONST = "999"

# Encabezados que reconoce el lector (normalizados: sin tildes, en mayúscula).
ALIAS_COLUMNAS: dict[str, set[str]] = {
    "factura": {"FACTURA", "NUMERO DE FACTURA", "NUMERO FACTURA", "NRO FACTURA"},
    "servicio": {"SERVICIO", "DESCRIPCION SERVICIO", "NOMBRE SERVICIO", "TECNOLOGIA"},
    "cantidad": {"CANTIDAD", "CANTIDAD FACTURADA", "CANT"},
    "glosa": {"DESCRIPCIONGLOSA", "DESCRIPCION GLOSA", "GLOSA", "CONCEPTO DE GLOSA"},
    "observacion": {"OBSERVACION", "OBSERVACIONES"},
    "valor": {"VALOR GLOSA", "VALOR GLOSADO", "VALOR OBJETADO"},
}

# Sin estas cuatro el archivo no sirve: se para antes de entregar algo malo.
COLUMNAS_OBLIGATORIAS = ("factura", "servicio", "glosa", "valor")

# "SO4201 - Existe ausencia…" → código + texto. También sin el guion y con el
# código partido por un espacio ("SO42 01"), como lo manda el Dispensario.
_RE_CODIGO_GLOSA = re.compile(r"^([A-Za-z]{2})\s*(\d{2})\s*(\d{2})\b\s*[-–:]?\s*(.*)$", re.DOTALL)

GRUPO_CLINICO = "CL"


def _texto(v: object) -> str:
    return " ".join(str(v if v is not None else "").split())


def codigo_y_concepto(descripcion: object) -> tuple[str, str]:
    """'SO4201 - Existe ausencia…' → ('SO4201', 'Existe ausencia…').

    CAPITAL SALUD no manda el código de glosa en columna propia: lo pega al
    principio de la descripción. Si la celda no empieza por un código, se
    devuelve código vacío y el texto entero: no se inventa un código.
    """
    t = _texto(descripcion)
    if not t:
        return "", ""
    m = _RE_CODIGO_GLOSA.match(t)
    if m:
        return (m.group(1) + m.group(2) + m.group(3)).upper(), _texto(m.group(4))
    return "", t


def nombre_servicio(valor: object) -> str:
    """'INTERNACION-ADULTOS-COMPLEJIDAD-ALTA' → 'INTERNACION ADULTOS COMPLEJIDAD ALTA'.

    CAPITAL escribe los nombres con guion donde va el espacio; el DGH los tiene
    con espacios. Se igualan al leer para que el cruce compare lo mismo.
    """
    return " ".join(str(valor if valor is not None else "").replace("-", " ").split())


def crotipobj_factura(grupos: set[str]) -> int:
    """Tipo de objeción de la factura según sus grupos de concepto (2 letras).

    Los tres valores que reconoce DGH: **0 = ADMINISTRATIVA** (sólo TA/FA/SO/
    AU/CO…), **1 = MEDICA** (sólo CL) y **2 = MIXTA** (CL junto con cualquier
    administrativa). Se decide por FACTURA, no por renglón.
    """
    tiene_cl = GRUPO_CLINICO in grupos
    tiene_admin = any(g and g != GRUPO_CLINICO for g in grupos)
    if tiene_cl and tiene_admin:
        return 2
    if tiene_cl:
        return 1
    return 0


def construir_crdobserv(codigo: str, servicio: str, observacion: str, valor: int) -> str:
    """CRDOBSERV = ``<código> <servicio>: <observación>$<valor>``."""
    partes = [p for p in (servicio, observacion) if p]
    texto = ": ".join(partes) if len(partes) == 2 else (partes[0] if partes else "")
    prefijo = f"{codigo} " if codigo else ""
    return f"{prefijo}{texto}${valor}"


# ─── Lectura del Excel de CAPITAL SALUD ──────────────────────────────────────


def mapear_columnas(encabezados: tuple) -> dict[str, int]:
    """{campo: posición} a partir de la fila de encabezados.

    Si falta alguna de las columnas indispensables se detiene diciendo cuál,
    en vez de entregar un archivo silenciosamente incompleto.
    """
    normalizados = [norm_header(h) for h in encabezados]
    indices: dict[str, int] = {}
    for campo, nombres in ALIAS_COLUMNAS.items():
        pos = next((i for i, h in enumerate(normalizados) if h in nombres), None)
        if pos is not None:
            indices[campo] = pos

    faltan = [c for c in COLUMNAS_OBLIGATORIAS if c not in indices]
    if faltan:
        raise ValueError(
            "El archivo de CAPITAL SALUD no trae las columnas "
            f"{', '.join(faltan)}. Encabezados leídos: {[h for h in normalizados if h]}. "
            "Se esperaba el export de CAPITAL SALUD (FACTURA, servicio, Cantidad, "
            "DescripcionGlosa, OBSERVACION, VALOR GLOSA)."
        )
    return indices


def leer_capital(ruta: Path, hoja: str | None = None) -> list[dict]:
    """Lee el Excel de CAPITAL SALUD y devuelve una lista de objeciones."""
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(ruta), data_only=True, read_only=True)
    try:
        if hoja:
            objetivo = _texto(hoja).upper()
            ws = next((x for x in wb.worksheets if _texto(x.title).upper() == objetivo), None)
            if ws is None:
                raise ValueError(f"La hoja '{hoja}' no existe en {ruta.name}.")
        else:
            ws = wb.active
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            raise ValueError(f"La hoja '{ws.title}' de {ruta.name} está vacía.")
        idx = mapear_columnas(filas[0])

        objeciones: list[dict] = []
        for n, fila in enumerate(filas[1:], start=2):
            if not any(fila):
                continue

            def dato(campo, f=fila):
                pos = idx.get(campo)
                return f[pos] if pos is not None and pos < len(f) else None

            factura = _texto(dato("factura"))
            if not factura:
                continue
            codigo, concepto = codigo_y_concepto(dato("glosa"))
            objeciones.append(
                {
                    "fila_excel": n,
                    "cxc": factura_larga(factura),
                    "codigo": codigo,
                    "concepto": concepto,
                    "servicio": nombre_servicio(dato("servicio")),
                    "cantidad": a_entero(dato("cantidad")),
                    "observacion": _texto(dato("observacion")),
                    "valor": a_entero(dato("valor")),
                }
            )
        return objeciones
    finally:
        wb.close()


# ─── Armado de los renglones ─────────────────────────────────────────────────


def construir_filas(
    objeciones: list[dict],
    fecha: _dt.datetime,
    servicios_dgh: dict | None = None,
    trazas: list[dict] | None = None,
) -> list[dict]:
    """Los renglones del formato OBJECIONES, uno por objeción."""
    consec: dict[str, int] = {}
    grupos: dict[str, set[str]] = defaultdict(set)
    for o in objeciones:
        grupos[o["cxc"]].add((o["codigo"] or "")[:2].upper())

    filas: list[dict] = []
    for o in objeciones:
        if o["cxc"] not in consec:
            consec[o["cxc"]] = len(consec) + 1

        slnserpro = ""
        if servicios_dgh is not None:
            unitario = o["valor"] // o["cantidad"] if o["cantidad"] else 0
            cruce = resolver_servicio(
                servicios_dgh.get(o["cxc"], []),
                # CAPITAL no manda código de servicio: sólo nombre y valor.
                codigo="",
                descripcion=o["servicio"],
                valor=o["valor"],
                valor_unitario=unitario,
                cantidad=o["cantidad"],
            )
            if cruce.linea is not None:
                slnserpro = cruce.linea.codigo
            if trazas is not None:
                trazas.append(
                    traza(
                        factura=o["cxc"],
                        codigo_objecion=o["codigo"],
                        valor=o["valor"],
                        cod_entidad="",
                        desc_entidad=o["servicio"],
                        unitario_entidad=unitario,
                        observacion=o["observacion"],
                        cruce=cruce,
                    )
                )

        filas.append(
            {
                "CDCONSEC": str(consec[o["cxc"]]),
                "CDFECDOC": fecha,
                "CRNCXC": o["cxc"],
                "CROFECOBJ": fecha,
                "CROREFERE": None,
                "CROOBSERV": None,
                "CROCLAOBJ": CROCLAOBJ_CONST,
                "CRNCLAOBJ": None,
                "GENUSUARIO4": GENUSUARIO4_CONST,
                "CRNCONOBJ": o["codigo"] or None,
                "SLNSERPRO": slnserpro or None,
                "IDRIPS": None,
                # CTNCENCOS va SIEMPRE vacía (regla del área).
                "CTNCENCOS": None,
                "CROVALOBJ": o["valor"],
                "CRDOBSERV": construir_crdobserv(
                    o["codigo"], o["servicio"], o["observacion"], o["valor"]
                ),
                "CROTIPOBJ": crotipobj_factura(grupos[o["cxc"]]),
            }
        )
    return filas


# ─── Escritura ───────────────────────────────────────────────────────────────


def escribir_objeciones(filas: list[dict], salida: Path) -> None:
    """Un .xlsx con hoja OBJECIONES y las 16 columnas del formato."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "OBJECIONES"
    fill_hdr = PatternFill("solid", fgColor="1F4E78")
    font_hdr = Font(bold=True, color="FFFFFF")
    for col, nombre in enumerate(COLUMNAS_OBJECIONES, start=1):
        c = ws.cell(row=1, column=col, value=nombre)
        c.fill = fill_hdr
        c.font = font_hdr

    for i, fila in enumerate(filas, start=2):
        for col, nombre in enumerate(COLUMNAS_OBJECIONES, start=1):
            celda = ws.cell(row=i, column=col, value=fila.get(nombre))
            celda.number_format = FORMATOS[nombre]

    ws.freeze_panes = "A2"
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(salida))


def escribir_por_factura(filas: list[dict], carpeta: Path, prefijo: str) -> list[Path]:
    """Un archivo por factura; cada uno reinicia su CDCONSEC en 1."""
    por_factura: dict[str, list[dict]] = defaultdict(list)
    for fila in filas:
        por_factura[fila["CRNCXC"]].append(fila)
    generados = []
    for cxc, grupo in sorted(por_factura.items()):
        destino = carpeta / f"{prefijo}_{cxc}.xlsx"
        escribir_objeciones([{**f, "CDCONSEC": "1"} for f in grupo], destino)
        generados.append(destino)
        logger.info("  %s: %d objeciones", destino.name, len(grupo))
    return generados


# ─── CLI ─────────────────────────────────────────────────────────────────────


def _parse_fecha(texto: str | None) -> _dt.datetime:
    if not texto:
        hoy = _dt.date.today()
        return _dt.datetime(hoy.year, hoy.month, hoy.day)
    try:
        return _dt.datetime.strptime(texto.strip(), "%Y-%m-%d")
    except ValueError:
        logger.error("--fecha inválida: %r (usá YYYY-MM-DD, p. ej. 2026-09-09)", texto)
        sys.exit(2)


PREFIJO_DEFAULT = "OBJECIONES_CAPITAL"


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--entrada", type=Path, required=True, help="Excel de glosas de CAPITAL SALUD")
    ap.add_argument("--hoja", default=None, help="Hoja del archivo (por defecto la primera)")
    ap.add_argument(
        "--salida",
        type=Path,
        required=True,
        help="Carpeta destino, o archivo .xlsx si se usa --consolidado",
    )
    ap.add_argument("--consolidado", action="store_true", help="Un solo Excel con todo")
    ap.add_argument(
        "--prefijo", default=PREFIJO_DEFAULT, help="Prefijo de los archivos por factura"
    )
    ap.add_argument(
        "--servicios-dgh",
        type=Path,
        default=None,
        help="Export de servicios facturados del DGH: CAPITAL no manda el código del "
        "servicio, así que sin esto SLNSERPRO sale vacío en todo el archivo",
    )
    ap.add_argument(
        "--reporte-cruce",
        type=Path,
        default=None,
        help="Excel de trabajo con el detalle del cruce (hojas CRUCE, REVISAR y RESUMEN). "
        "Requiere --servicios-dgh",
    )
    ap.add_argument("--fecha", default=None, help="Fecha de la objeción (YYYY-MM-DD)")
    ap.add_argument("--log", type=Path, default=None, help="Guarda un log adicional a archivo")
    args = ap.parse_args(argv)

    handlers: list[logging.Handler] = [logging.StreamHandler(sys.stdout)]
    if args.log is not None:
        args.log.parent.mkdir(parents=True, exist_ok=True)
        handlers.append(logging.FileHandler(args.log, encoding="utf-8"))
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s", handlers=handlers
    )

    if not args.entrada.is_file():
        logger.error("No existe el archivo de entrada: %s", args.entrada)
        return 1
    if args.reporte_cruce is not None and args.servicios_dgh is None:
        logger.error("--reporte-cruce necesita --servicios-dgh (es el detalle de ese cruce).")
        return 1

    servicios_dgh = None
    if args.servicios_dgh is not None:
        if not args.servicios_dgh.is_file():
            logger.error("No existe el export del DGH: %s", args.servicios_dgh)
            return 1
        logger.info("Leyendo servicios facturados del DGH: %s", args.servicios_dgh.name)
        try:
            servicios_dgh = leer_servicios_dgh(args.servicios_dgh, avisar=logger.warning)
        except ValueError as e:
            logger.error(str(e))
            return 1
        logger.info(
            "  %d renglones de servicio en %d factura(s).",
            sum(len(v) for v in servicios_dgh.values()),
            len(servicios_dgh),
        )
    else:
        logger.warning(
            "Sin --servicios-dgh no hay con qué llenar SLNSERPRO: CAPITAL SALUD no "
            "manda el código del servicio. El archivo saldrá con esa columna vacía."
        )

    logger.info("Leyendo glosas de CAPITAL SALUD: %s", args.entrada.name)
    try:
        objeciones = leer_capital(args.entrada, args.hoja)
    except ValueError as e:
        logger.error(str(e))
        return 1
    if not objeciones:
        logger.error("No se encontró ninguna objeción en %s", args.entrada.name)
        return 1

    sin_codigo = [o for o in objeciones if not o["codigo"]]
    if sin_codigo:
        logger.warning(
            "  ⚠ %d objeciones sin código de glosa legible en DescripcionGlosa: "
            "CRNCONOBJ queda vacío (filas %s).",
            len(sin_codigo),
            ", ".join(str(o["fila_excel"]) for o in sin_codigo[:10]),
        )
    sin_servicio = [o for o in objeciones if not o["servicio"]]
    if sin_servicio:
        logger.warning(
            "  ⚠ %d objeciones sin nombre de servicio: no hay con qué cruzar (filas %s).",
            len(sin_servicio),
            ", ".join(str(o["fila_excel"]) for o in sin_servicio[:10]),
        )

    trazas: list[dict] = []
    filas = construir_filas(objeciones, _parse_fecha(args.fecha), servicios_dgh, trazas)

    if servicios_dgh is not None:
        cuenta = {
            k: sum(1 for t in trazas if t["confianza"] == k)
            for k in ("ALTA", "MEDIA", "BAJA", "SIN CRUCE")
        }
        ubicados = cuenta["ALTA"] + cuenta["MEDIA"]
        logger.info(
            "  Cruce contra los servicios del DGH: %d de %d servicios ubicados con "
            "confianza alta/media (%.0f%%).",
            ubicados,
            len(trazas),
            ubicados / (len(trazas) or 1) * 100,
        )
        logger.info(
            "    ALTA=%d  MEDIA=%d  BAJA=%d  SIN CRUCE=%d  → revisar %d.",
            cuenta["ALTA"],
            cuenta["MEDIA"],
            cuenta["BAJA"],
            cuenta["SIN CRUCE"],
            cuenta["BAJA"] + cuenta["SIN CRUCE"],
        )

    verificar_reglas(
        [
            {
                "factura": f["CRNCXC"],
                "slnserpro": f["SLNSERPRO"],
                "ctncencos": f["CTNCENCOS"],
                "crotipobj": f["CROTIPOBJ"],
                "codigo_glosa": f["CRNCONOBJ"],
            }
            for f in filas
        ],
        servicios_dgh,
        avisar=logger.info,
    )

    facturas = sorted({o["cxc"] for o in objeciones})
    if args.consolidado:
        escribir_objeciones(filas, args.salida)
        logger.info(
            "\nConsolidado: %d objeciones de %d factura(s) → %s",
            len(filas),
            len(facturas),
            args.salida,
        )
    else:
        generados = escribir_por_factura(filas, args.salida, args.prefijo)
        logger.info("\n%d archivo(s) de CAPITAL SALUD en: %s", len(generados), args.salida)

    if args.reporte_cruce is not None:
        escribir_reporte_cruce(trazas, args.reporte_cruce, entidad="CAPITAL SALUD")
        pendientes = sum(1 for t in trazas if t["aviso"] or t["confianza"] in ("BAJA", "SIN CRUCE"))
        logger.info(
            "Detalle del cruce: %s (%d renglón(es) en la hoja REVISAR).",
            args.reporte_cruce,
            pendientes,
        )

    logger.info(
        "  Facturas: %d  |  Objeciones: %d  |  Valor glosado total: $%s",
        len(facturas),
        len(filas),
        f"{sum(o['valor'] or 0 for o in objeciones):,}",
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
