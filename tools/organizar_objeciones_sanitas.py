"""organizar_objeciones_sanitas.py — Objeciones de SANITAS → Excel OBJECIONES (DGH).

Toma el Excel de glosas que entrega **SANITAS** (hoja `Glosa`, 7 columnas) y lo
convierte al formato de trabajo de 16 columnas (hoja `OBJECIONES`) que se
importa en Dinámica Gerencial, el mismo de FAMISANAR, SAVIA y el Dispensario.

OJO CON EL ARCHIVO DE SANITAS. Sus encabezados están mal rotulados: la SEGUNDA
columna se llama otra vez «NUMERO DE FACTURA» pero lo que trae es el **código
de glosa** (CO2301, TA0801…). Por eso las columnas NO se pueden resolver por
nombre —un lector por encabezado metería el código del procedimiento en la
columna del código de glosa y el archivo saldría mal sin que nadie lo note—:
se leen por POSICIÓN y se verifica el contenido antes de seguir
(`verificar_columnas`). Si el archivo llega con otro orden, el bot se detiene
con un mensaje claro en vez de entregar un archivo silenciosamente malo.

    col 0  NUMERO DE FACTURA        → HUS0000548650
    col 1  «NUMERO DE FACTURA»      → CÓDIGO DE GLOSA (CO2301, TA0801, …)
    col 2  VALOR REAL GLOSA         → valor objetado
    col 3  CODIGO PROCEDIMIENTO     → código del servicio
    col 4  NOMBRE PROCEDIMIENTO     → nombre del servicio
    col 5  CANTIDAD PROCEDIMIENTO   → cantidad
    col 6  OBSERVACIONES            → motivo escrito por el auditor de SANITAS

A diferencia de FAMISANAR, SANITAS **sí** manda el código del servicio en su
propia columna. Aun así se cruza contra el export del DGH (`--servicios-dgh`)
para confirmar que ese código existe en esa factura: la regla del área es que
en `SLNSERPRO` no puede ir un código que el DGH no reconozca.

REGLAS FIJAS DEL FORMATO (ver CLAUDE.md, valen para todas las entidades):
    CTNCENCOS  → siempre vacía
    CROTIPOBJ  → por factura: 0 = ADMINISTRATIVA, 1 = MEDICA, 2 = MIXTA
    SLNSERPRO  → prohibido inventar: sin cruce confiable, celda vacía
    El archivo lleva el 100% de los renglones; los no cruzados van con la
    celda vacía para completarlos a mano.

USO:
    py tools\\organizar_objeciones_sanitas.py ^
        --entrada       "D:\\...\\SANITAS_4_SEPTIEMBRE.xlsx" ^
        --servicios-dgh "D:\\...\\SERVICIOS_FACTURADOS_DGH.xlsx" ^
        --salida        "D:\\...\\OBJECIONES_SANITAS_04092026.xlsx" --consolidado ^
        --reporte-cruce "D:\\...\\CRUCE_SANITAS_04092026.xlsx" ^
        --fecha 2026-09-04

INSTALACIÓN (una vez):  py -m pip install openpyxl
"""

from __future__ import annotations

import argparse
import datetime as _dt
import logging
import re
import sys
from collections import defaultdict
from pathlib import Path

# Motor del cruce y lector de pesos, comunes a todos los bots.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from _cruce_dgh import (  # noqa: E402
    escribir_reporte_cruce,
    factura_larga,
    leer_servicios_dgh,
    resolver_servicio,
    traza,
    verificar_reglas,
)
from _dinero import a_entero  # noqa: E402

logger = logging.getLogger("organizar_sanitas")


# ─── Formato de salida: layout de 16 columnas (hoja OBJECIONES) ──────────────

COLUMNAS_OBJECIONES: tuple[str, ...] = (
    "CDCONSEC",
    "CDFECDOC",
    "CRNCXC",
    "CROFECOBJ",
    "CROREFERE",
    "CROOBSERV",
    "CROCLAOBJ",
    "CRNCLAOBJ",
    "GENUSUARIO4",
    "CRNCONOBJ",
    "SLNSERPRO",
    "IDRIPS",
    "CTNCENCOS",
    "CROVALOBJ",
    "CRDOBSERV",
    "CROTIPOBJ",
)

# number_format por columna, igual que en los archivos reales.
FORMATOS: dict[str, str] = {
    "CDCONSEC": "@",
    "CDFECDOC": "mm-dd-yy",
    "CRNCXC": "@",
    "CROFECOBJ": "mm-dd-yy",
    "CROREFERE": "@",
    "CROOBSERV": "@",
    "CROCLAOBJ": "General",
    "CRNCLAOBJ": "@",
    "GENUSUARIO4": "@",
    "CRNCONOBJ": "@",
    "SLNSERPRO": "@",
    "IDRIPS": "@",
    "CTNCENCOS": "@",
    "CROVALOBJ": '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-',
    "CRDOBSERV": "@",
    "CROTIPOBJ": "0",
}

CROCLAOBJ_CONST = 0
GENUSUARIO4_CONST = "999"

# Posición fija de cada dato en el Excel de SANITAS (los encabezados no sirven).
IDX_FACTURA = 0
IDX_CODIGO_GLOSA = 1
IDX_VALOR = 2
IDX_COD_SERVICIO = 3
IDX_NOMBRE_SERVICIO = 4
IDX_CANTIDAD = 5
IDX_OBSERVACION = 6

_RE_CODIGO_GLOSA = re.compile(r"^([A-Z]{2}\d{2})\s*(\d{2})$")
_RE_CODIGO_GLOSA_LIMPIO = re.compile(r"^[A-Z]{2}\d{4}$")
# La factura lleva prefijo y MUCHOS dígitos (HUS0000548650). Se exigen cinco o
# más para que un código de glosa —dos letras y cuatro dígitos, "CO2301"— no
# pase por factura y la verificación de columnas sirva de algo.
_RE_FACTURA = re.compile(r"^[A-Z]{2,5}\d{5,}$")

GRUPO_CLINICO = "CL"


def _texto(v: object) -> str:
    return " ".join(str(v if v is not None else "").split())


def codigo_glosa(valor: object) -> str:
    """Deja el código de glosa limpio: 'TA08 01' → 'TA0801'; 'CO2301' tal cual."""
    t = _texto(valor).upper()
    m = _RE_CODIGO_GLOSA.match(t)
    if m:
        return m.group(1) + m.group(2)
    return t.replace(" ", "")


def crotipobj_factura(grupos: set[str]) -> int:
    """Tipo de objeción de la factura según sus grupos de concepto (2 letras).

    Los tres valores que reconoce DGH: **0 = ADMINISTRATIVA** (sólo TA/FA/SO/
    AU/CO…), **1 = MEDICA** (sólo CL) y **2 = MIXTA** (CL junto con cualquier
    administrativa). Se decide por FACTURA, no por renglón.
    """
    tiene_cl = GRUPO_CLINICO in grupos
    tiene_admin = any(g and g != GRUPO_CLINICO for g in grupos)
    if tiene_cl and tiene_admin:
        return 2
    if tiene_cl:
        return 1
    return 0


def construir_crdobserv(codigo: str, servicio: str, observacion: str, valor: int) -> str:
    """CRDOBSERV = ``<código> <servicio>: <observación>$<valor>``."""
    partes = [p for p in (servicio, observacion) if p]
    texto = ": ".join(partes) if len(partes) == 2 else (partes[0] if partes else "")
    prefijo = f"{codigo} " if codigo else ""
    return f"{prefijo}{texto}${valor}"


# ─── Lectura del Excel de SANITAS ────────────────────────────────────────────


def verificar_columnas(filas: list[tuple]) -> None:
    """Comprueba que las columnas estén donde se esperan antes de procesar.

    El archivo de SANITAS trae dos columnas con el mismo encabezado y una de
    ellas mal rotulada, así que la única defensa es mirar el CONTENIDO: la
    primera debe ser la factura y la segunda un código de glosa. Si no cuadra,
    mejor detenerse que entregar un archivo malo.
    """
    muestra = [f for f in filas if any(f)][:25]
    if not muestra:
        raise ValueError("El archivo de SANITAS no tiene renglones de glosa.")

    facturas = [_texto(f[IDX_FACTURA]).upper() for f in muestra if len(f) > IDX_FACTURA]
    if not facturas or sum(bool(_RE_FACTURA.match(x)) for x in facturas) < len(facturas) * 0.8:
        raise ValueError(
            "La primera columna no parece el número de factura "
            f"(ejemplos: {facturas[:3]}). Revisá el archivo de SANITAS."
        )

    codigos = [codigo_glosa(f[IDX_CODIGO_GLOSA]) for f in muestra if len(f) > IDX_CODIGO_GLOSA]
    if (
        not codigos
        or sum(bool(_RE_CODIGO_GLOSA_LIMPIO.match(x)) for x in codigos) < len(codigos) * 0.8
    ):
        raise ValueError(
            "La segunda columna no trae códigos de glosa tipo CO2301 "
            f"(ejemplos: {codigos[:3]}). En el archivo de SANITAS esa columna viene "
            "rotulada «NUMERO DE FACTURA» pero debe contener el código de la glosa; "
            "si cambió el orden de las columnas, hay que ajustar el bot."
        )


def leer_sanitas(ruta: Path, hoja: str | None = None) -> list[dict]:
    """Lee el Excel de SANITAS y devuelve una lista de objeciones."""
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(ruta), data_only=True, read_only=True)
    try:
        if hoja:
            objetivo = _texto(hoja).upper()
            ws = next((x for x in wb.worksheets if _texto(x.title).upper() == objetivo), None)
            if ws is None:
                raise ValueError(f"La hoja '{hoja}' no existe en {ruta.name}.")
        else:
            ws = next((x for x in wb.worksheets if _texto(x.title).upper() == "GLOSA"), wb.active)
        filas = list(ws.iter_rows(values_only=True))[1:]  # saltar encabezados
        verificar_columnas(filas)

        objeciones: list[dict] = []
        for n, fila in enumerate(filas, start=2):
            if not any(fila):
                continue
            factura = _texto(fila[IDX_FACTURA])
            if not factura:
                continue

            def dato(i, f=fila):
                return f[i] if i < len(f) else None

            objeciones.append(
                {
                    "fila_excel": n,
                    "cxc": factura_larga(factura),
                    "codigo": codigo_glosa(dato(IDX_CODIGO_GLOSA)),
                    "cod_servicio": _texto(dato(IDX_COD_SERVICIO)),
                    "servicio": _texto(dato(IDX_NOMBRE_SERVICIO)),
                    "cantidad": a_entero(dato(IDX_CANTIDAD)),
                    "observacion": _texto(dato(IDX_OBSERVACION)),
                    "valor": a_entero(dato(IDX_VALOR)),
                }
            )
        return objeciones
    finally:
        wb.close()


# ─── Armado de los renglones ─────────────────────────────────────────────────


def construir_filas(
    objeciones: list[dict],
    fecha: _dt.datetime,
    servicios_dgh: dict | None = None,
    trazas: list[dict] | None = None,
) -> list[dict]:
    """Los renglones del formato OBJECIONES, uno por objeción."""
    consec: dict[str, int] = {}
    grupos: dict[str, set[str]] = defaultdict(set)
    for o in objeciones:
        grupos[o["cxc"]].add((o["codigo"] or "")[:2].upper())

    filas: list[dict] = []
    for o in objeciones:
        if o["cxc"] not in consec:
            consec[o["cxc"]] = len(consec) + 1

        slnserpro = ""
        if servicios_dgh is not None:
            unitario = o["valor"] // o["cantidad"] if o["cantidad"] else 0
            cruce = resolver_servicio(
                servicios_dgh.get(o["cxc"], []),
                codigo=o["cod_servicio"],
                descripcion=o["servicio"],
                valor=o["valor"],
                valor_unitario=unitario,
                cantidad=o["cantidad"],
            )
            if cruce.linea is not None:
                slnserpro = cruce.linea.codigo
            if trazas is not None:
                trazas.append(
                    traza(
                        factura=o["cxc"],
                        codigo_objecion=o["codigo"],
                        valor=o["valor"],
                        cod_entidad=o["cod_servicio"],
                        desc_entidad=o["servicio"],
                        unitario_entidad=unitario,
                        observacion=o["observacion"],
                        cruce=cruce,
                    )
                )

        filas.append(
            {
                "CDCONSEC": str(consec[o["cxc"]]),
                "CDFECDOC": fecha,
                "CRNCXC": o["cxc"],
                "CROFECOBJ": fecha,
                "CROREFERE": None,
                "CROOBSERV": None,
                "CROCLAOBJ": CROCLAOBJ_CONST,
                "CRNCLAOBJ": None,
                "GENUSUARIO4": GENUSUARIO4_CONST,
                "CRNCONOBJ": o["codigo"] or None,
                "SLNSERPRO": slnserpro or None,
                "IDRIPS": None,
                # CTNCENCOS va SIEMPRE vacía (regla del área).
                "CTNCENCOS": None,
                "CROVALOBJ": o["valor"],
                "CRDOBSERV": construir_crdobserv(
                    o["codigo"], o["servicio"], o["observacion"], o["valor"]
                ),
                "CROTIPOBJ": crotipobj_factura(grupos[o["cxc"]]),
            }
        )
    return filas


# ─── Escritura ───────────────────────────────────────────────────────────────


def escribir_objeciones(filas: list[dict], salida: Path) -> None:
    """Un .xlsx con hoja OBJECIONES y las 16 columnas del formato."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "OBJECIONES"
    fill_hdr = PatternFill("solid", fgColor="1F4E78")
    font_hdr = Font(bold=True, color="FFFFFF")
    for col, nombre in enumerate(COLUMNAS_OBJECIONES, start=1):
        c = ws.cell(row=1, column=col, value=nombre)
        c.fill = fill_hdr
        c.font = font_hdr

    for i, fila in enumerate(filas, start=2):
        for col, nombre in enumerate(COLUMNAS_OBJECIONES, start=1):
            celda = ws.cell(row=i, column=col, value=fila.get(nombre))
            celda.number_format = FORMATOS[nombre]

    ws.freeze_panes = "A2"
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(salida))


def escribir_por_factura(filas: list[dict], carpeta: Path, prefijo: str) -> list[Path]:
    """Un archivo por factura; cada uno reinicia su CDCONSEC en 1."""
    por_factura: dict[str, list[dict]] = defaultdict(list)
    for fila in filas:
        por_factura[fila["CRNCXC"]].append(fila)
    generados = []
    for cxc, grupo in sorted(por_factura.items()):
        destino = carpeta / f"{prefijo}_{cxc}.xlsx"
        escribir_objeciones([{**f, "CDCONSEC": "1"} for f in grupo], destino)
        generados.append(destino)
        logger.info("  %s: %d objeciones", destino.name, len(grupo))
    return generados


# ─── CLI ─────────────────────────────────────────────────────────────────────


def _parse_fecha(texto: str | None) -> _dt.datetime:
    if not texto:
        hoy = _dt.date.today()
        return _dt.datetime(hoy.year, hoy.month, hoy.day)
    try:
        return _dt.datetime.strptime(texto.strip(), "%Y-%m-%d")
    except ValueError:
        logger.error("--fecha inválida: %r (usá YYYY-MM-DD, p. ej. 2026-09-04)", texto)
        sys.exit(2)


PREFIJO_DEFAULT = "OBJECIONES_SANITAS"


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--entrada", type=Path, required=True, help="Excel de glosas de SANITAS")
    ap.add_argument("--hoja", default=None, help="Hoja del archivo (por defecto 'Glosa')")
    ap.add_argument(
        "--salida",
        type=Path,
        required=True,
        help="Carpeta destino, o archivo .xlsx si se usa --consolidado",
    )
    ap.add_argument("--consolidado", action="store_true", help="Un solo Excel con todo")
    ap.add_argument(
        "--prefijo", default=PREFIJO_DEFAULT, help="Prefijo de los archivos por factura"
    )
    ap.add_argument(
        "--servicios-dgh",
        type=Path,
        default=None,
        help="Export de servicios facturados del DGH: con él SLNSERPRO queda con el "
        "código que Dinámica Gerencial reconoce",
    )
    ap.add_argument(
        "--reporte-cruce",
        type=Path,
        default=None,
        help="Excel de trabajo con el detalle del cruce (hojas CRUCE, REVISAR y RESUMEN). "
        "Requiere --servicios-dgh",
    )
    ap.add_argument("--fecha", default=None, help="Fecha de la objeción (YYYY-MM-DD)")
    ap.add_argument("--log", type=Path, default=None, help="Guarda un log adicional a archivo")
    args = ap.parse_args(argv)

    handlers: list[logging.Handler] = [logging.StreamHandler(sys.stdout)]
    if args.log is not None:
        args.log.parent.mkdir(parents=True, exist_ok=True)
        handlers.append(logging.FileHandler(args.log, encoding="utf-8"))
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s", handlers=handlers
    )

    if not args.entrada.is_file():
        logger.error("No existe el archivo de entrada: %s", args.entrada)
        return 1
    if args.reporte_cruce is not None and args.servicios_dgh is None:
        logger.error("--reporte-cruce necesita --servicios-dgh (es el detalle de ese cruce).")
        return 1

    servicios_dgh = None
    if args.servicios_dgh is not None:
        if not args.servicios_dgh.is_file():
            logger.error("No existe el export del DGH: %s", args.servicios_dgh)
            return 1
        logger.info("Leyendo servicios facturados del DGH: %s", args.servicios_dgh.name)
        try:
            servicios_dgh = leer_servicios_dgh(args.servicios_dgh, avisar=logger.warning)
        except ValueError as e:
            logger.error(str(e))
            return 1
        logger.info(
            "  %d renglones de servicio en %d factura(s).",
            sum(len(v) for v in servicios_dgh.values()),
            len(servicios_dgh),
        )

    logger.info("Leyendo glosas de SANITAS: %s", args.entrada.name)
    try:
        objeciones = leer_sanitas(args.entrada, args.hoja)
    except ValueError as e:
        logger.error(str(e))
        return 1
    if not objeciones:
        logger.error("No se encontró ninguna objeción en %s", args.entrada.name)
        return 1

    sin_codigo = [o for o in objeciones if not o["codigo"]]
    if sin_codigo:
        logger.warning(
            "  ⚠ %d objeciones sin código de glosa: CRNCONOBJ queda vacío (filas %s).",
            len(sin_codigo),
            ", ".join(str(o["fila_excel"]) for o in sin_codigo[:10]),
        )

    trazas: list[dict] = []
    filas = construir_filas(objeciones, _parse_fecha(args.fecha), servicios_dgh, trazas)

    if servicios_dgh is not None:
        cuenta = {
            k: sum(1 for t in trazas if t["confianza"] == k)
            for k in ("ALTA", "MEDIA", "BAJA", "SIN CRUCE")
        }
        ubicados = cuenta["ALTA"] + cuenta["MEDIA"]
        logger.info(
            "  Cruce contra los servicios del DGH: %d de %d servicios ubicados con "
            "confianza alta/media (%.0f%%).",
            ubicados,
            len(trazas),
            ubicados / (len(trazas) or 1) * 100,
        )
        logger.info(
            "    ALTA=%d  MEDIA=%d  BAJA=%d  SIN CRUCE=%d  → revisar %d.",
            cuenta["ALTA"],
            cuenta["MEDIA"],
            cuenta["BAJA"],
            cuenta["SIN CRUCE"],
            cuenta["BAJA"] + cuenta["SIN CRUCE"],
        )

    verificar_reglas(
        [
            {
                "factura": f["CRNCXC"],
                "slnserpro": f["SLNSERPRO"],
                "ctncencos": f["CTNCENCOS"],
                "crotipobj": f["CROTIPOBJ"],
                "codigo_glosa": f["CRNCONOBJ"],
            }
            for f in filas
        ],
        servicios_dgh,
        avisar=logger.info,
    )

    facturas = sorted({o["cxc"] for o in objeciones})
    if args.consolidado:
        escribir_objeciones(filas, args.salida)
        logger.info(
            "\nConsolidado: %d objeciones de %d factura(s) → %s",
            len(filas),
            len(facturas),
            args.salida,
        )
    else:
        generados = escribir_por_factura(filas, args.salida, args.prefijo)
        logger.info("\n%d archivo(s) de SANITAS en: %s", len(generados), args.salida)

    if args.reporte_cruce is not None:
        escribir_reporte_cruce(trazas, args.reporte_cruce, entidad="SANITAS")
        pendientes = sum(1 for t in trazas if t["aviso"] or t["confianza"] in ("BAJA", "SIN CRUCE"))
        logger.info(
            "Detalle del cruce: %s (%d renglón(es) en la hoja REVISAR).",
            args.reporte_cruce,
            pendientes,
        )

    logger.info(
        "  Facturas: %d  |  Objeciones: %d  |  Valor glosado total: $%s",
        len(facturas),
        len(filas),
        f"{sum(o['valor'] for o in objeciones):,}",
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
