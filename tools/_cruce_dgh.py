"""_cruce_dgh.py — Motor común para ubicar el servicio de una objeción en el DGH.

EL PROBLEMA. Las entidades (FAMISANAR, el Dispensario Médico, …) mandan las
objeciones diciendo el servicio a su manera: con el catálogo IUM
(``91022534``, «LINEA INFUSION E INYECCION - JERINGA 1ML 25G»), con el código
del medicamento rellenado con ceros (``P32606-02`` donde el HUS tiene
``32606-2``) o simplemente con el nombre («MONITOREO ELECTROCARDIOGRAFICO
CONTINUO (HOLTER)»). Dinámica Gerencial sólo reconoce **su** código, así que
sin traducirlo el renglón no se puede cargar.

QUÉ HACE ESTE MÓDULO. Con el export de servicios facturados del DGH busca,
**dentro de la misma factura**, de qué servicio habla cada objeción, puntuando
cada renglón por código, nombre y valor. Devuelve el renglón encontrado con su
NIVEL DE CONFIANZA (ALTA / MEDIA / BAJA) o vacío. **Nunca inventa un
servicio:** sin candidato con puntaje suficiente devuelve un cruce vacío y el
bot que lo llama deja lo que ya sabía y manda el renglón a revisión.

Lo usan `organizar_objeciones_famisanar.py` y
`organizar_objeciones_dispensario.py`. Vive aparte —como `_dinero.py`— para
que la lógica del cruce no se duplique y se separe entre bots.
"""

from __future__ import annotations

import difflib
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from _dinero import a_entero

# ─── Columnas del export de servicios facturados del DGH ─────────────────────

COLUMNAS_DGH_SERVICIOS = {
    # "SERVICOS DGH" (sin la I) es como sale el encabezado en algunos exports:
    # sin ese alias la columna del código se perdía y el archivo salía con
    # SLNSERPRO vacío aunque el cruce sí hubiera encontrado el renglón.
    "codigo": {
        "SERVICIOS DGH",
        "SERVICIOS_DGH",
        "SERVICOS DGH",
        "SERVICOS_DGH",
        "SLNSERPRO SERVICIO",
        "SLNSERPRO_SERVICIO",
    },
    "desc_institucional": {"DESCRIPCION INSTITUCIONAL", "DESCRIPCION_INSTITUCIONAL"},
    "cups": {"SLNSERPRO CUPS", "SLNSERPRO_CUPS"},
    "desc_cups": {"DESCRIPCION CUPS", "DESCRIPCION_CUPS"},
    "cod_medicamento": {"CODIGO MEDICAMENTO", "CODIGO_MEDICAMENTO"},
    "nombre_medicamento": {"NOMBRE MEDICAMENTO", "NOMBRE_MEDICAMENTO"},
    "centro_costo": {"NOM CENTRO COSTO", "NOM_CENTRO_COSTO", "CENTRO COSTO", "CENTRO_COSTO"},
    "factura": {"FACTURA", "NRO_FACTURA", "NRO FACTURA", "NUMERO FACTURA"},
    "cantidad": {"CAT SERVICIOS", "CAT_SERVICIOS", "CANTIDAD", "CANT"},
    "valor": {"VR SERVICIO", "VR_SERVICIO", "VALOR SERVICIO"},
}


# ─── Normalización de texto y códigos ────────────────────────────────────────

# Palabras que no distinguen un servicio de otro.
VACIAS_DESC = frozenset(
    {
        "DE", "DEL", "LA", "EL", "LOS", "LAS", "POR", "CON", "SIN", "PARA", "EN",
        "A", "Y", "O", "X", "AL", "UN", "UNA", "SU", "MAS", "REF", "TIPO",
    }
)  # fmt: skip

_RE_NO_ALFA = re.compile(r"[^A-Z0-9 ]+")
_RE_NUM_LETRA = re.compile(r"(?<=\d)(?=[A-Z])")
_RE_LETRA_NUM = re.compile(r"(?<=[A-Z])(?=\d)")


def norm_header(h: object) -> str:
    """Normaliza un encabezado: mayúsculas, sin tildes, sin espacios de más."""
    s = unicodedata.normalize("NFKD", str(h or "").strip().upper())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.split())


def norm_desc(texto: object) -> str:
    """MAYÚSCULAS, sin tildes, sin puntuación, espacios colapsados."""
    if texto is None:
        return ""
    return " ".join(_RE_NO_ALFA.sub(" ", norm_header(texto)).split())


def palabras_desc(desc: str) -> set[str]:
    """Palabras significativas, separando número y unidad pegados (1ML → 1 ML),
    para que 'JERINGA DESECHABLE 1ML' del DGH y 'JERINGA 1 ML CON AGUJA' de la
    entidad compartan las mismas palabras."""
    if not desc:
        return set()
    t = _RE_LETRA_NUM.sub(" ", _RE_NUM_LETRA.sub(" ", desc))
    return {w for w in t.split() if w and w not in VACIAS_DESC}


def parecido_desc(a: str, b: str) -> float:
    """0..1 — parecido entre dos nombres de servicio ya normalizados."""
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    ratio = difflib.SequenceMatcher(None, a, b).ratio()
    if a in b or b in a:
        ratio = max(ratio, 0.90)
    ta, tb = palabras_desc(a), palabras_desc(b)
    if ta and tb:
        comunes = ta & tb
        # Sin ninguna palabra de fondo en común, el parecido letra a letra es
        # ruido ('VITAMINA D3 CAP X 1 000' vs. 'JERINGA 1ML 25G'): se topa.
        if not {w for w in comunes if len(w) >= 3 and not w.isdigit()}:
            ratio = min(ratio, 0.30)
        ratio = max(ratio, len(comunes) / len(ta | tb) * 0.95)
        # Cobertura: qué tanto del nombre más corto está dentro del más largo.
        # La entidad antepone la categoría ('LINEA INFUSION E INYECCION - …') y
        # el DGH usa el nombre pelado.
        corto = min(len(ta), len(tb))
        cobertura = len(comunes) / corto if corto else 0.0
        if corto < 2 or len(comunes) < 2:
            cobertura = min(cobertura, 0.55)
        ratio = max(ratio, cobertura * 0.85)
    return ratio


def variantes_codigo(cod: object) -> set[str]:
    """Todas las formas en que la misma referencia puede venir escrita:
    la letra que antepone la entidad (P32606-02), la H que agrega el DGH al
    CUPS (903437H) y el sufijo rellenado con ceros (32606-02 = 32606-2)."""
    crudo = norm_header(cod)
    base = re.sub(r"[^A-Z0-9]", "", crudo)
    if not base:
        return set()
    v = {base}
    if len(base) > 1 and base[0] in ("P", "U") and base[1].isdigit():
        v.add(base[1:])
    if base.endswith("H") and base[:-1].isdigit():
        v.add(base[:-1])
    m = re.match(r"^([PU]?)([A-Z0-9]+)-0*(\d+)$", crudo)
    if m:
        v.add(re.sub(r"[^A-Z0-9]", "", f"{m.group(2)}{m.group(3)}"))
        v.add(re.sub(r"[^A-Z0-9]", "", f"{m.group(1)}{m.group(2)}{m.group(3)}"))
    return {x for x in v if x}


# Código largo pegado adelante del nombre ("1O1044511000101 OXIGENO").
_RE_COD_LARGO_PEGADO = re.compile(r"^[0-9A-Z]{7,}\s+(?=[A-Z])")


def formas_descripcion(desc: str) -> list[str]:
    """Formas alternativas de leer el nombre que escribió la entidad: completo,
    sin el código pegado adelante y sólo lo que va después del guion
    ('LINEA INFUSION E INYECCION - JERINGA 5 ML' → 'JERINGA 5 ML')."""
    base = norm_desc(desc)
    if not base:
        return []
    formas = [base]
    sin_cod = _RE_COD_LARGO_PEGADO.sub("", base)
    if sin_cod and sin_cod != base:
        formas.append(sin_cod)
    if " - " in (desc or ""):
        cola = norm_desc(desc.split(" - ")[-1])
        if cola and cola not in formas:
            formas.append(cola)
    return formas


_RE_FACTURA = re.compile(r"^([A-Za-z]+)0*(\d+)$")


def factura_larga(fac: object, ancho: int = 10) -> str:
    """HUS532670 → HUS0000532670 (rellena con ceros a `ancho` dígitos).
    Idempotente. Si no matchea el patrón, devuelve el texto tal cual."""
    s = str(fac or "").strip()
    m = _RE_FACTURA.match(s)
    if not m:
        return s
    return m.group(1).upper() + m.group(2).zfill(ancho)


# ─── Lectura del export del DGH ──────────────────────────────────────────────


class LineaDgh:
    """Un renglón del export de servicios facturados del DGH."""

    __slots__ = (
        "codigo",
        "descripcion",
        "descripciones",
        "codigos",
        "centro_costo",
        "cantidad",
        "valor",
        "unitario",
        "usos",
    )

    def __init__(
        self, codigo, descripcion, desc_cups, nombre_med, cups, cod_med, centro, cant, valor
    ):
        # Si el export no trae la columna de servicio, el CUPS o el código del
        # medicamento identifican la misma línea en DGH (igual que en el bot
        # del ADRES). Antes el renglón salía sin código aunque el cruce
        # acertara.
        self.codigo = codigo or cups or cod_med
        self.descripcion = descripcion or desc_cups or nombre_med
        self.descripciones = {norm_desc(d) for d in (descripcion, desc_cups, nombre_med) if d}
        self.codigos: set[str] = set()
        for c in (codigo, cups, cod_med):
            self.codigos |= variantes_codigo(c)
        self.centro_costo = centro
        self.cantidad = cant
        self.valor = valor
        self.unitario = (valor / cant) if cant else valor
        self.usos = 0


def leer_servicios_dgh(ruta: Path, avisar=None) -> dict[str, list[LineaDgh]]:
    """{factura: [servicios facturados]} desde el export del DGH.

    `avisar` es una función tipo `logger.warning` para los avisos no fatales.
    """
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(ruta), data_only=True, read_only=True)
    try:
        ws = wb.active
        filas = ws.iter_rows(values_only=True)
        headers = list(next(filas, ()) or ())
        norm = [norm_header(h) for h in headers]
        idx: dict[str, int] = {}
        for clave, alias in COLUMNAS_DGH_SERVICIOS.items():
            i = next((n for n, h in enumerate(norm) if h in alias), None)
            if i is not None:
                idx[clave] = i
        if "factura" not in idx or "valor" not in idx:
            raise ValueError(
                f"{ruta.name} no parece el export de servicios del DGH: no encontré "
                "las columnas FACTURA y Vr_SERVICIO."
            )
        if not {"codigo", "cups", "cod_medicamento"} & set(idx):
            raise ValueError(
                f"{ruta.name} no trae ninguna columna de código de servicio "
                "(SERVICIOS DGH / SLNSERPRO_CUPS / CODIGO_MEDICAMENTO): sin eso el "
                f"cruce dejaría SLNSERPRO vacío. Encabezados leídos: {[h for h in headers if h]}"
            )
        if "codigo" not in idx and avisar is not None:
            avisar(
                f"  ⚠ {ruta.name} no trae la columna SERVICIOS DGH; se usará el CUPS o "
                "el código de medicamento como código del servicio."
            )

        def dato(fila, clave):
            i = idx.get(clave)
            return fila[i] if i is not None and i < len(fila) else None

        fuera: dict[str, list[LineaDgh]] = defaultdict(list)
        for fila in filas:
            if fila is None:
                continue
            factura = str(dato(fila, "factura") or "").strip()
            if not factura:
                continue
            fuera[factura_larga(factura)].append(
                LineaDgh(
                    codigo=str(dato(fila, "codigo") or "").strip(),
                    descripcion=str(dato(fila, "desc_institucional") or "").strip(),
                    desc_cups=str(dato(fila, "desc_cups") or "").strip(),
                    nombre_med=str(dato(fila, "nombre_medicamento") or "").strip(),
                    cups=str(dato(fila, "cups") or "").strip(),
                    cod_med=str(dato(fila, "cod_medicamento") or "").strip(),
                    centro=str(dato(fila, "centro_costo") or "").strip(),
                    cant=a_entero(dato(fila, "cantidad")),
                    valor=a_entero(dato(fila, "valor")),
                )
            )
        return dict(fuera)
    finally:
        wb.close()


# ─── El cruce ────────────────────────────────────────────────────────────────

TOLERANCIA_PESOS = 0.5
PUNTAJE_MINIMO = 3.0
UMBRAL_ALTA = 6.0
UMBRAL_MEDIA = 4.5

AVISO_SIN_FACTURA = "la factura no está en el export del DGH"
AVISO_SIN_CRUCE = "no se identificó el servicio: completar a mano"
AVISO_DEBIL = "cruce débil: verificar antes de subir"
AVISO_NOMBRE = "el nombre del servicio no coincide con el del DGH: confirmar el renglón"


class Cruce:
    """Resultado de buscar el servicio de una objeción en el export del DGH."""

    __slots__ = ("linea", "confianza", "motivos", "puntaje", "aviso")

    def __init__(self, linea=None, confianza="SIN CRUCE", motivos="", puntaje=0.0, aviso=""):
        self.linea = linea
        self.confianza = confianza
        self.motivos = motivos
        self.puntaje = puntaje
        self.aviso = aviso


def _mismo_valor(a: float, b: float) -> bool:
    return abs(a - b) < TOLERANCIA_PESOS


def _puntuar(
    linea: LineaDgh,
    cods: set[str],
    formas: list[str],
    unitario: int,
    valor: int,
    cantidad: int,
    ctx: dict,
):
    """Qué tanto esta línea del DGH se parece a lo que dice la objeción."""
    puntos = 0.0
    motivos: list[str] = []

    if cods and linea.codigos:
        if cods & linea.codigos:
            puntos += 3.0
            motivos.append("código")
        elif any(a.startswith(b) or b.startswith(a) for a in cods for b in linea.codigos):
            puntos += 1.5
            motivos.append("código parcial")

    mejor_desc = 0.0
    for forma in formas:
        for d in linea.descripciones:
            mejor_desc = max(mejor_desc, parecido_desc(forma, d))
    if mejor_desc >= 0.99:
        puntos += 3.0
        motivos.append("nombre exacto")
    elif mejor_desc >= 0.55:
        puntos += 3.0 * mejor_desc
        motivos.append(f"nombre {mejor_desc:.0%}")

    puntos_valor = 0.0
    if unitario and _mismo_valor(linea.unitario, unitario):
        puntos_valor = 2.0
        motivos.append("valor unitario")
    if valor and _mismo_valor(linea.valor, valor):
        puntos_valor = max(puntos_valor, 2.0)
        motivos.append("valor del renglón")
    if unitario and _mismo_valor(linea.valor, unitario):
        puntos_valor = max(puntos_valor, 1.5)
        motivos.append("valor")
    puntos += puntos_valor

    if cantidad and linea.cantidad and _mismo_valor(linea.cantidad, cantidad):
        puntos += 0.5
        motivos.append("cantidad")

    # Un valor que en toda la factura sólo lo tiene UN servicio identifica el
    # renglón aunque la entidad use otro código y otro nombre.
    if puntos_valor:
        if ctx.get("codigos_con_ese_valor") == 1:
            puntos += 1.5
            motivos.append("valor único en la factura")
        elif ctx.get("codigo_preferido") and linea.codigo == ctx["codigo_preferido"]:
            puntos += 1.5
            motivos.append("valor + nombre más parecido")

    # Y al revés: un nombre que en toda la factura lo tiene UN solo servicio
    # también identifica el renglón. Es el caso del Dispensario, que manda el
    # nombre en su propia columna y objeta la DIFERENCIA de tarifa, así que el
    # valor de la objeción nunca coincide con el del renglón.
    if mejor_desc >= 0.99 and ctx.get("codigos_con_ese_nombre") == 1:
        puntos += 1.5
        motivos.append("nombre único en la factura")

    # Código igual y nombre distinto: si el valor coincide es la misma línea y
    # la entidad escribió mal el nombre (busca el código en CUPS y no en el
    # catálogo del hospital); si el valor no coincide, no es el mismo servicio.
    if "código" in motivos and formas and mejor_desc < 0.35:
        if not puntos_valor:
            puntos -= 2.5
        motivos.append("¡el nombre no concuerda!")
    return puntos, motivos


def resolver_servicio(
    lineas: list[LineaDgh],
    *,
    codigo: str = "",
    descripcion: str = "",
    valor: int = 0,
    valor_unitario: int = 0,
    cantidad: int = 0,
) -> Cruce:
    """Busca en las líneas del DGH de ESA factura el servicio del que habla la
    objeción. Si no hay un candidato con puntaje suficiente devuelve un cruce
    vacío: nunca escribe un servicio del que no está seguro."""
    if not lineas:
        return Cruce(aviso=AVISO_SIN_FACTURA)

    cods = variantes_codigo(codigo)
    formas = formas_descripcion(descripcion)

    def coincide_valor(linea: LineaDgh) -> bool:
        return bool(
            (valor_unitario and _mismo_valor(linea.unitario, valor_unitario))
            or (valor and _mismo_valor(linea.valor, valor))
        )

    def nombre_exacto(linea: LineaDgh) -> bool:
        return any(parecido_desc(f, d) >= 0.99 for f in formas for d in linea.descripciones)

    con_ese_valor = [x for x in lineas if coincide_valor(x)]
    ctx = {
        "codigos_con_ese_valor": len({x.codigo for x in con_ese_valor}),
        "codigos_con_ese_nombre": len({x.codigo for x in lineas if nombre_exacto(x)}),
        "codigo_preferido": "",
    }
    # Si varias líneas comparten el valor, el nombre desempata.
    if ctx["codigos_con_ese_valor"] > 1 and formas:
        por_codigo: dict[str, float] = {}
        for linea in con_ese_valor:
            for forma in formas:
                for d in linea.descripciones:
                    r = parecido_desc(forma, d)
                    if r > por_codigo.get(linea.codigo, 0.0):
                        por_codigo[linea.codigo] = r
        orden = sorted(por_codigo.items(), key=lambda kv: kv[1], reverse=True)
        if orden and orden[0][1] >= 0.30 and (len(orden) == 1 or orden[0][1] - orden[1][1] >= 0.15):
            ctx["codigo_preferido"] = orden[0][0]

    mejor, mejor_pts, mejor_motivos = None, 0.0, []
    for linea in lineas:
        pts, motivos = _puntuar(linea, cods, formas, valor_unitario, valor, cantidad, ctx)
        # A igualdad de puntaje, preferir una línea todavía no usada: la entidad
        # manda una objeción por unidad y el DGH tiene un renglón por unidad.
        pts_desempate = pts - min(linea.usos, 3) * 0.15
        if pts_desempate > mejor_pts:
            mejor, mejor_pts, mejor_motivos = linea, pts_desempate, motivos

    if mejor is None or mejor_pts < PUNTAJE_MINIMO:
        return Cruce(
            motivos=", ".join(mejor_motivos), puntaje=round(mejor_pts, 2), aviso=AVISO_SIN_CRUCE
        )

    mejor.usos += 1
    confianza = (
        "ALTA" if mejor_pts >= UMBRAL_ALTA else "MEDIA" if mejor_pts >= UMBRAL_MEDIA else "BAJA"
    )
    aviso = ""
    if "¡el nombre no concuerda!" in mejor_motivos:
        confianza = "MEDIA" if confianza == "ALTA" else confianza
        aviso = AVISO_NOMBRE
    elif confianza == "BAJA":
        aviso = AVISO_DEBIL
    return Cruce(mejor, confianza, ", ".join(mejor_motivos), round(mejor_pts, 2), aviso)


# ─── Verificación de las reglas fijas del archivo de OBJECIONES ──────────────


def verificar_reglas(
    renglones: list[dict], servicios_dgh: dict[str, list[LineaDgh]] | None, avisar=None
) -> list[str]:
    """Revisa el archivo terminado contra las reglas que fijó el área.

    No cambia nada: sólo mira lo que se va a entregar y devuelve la lista de
    incumplimientos. Existe para que las reglas no dependan de que alguien se
    acuerde de revisarlas a mano:

    1. `CTNCENCOS` vacía en todos los renglones.
    2. Todo `SLNSERPRO` escrito existe en el export del DGH de ESA factura
       (prohibido inventar códigos).
    3. `CROTIPOBJ` igual para toda la factura y acorde a sus grupos de glosa:
       0 = ADMINISTRATIVA (sólo TA/FA/SO/AU/CO…), 1 = MEDICA (sólo CL),
       2 = MIXTA (CL junto con administrativas).

    Cada renglón es un dict con: factura, slnserpro, ctncencos, crotipobj y
    codigo_glosa. Puede traer además `grupos` (los grupos TA/CL/FA… que el
    propio bot usó para decidir el tipo) para las entidades cuyo código de
    glosa no dice el grupo: el ADRES usa códigos de cuatro dígitos (3106,
    3209…) y el grupo sale de su columna de clasificación.
    """
    fallas: list[str] = []

    con_centro = [r for r in renglones if r.get("ctncencos") not in (None, "")]
    if con_centro:
        fallas.append(
            f"CTNCENCOS: {len(con_centro)} renglón(es) traen dato y debe ir vacía "
            f"(primera factura: {con_centro[0].get('factura')})."
        )

    if servicios_dgh is not None:
        inventados = []
        for r in renglones:
            cod = r.get("slnserpro")
            if not cod:
                continue
            lineas = servicios_dgh.get(r.get("factura", ""), [])
            if not any(variantes_codigo(cod) & linea.codigos for linea in lineas):
                inventados.append((r.get("factura"), cod))
        if inventados:
            muestra = ", ".join(f"{f}:{c}" for f, c in inventados[:5])
            fallas.append(
                f"SLNSERPRO: {len(inventados)} código(s) no existen en el export del "
                f"DGH de su factura ({muestra})."
            )

    grupos: dict[str, set[str]] = defaultdict(set)
    tipos: dict[str, set] = defaultdict(set)
    for r in renglones:
        factura = r.get("factura", "")
        propios = r.get("grupos")
        if propios:
            grupos[factura].update(str(g)[:2].upper() for g in propios)
        else:
            grupos[factura].add(str(r.get("codigo_glosa") or "")[:2].upper())
        tipos[factura].add(r.get("crotipobj"))
    for factura, gs in grupos.items():
        tiene_cl = "CL" in gs
        tiene_admin = any(g and g != "CL" for g in gs)
        esperado = 2 if (tiene_cl and tiene_admin) else 1 if tiene_cl else 0
        if tipos[factura] != {esperado}:
            fallas.append(
                f"CROTIPOBJ: la factura {factura} tiene grupos {sorted(g for g in gs if g)} "
                f"→ debía ser {esperado} y quedó {sorted(tipos[factura])}."
            )

    if avisar is not None:
        if fallas:
            for f in fallas:
                avisar(f"  ✗ REGLA INCUMPLIDA — {f}")
        else:
            avisar(
                "  ✓ Reglas verificadas: CTNCENCOS vacía, ningún SLNSERPRO inventado "
                "y CROTIPOBJ correcto en todas las facturas."
            )
    return fallas


# ─── Reporte de trabajo del auditor ──────────────────────────────────────────

COLORES_CONFIANZA = {
    "ALTA": "E2EFDA",
    "MEDIA": "FFF2CC",
    "BAJA": "FCE4D6",
    "SIN CRUCE": "F8CBAD",
}


def _columnas_reporte(entidad: str) -> tuple[tuple[str, int], ...]:
    return (
        ("Factura", 18),
        ("Cód. glosa", 11),
        ("Valor objetado", 14),
        (f"Servicio según {entidad}", 42),
        (f"Cód. según {entidad}", 18),
        (f"Vr. unitario {entidad}", 16),
        ("Servicio en el DGH", 42),
        ("Cód. en el DGH (SLNSERPRO)", 20),
        ("Vr. unitario DGH", 15),
        ("Centro de costo", 28),
        ("Confianza", 11),
        ("Por qué cruzó", 34),
        ("Puntaje", 9),
        ("Qué revisar", 46),
        (f"Observación de {entidad}", 70),
    )


def escribir_reporte_cruce(trazas: list[dict], salida: Path, entidad: str = "la entidad") -> None:
    """Excel de trabajo del auditor: CRUCE (todo), REVISAR (lo que hay que
    confirmar a mano) y RESUMEN por factura. No es el archivo que se sube:
    es el respaldo de por qué cada objeción quedó con ese servicio."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    columnas = _columnas_reporte(entidad)
    wb = Workbook()
    fill_hdr = PatternFill("solid", fgColor="1F4E78")
    font_hdr = Font(bold=True, color="FFFFFF")

    def encabezar(ws, cols) -> None:
        for col, (nombre, ancho) in enumerate(cols, start=1):
            c = ws.cell(row=1, column=col, value=nombre)
            c.fill = fill_hdr
            c.font = font_hdr
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = ancho
        ws.freeze_panes = "A2"

    def fila(t: dict) -> list:
        return [
            t["factura"],
            t["codigo_objecion"],
            t["valor"],
            t["desc_entidad"],
            t["cod_entidad"],
            t["unitario_entidad"] or "",
            t["servicio_dgh"],
            t["cod_dgh"],
            t["unitario_dgh"],
            t["centro_costo"],
            t["confianza"],
            t["motivos"],
            t["puntaje"],
            t["aviso"],
            t["observacion"],
        ]

    ws = wb.active
    ws.title = "CRUCE"
    encabezar(ws, columnas)
    for i, t in enumerate(trazas, start=2):
        for col, valor in enumerate(fila(t), start=1):
            celda = ws.cell(row=i, column=col, value=valor)
            if col == 11:
                celda.fill = PatternFill(
                    "solid", fgColor=COLORES_CONFIANZA.get(t["confianza"], "F8CBAD")
                )

    ws2 = wb.create_sheet("REVISAR")
    encabezar(ws2, columnas)
    pendientes = [t for t in trazas if t["aviso"] or t["confianza"] in ("BAJA", "SIN CRUCE")]
    for i, t in enumerate(pendientes, start=2):
        for col, valor in enumerate(fila(t), start=1):
            ws2.cell(row=i, column=col, value=valor)

    ws3 = wb.create_sheet("RESUMEN")
    encabezar(
        ws3,
        (
            ("Factura", 18),
            ("Objeciones", 11),
            ("Valor objetado", 16),
            ("ALTA", 8),
            ("MEDIA", 8),
            ("BAJA", 8),
            ("Sin cruce", 10),
            ("Estado", 32),
        ),
    )
    por_factura: dict[str, list[dict]] = defaultdict(list)
    for t in trazas:
        por_factura[t["factura"]].append(t)
    linea = 2
    for factura in sorted(por_factura):
        grupo = por_factura[factura]
        cuenta = {k: sum(1 for t in grupo if t["confianza"] == k) for k in COLORES_CONFIANZA}
        pendiente = cuenta["BAJA"] + cuenta["SIN CRUCE"]
        estado = "Listo para subir" if not pendiente else f"Revisar {pendiente} objeción(es)"
        valores = [
            factura,
            len(grupo),
            sum(t["valor"] for t in grupo),
            cuenta["ALTA"],
            cuenta["MEDIA"],
            cuenta["BAJA"],
            cuenta["SIN CRUCE"],
            estado,
        ]
        for col, valor in enumerate(valores, start=1):
            ws3.cell(row=linea, column=col, value=valor)
        linea += 1
    totales = [
        "TOTAL",
        len(trazas),
        sum(t["valor"] for t in trazas),
        *[sum(1 for t in trazas if t["confianza"] == k) for k in COLORES_CONFIANZA],
        "",
    ]
    for col, valor in enumerate(totales, start=1):
        ws3.cell(row=linea, column=col, value=valor).font = Font(bold=True)

    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(salida))


def traza(
    *,
    factura: str,
    codigo_objecion: str,
    valor: int,
    cod_entidad: str,
    desc_entidad: str,
    unitario_entidad: int,
    observacion: str,
    cruce: Cruce,
) -> dict:
    """Una fila del reporte de cruce, con lo que se leyó y lo que se encontró."""
    return {
        "factura": factura,
        "codigo_objecion": codigo_objecion,
        "valor": valor,
        "cod_entidad": cod_entidad,
        "desc_entidad": desc_entidad,
        "unitario_entidad": unitario_entidad,
        "cod_dgh": cruce.linea.codigo if cruce.linea else "",
        "servicio_dgh": cruce.linea.descripcion if cruce.linea else "",
        "unitario_dgh": cruce.linea.unitario if cruce.linea else "",
        "centro_costo": cruce.linea.centro_costo if cruce.linea else "",
        "confianza": cruce.confianza,
        "motivos": cruce.motivos,
        "puntaje": cruce.puntaje,
        "aviso": cruce.aviso,
        "observacion": observacion,
    }
