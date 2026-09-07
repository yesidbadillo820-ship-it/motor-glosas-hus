"""Armar el ACTA SINAC desde la lista de facturas y el archivo de la EPS.

EL TRABAJO QUE REEMPLAZA. Antes de cada mesa alguien copiaba a mano, renglón
por renglón, los datos del consolidado de la EPS al formato del acta. Cien
facturas son más de doscientos renglones con nueve datos cada uno: media
jornada, y un número mal copiado se discute en la mesa como si fuera cierto.

LO QUE ESTAS PRUEBAS CUIDAN:

  · Que el cruce por número de factura funcione **en los tres formatos** en
    que aparece (`HUS0000542497`, `542497`, `HUS542497`). Si la llave falla,
    el acta sale vacía y no se sabe por qué.
  · Que la tipificación deducida sea la del Manual Único, verificada contra
    el acta 709 del Dispensario: CL→PERTINENCIA, FA→FACTURACIÓN,
    SO→SOPORTES, TA→TARIFAS.
  · Que **las de pertinencia NO se rellenen solas**. Se reparten entre MIXTA
    y MÉDICO según el caso clínico; en una mesa, un tipo mal puesto manda la
    glosa al abogado equivocado.
  · Que el acta generada la pueda leer y cuadrar el módulo que ya existía.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base

# El armador importa este modelo perezosamente (dentro de las funciones), así
# que sin traerlo acá `create_all` no crea su tabla y la prueba de la memoria
# falla con «no such table».
from app.models.db import ConciliacionTipificacionRecord  # noqa: F401
from app.services import acta_conciliacion_armar as armador
from app.services.acta_conciliacion_armar import (
    Encabezado,
    armar,
    clave_factura,
    cod_glosa_normalizado,
    factura_larga,
    leer_archivo_eps,
    leer_lista_facturas,
)
from app.services.acta_conciliacion_excel import leer_acta, revisar

MODELO = Path(__file__).resolve().parents[2] / "plantillas" / "ACTA_SINAC_modelo.xlsm"

# Los encabezados tal como los manda la EPS del Dispensario.
CABECERA_EPS = [
    "ESM REGIONALIZADORA",
    "UNIDAD SATELITE",
    "NIT IPS",
    "NOMBRE IPS",
    "MODALIDAD",
    "NUMERO CONTRATO",
    "LOTE",
    "RADICADO",
    "PREFIJO",
    "FACTURA",
    "FECHA ATENCION",
    "FECHA RADICACION",
    "MES RADICACION",
    "TIPO IDENTIFICACION",
    "IDENTIFICACION",
    "NOMBRES Y APELLIDOS",
    "VALOR FACTURA",
    "# GLOSA",
    "CODIGO CONCEPTO DE GLOSA",
    "MOTIVO DE GLOSA",
    "VALOR OBJETADO",
    "SERVICIO OBJETADO",
    "CODIGO RESPUESTA",
    "VALOR ACEPTADO RESPUESTA IPS",
    "VALOR LEVANTADADO RESPUESTA IPS",
    "VALOR RATIFICADO RESPUESTA IPS",
    "FECHA ACTA RESPUESTA",
    "NUMERO ACTA RESPUESTA",
]


def _excel(filas: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for fila in filas:
        ws.append(fila)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _glosa_eps(factura, cod, valor, motivo="MOTIVO", valor_factura=1_000_000, acta="AR002328"):
    fila = [""] * len(CABECERA_EPS)
    fila[7] = "560611"  # RADICADO
    fila[8] = "HUS"  # PREFIJO
    fila[9] = factura  # FACTURA
    fila[10] = date(2025, 9, 1)  # FECHA ATENCION
    fila[11] = date(2025, 10, 7)  # FECHA RADICACION
    fila[16] = valor_factura  # VALOR FACTURA
    fila[18] = cod  # CODIGO CONCEPTO DE GLOSA
    fila[19] = motivo  # MOTIVO DE GLOSA
    fila[20] = valor  # VALOR OBJETADO
    fila[27] = acta  # NUMERO ACTA RESPUESTA
    return fila


def _archivo_eps(glosas: list[list]) -> bytes:
    return _excel([CABECERA_EPS] + glosas)


@pytest.fixture
def db():
    eng = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(eng)
    s = sessionmaker(bind=eng)()
    try:
        yield s
    finally:
        s.close()
        eng.dispose()


# ══════════════════════════════════════════════════════════════════════════
class TestLaLlaveDelCruce:
    """Si la llave falla, el acta sale vacía y en silencio."""

    @pytest.mark.parametrize(
        "escrito", ["HUS0000542497", "542497", "HUS542497", "hus0000542497", 542497]
    )
    def test_la_misma_factura_escrita_de_cinco_formas(self, escrito):
        assert clave_factura(escrito) == "542497"

    def test_el_acta_la_quiere_con_ceros_y_prefijo(self):
        assert factura_larga("542497") == "HUS0000542497"
        assert factura_larga("HUS542497") == "HUS0000542497"

    def test_lo_que_no_es_una_factura_no_inventa_una(self):
        assert clave_factura("") == "" and clave_factura(None) == ""
        assert factura_larga("SIN NUMERO") == ""


class TestElCodigoDeGlosa:
    @pytest.mark.parametrize(
        "crudo,esperado",
        [
            ("CL03 01 CALIDAD-HONORARIOS PROFESIONALES", "CL0301"),
            ("so23 01 SOPORTES-OTROS PROCEDIMIENTOS", "SO2301"),
            ("TA02 01", "TA0201"),
            ("CO46 01 COBERTURA-COBERTURA SIN AGOTAR", "CO4601"),
        ],
    )
    def test_se_lee_como_lo_escribe_la_eps(self, crudo, esperado):
        assert cod_glosa_normalizado(crudo) == esperado

    def test_lo_ilegible_devuelve_vacio_no_un_invento(self):
        assert cod_glosa_normalizado("GLOSA TOTAL") == ""
        assert cod_glosa_normalizado(None) == ""


# ══════════════════════════════════════════════════════════════════════════
class TestLeerLosDosArchivos:
    def test_la_lista_es_una_columna_sin_encabezado(self):
        crudo = _excel([["HUS0000542497"], ["HUS0000542501"], ["HUS0000542627"]])
        assert leer_lista_facturas(crudo) == ["542497", "542501", "542627"]

    def test_si_trae_encabezado_no_se_cuela_como_factura(self):
        crudo = _excel([["FACTURA"], ["HUS0000542497"]])
        assert leer_lista_facturas(crudo) == ["542497"]

    def test_una_factura_repetida_se_cuenta_una_vez(self):
        crudo = _excel([["542497"], ["HUS0000542497"], ["542501"]])
        assert leer_lista_facturas(crudo) == ["542497", "542501"]

    def test_se_conserva_el_orden_del_auditor(self):
        """Es el orden con el que va a trabajar en la mesa."""
        crudo = _excel([["543084"], ["542497"], ["542501"]])
        assert leer_lista_facturas(crudo) == ["543084", "542497", "542501"]

    def test_las_columnas_de_la_eps_se_buscan_por_nombre(self):
        """Cada EPS manda el consolidado en otro orden; un índice fijo lee
        todo corrido en cuanto agregan una columna al principio."""
        filas, idx = leer_archivo_eps(_archivo_eps([_glosa_eps("542497", "TA0201", 6685)]))
        assert len(filas) == 1
        assert idx["factura"] == 9 and idx["valor_objetado"] == 20
        assert filas[0]["valor_objetado"] == 6685

    def test_valor_factura_no_le_roba_la_columna_a_factura(self):
        """«FACTURA» aparece dentro de «VALOR FACTURA»: si el emparejado no
        es exacto, el número de factura sale siendo un valor en pesos."""
        filas, _ = leer_archivo_eps(_archivo_eps([_glosa_eps("542497", "TA0201", 100)]))
        assert clave_factura(filas[0]["factura"]) == "542497"

    def test_un_archivo_sin_encabezado_reconocible_no_revienta(self):
        filas, idx = leer_archivo_eps(_excel([["hola"], ["mundo"]]))
        assert filas == [] and idx == {}


# ══════════════════════════════════════════════════════════════════════════
class TestLaTipificacionQueSeDeduce:
    @pytest.mark.parametrize(
        "cod,tipificacion,tipo",
        [
            ("FA0301", "FACTURACION", "ADMINISTRATIVA"),
            ("SO3601", "SOPORTES", "ADMINISTRATIVA"),
            ("TA0201", "TARIFAS", "ADMINISTRATIVA"),
        ],
    )
    def test_las_tres_deterministas_salen_completas(self, cod, tipificacion, tipo):
        r = armar(
            ["542497"],
            leer_archivo_eps(_archivo_eps([_glosa_eps("542497", cod, 100)]))[0],
            Encabezado(),
        )
        linea = r.acta.lineas[0]
        assert (linea.tipificacion, linea.tipo_glosa) == (tipificacion, tipo)
        assert r.avisos == []

    def test_pertinencia_se_tipifica_pero_NO_se_reparte_sola(self):
        """Mixta o médica lo decide un médico auditor. Rellenarlo con la más
        común manda la glosa al abogado equivocado."""
        r = armar(
            ["542497"],
            leer_archivo_eps(_archivo_eps([_glosa_eps("542497", "CL0301", 100)]))[0],
            Encabezado(),
        )
        linea = r.acta.lineas[0]
        assert linea.tipificacion == "PERTINENCIA"
        assert linea.tipo_glosa == "", "se rellenó una decisión clínica"
        assert len(r.avisos) == 1 and "MIXTA" in r.avisos[0].motivo

    def test_una_familia_que_el_acta_modelo_no_usa_se_avisa(self):
        """CO (cobertura) aparece en el archivo de la EPS y no en el acta 709.
        No se le inventa una tipificación."""
        r = armar(
            ["542497"],
            leer_archivo_eps(_archivo_eps([_glosa_eps("542497", "CO4601", 100)]))[0],
            Encabezado(),
        )
        assert r.acta.lineas[0].tipificacion == ""
        assert len(r.avisos) == 1 and "CO" in r.avisos[0].motivo


# ══════════════════════════════════════════════════════════════════════════
class TestElCruce:
    def _dos_facturas(self):
        return _archivo_eps(
            [
                _glosa_eps("542497", "TA0201", 6685),
                _glosa_eps("542497", "SO3601", 12000),
                _glosa_eps("999999", "FA0301", 50000),
            ]
        )

    def test_solo_entran_las_facturas_de_la_lista(self):
        r = armar(["542497"], leer_archivo_eps(self._dos_facturas())[0], Encabezado())
        assert len(r.acta.lineas) == 2
        assert r.fuera_de_lista == ["999999"]

    def test_una_factura_de_la_lista_sin_glosas_se_dice(self):
        r = armar(["542497", "542501"], leer_archivo_eps(self._dos_facturas())[0], Encabezado())
        assert r.sin_glosas == ["542501"]

    def test_el_acta_sigue_el_orden_de_la_lista_no_el_de_la_eps(self):
        eps = _archivo_eps([_glosa_eps("111111", "TA0201", 10), _glosa_eps("222222", "TA0201", 20)])
        r = armar(["222222", "111111"], leer_archivo_eps(eps)[0], Encabezado())
        assert [x.factura for x in r.acta.lineas] == ["HUS0000222222", "HUS0000111111"]

    def test_los_totales_los_calcula_el_motor(self):
        r = armar(["542497"], leer_archivo_eps(self._dos_facturas())[0], Encabezado())
        assert r.acta.cantidad_facturas_declarada == 1
        assert r.acta.valor_a_conciliar_declarado == 18685

    def test_nada_llega_repartido_a_la_mesa(self):
        """Aceptar, levantar y ratificar se escriben EN la audiencia."""
        r = armar(["542497"], leer_archivo_eps(self._dos_facturas())[0], Encabezado())
        for x in r.acta.lineas:
            assert (x.acepta_ips, x.levanta_entidad, x.ratificado) == (0.0, 0.0, 0.0)
            assert x.pendiente == x.glosa_inicial


# ══════════════════════════════════════════════════════════════════════════
class TestLaMemoria:
    """«Son las mismas cuentas de siempre»: lo que se decidió una vez no se
    vuelve a preguntar."""

    def test_lo_que_decidio_una_persona_llena_la_proxima_acta(self, db):
        eps = leer_archivo_eps(_archivo_eps([_glosa_eps("542497", "CL0301", 100)]))[0]

        primera = armar(["542497"], eps, Encabezado())
        assert primera.acta.lineas[0].tipo_glosa == ""
        assert len(primera.avisos) == 1

        # El médico auditor la reparte en la mesa y se guarda.
        primera.acta.lineas[0].tipo_glosa = "MEDICO"
        parte = armador.aprender(
            db, primera.acta.lineas, numero_acta="710", usuario="med@hus.gov.co"
        )
        assert parte["nuevas"] == 1

        segunda = armar(["542497"], eps, Encabezado(), memoria=armador.memoria_de(db, ["542497"]))
        assert segunda.acta.lineas[0].tipo_glosa == "MEDICO"
        assert segunda.avisos == [], "volvió a preguntar algo que ya sabía"

    def test_no_se_aprende_de_una_casilla_sin_resolver(self, db):
        eps = leer_archivo_eps(_archivo_eps([_glosa_eps("542497", "CL0301", 100)]))[0]
        r = armar(["542497"], eps, Encabezado())
        parte = armador.aprender(db, r.acta.lineas)
        assert parte["nuevas"] == 0 and parte["sin_dato"] == 1

    def test_tampoco_de_la_marca_de_pendiente(self, db):
        eps = leer_archivo_eps(_archivo_eps([_glosa_eps("542497", "CL0301", 100)]))[0]
        r = armar(["542497"], eps, Encabezado())
        r.acta.lineas[0].tipo_glosa = armador._MARCA_FALTA_HUMANO
        assert armador.aprender(db, r.acta.lineas)["nuevas"] == 0

    def test_cambiar_de_opinion_actualiza_en_vez_de_duplicar(self, db):
        eps = leer_archivo_eps(_archivo_eps([_glosa_eps("542497", "CL0301", 100)]))[0]
        r = armar(["542497"], eps, Encabezado())
        r.acta.lineas[0].tipo_glosa = "MIXTA"
        armador.aprender(db, r.acta.lineas)
        r.acta.lineas[0].tipo_glosa = "MEDICO"
        parte = armador.aprender(db, r.acta.lineas)
        assert parte["actualizadas"] == 1
        assert armador.memoria_de(db, ["542497"])[("542497", "CL0301")] == "MEDICO"

    def test_sin_facturas_no_consulta_la_base(self, db):
        assert armador.memoria_de(db, []) == {}


# ══════════════════════════════════════════════════════════════════════════
@pytest.mark.skipif(not MODELO.is_file(), reason="falta plantillas/ACTA_SINAC_modelo.xlsm")
class TestElArchivoQueSaleAlaMesa:
    def _acta(self):
        eps = leer_archivo_eps(
            _archivo_eps(
                [
                    _glosa_eps("542497", "TA0201", 6685, valor_factura=82085),
                    _glosa_eps("542501", "SO3601", 12000, valor_factura=50000),
                    _glosa_eps("542627", "CL0301", 79407, valor_factura=3757260),
                ]
            )
        )[0]
        enc = Encabezado(
            nit="901541137",
            razon_social="DISPENSARIO MEDICO",
            fecha_conciliacion=date(2026, 9, 8),
            numero_acta="710",
            periodo="SEP 2026",
        )
        r = armar(["542497", "542501", "542627"], eps, enc)
        return armador.escribir_en_modelo(r, MODELO.read_bytes(), enc), r

    def test_las_macros_del_formato_oficial_sobreviven(self):
        """El acta que se lleva a la mesa es la del formato, con sus botones."""
        import zipfile

        libro, _ = self._acta()
        assert any("vbaProject" in n for n in zipfile.ZipFile(BytesIO(libro)).namelist())

    def test_el_modulo_que_ya_existia_la_puede_leer(self):
        libro, _ = self._acta()
        acta = leer_acta(libro)
        assert len(acta.lineas) == 3
        assert acta.nit == "901541137"
        assert acta.lineas[0].factura == "HUS0000542497"
        assert acta.lineas[0].cod_glosa == "TA0201"
        assert acta.lineas[0].glosa_inicial == 6685

    def test_y_cuadra_sin_un_solo_hallazgo(self):
        """No basta con llenarla: tiene que salir cuadrada."""
        libro, _ = self._acta()
        assert revisar(leer_acta(libro))["hallazgos"] == []

    def test_lo_que_falta_decidir_queda_a_la_vista(self):
        """Una celda vacía en un acta de cien renglones se pasa por alto."""
        libro, _ = self._acta()
        acta = leer_acta(libro)
        pertinencia = next(x for x in acta.lineas if x.cod_glosa == "CL0301")
        assert armador._MARCA_FALTA_HUMANO in pertinencia.tipo_glosa

    def test_el_encabezado_se_escribe_pese_a_las_celdas_combinadas(self):
        """El formato trae casi todo el encabezado combinado, y openpyxl solo
        deja escribir en la celda de arriba a la izquierda del grupo."""
        libro, r = self._acta()
        acta = leer_acta(libro)
        assert acta.razon_social and "DISPENSARIO" in acta.razon_social.upper()
        assert acta.cantidad_facturas_declarada == 3
        assert acta.valor_a_conciliar_declarado == r.acta.valor_a_conciliar_declarado


# ══════════════════════════════════════════════════════════════════════════
@pytest.mark.skipif(not MODELO.is_file(), reason="falta plantillas/ACTA_SINAC_modelo.xlsm")
class TestElActaNoSaleReparada:
    """Excel abría el acta generada diciendo «[Reparado]» (07-09-2026).

    La causa no era un detalle de openpyxl que se pudiera ajustar: openpyxl
    **no edita** el .xlsm, lo RECONSTRUYE, y al hacerlo se lleva por delante
    3 de los 5 nombres definidos del modelo —los autofiltros de ACTA, GLOSAS
    y TRAMITES— dejando el único que sobrevive apuntando a otra hoja.

    Un acta que Excel tiene que reparar es un acta que el auditor no sabe si
    puede firmar, y esta sale de una mesa de conciliación con una EPS.
    """

    def _generar(self, n_lineas: int = 3) -> bytes:
        glosas = [_glosa_eps(f"54{i:04d}", "TA0201", 1000 + i) for i in range(n_lineas)]
        eps = leer_archivo_eps(_archivo_eps(glosas))[0]
        enc = Encabezado(nit="901541137", razon_social="DISPENSARIO MEDICO", numero_acta="710")
        r = armar([armador.clave_factura(g[9]) for g in glosas], eps, enc)
        return armador.escribir_en_modelo(r, MODELO.read_bytes(), enc)

    def _nombres(self, crudo: bytes) -> list[tuple[str, str]]:
        import re
        import zipfile

        xml = zipfile.ZipFile(BytesIO(crudo)).read("xl/workbook.xml").decode("utf-8")
        m = re.search(r"<definedNames>(.*?)</definedNames>", xml, re.S)
        return re.findall(r'name="([^"]+)"[^>]*>([^<]*)<', m.group(1)) if m else []

    def test_los_cinco_nombres_definidos_sobreviven(self):
        del_modelo = self._nombres(MODELO.read_bytes())
        del_acta = self._nombres(self._generar())
        assert len(del_acta) == len(del_modelo) == 5
        assert sorted(del_acta) == sorted(del_modelo), "openpyxl volvió a comerse los autofiltros"

    def test_cada_autofiltro_sigue_en_SU_hoja(self):
        """El síntoma más feo del defecto: el filtro de ACTA quedaba
        apuntando a Hoja3."""
        hojas = {v.split("!")[0] for n, v in self._nombres(self._generar()) if "Filter" in n}
        assert {"ACTA", "GLOSAS", "TRAMITES", "Hoja3"} == hojas

    def test_no_falta_ninguna_parte_salvo_las_cachés(self):
        """`calcChain` y `sharedStrings` se dejan fuera A PROPÓSITO: Excel las
        rehace, y un calcChain viejo —de antes de escribir— es otra de las
        cosas que disparan la reparación."""
        import zipfile

        del_modelo = set(zipfile.ZipFile(MODELO.open("rb")).namelist())
        del_acta = set(zipfile.ZipFile(BytesIO(self._generar())).namelist())
        assert del_modelo - del_acta == {"xl/calcChain.xml", "xl/sharedStrings.xml"}

    def test_la_configuracion_de_impresora_vuelve(self):
        """El acta se imprime para firmarla: perder el área de impresión y la
        configuración de página no es cosmético."""
        import zipfile

        partes = set(zipfile.ZipFile(BytesIO(self._generar())).namelist())
        assert any("printerSettings" in p for p in partes)
        assert any(n == "_xlnm.Print_Area" for n, _ in self._nombres(self._generar()))


@pytest.mark.skipif(not MODELO.is_file(), reason="falta plantillas/ACTA_SINAC_modelo.xlsm")
class TestSinCuadriculaVaciaDebajo:
    """El modelo trae 260 renglones con bordes ya puestos: un acta de tres
    líneas salía con 257 filas de cuadrícula vacía debajo."""

    def _acta(self, n_lineas: int):
        glosas = [_glosa_eps(f"54{i:04d}", "TA0201", 1000 + i) for i in range(n_lineas)]
        eps = leer_archivo_eps(_archivo_eps(glosas))[0]
        enc = Encabezado(nit="901541137", razon_social="DISPENSARIO MEDICO")
        r = armar([armador.clave_factura(g[9]) for g in glosas], eps, enc)
        libro = armador.escribir_en_modelo(r, MODELO.read_bytes(), enc)
        return openpyxl.load_workbook(BytesIO(libro))["ACTA"], len(r.acta.lineas)

    @staticmethod
    def _tiene_borde(celda) -> bool:
        b = celda.border
        return bool(
            b
            and any(
                getattr(b, lado) and getattr(b, lado).style
                for lado in ("left", "right", "top", "bottom")
            )
        )

    @pytest.mark.parametrize("n", [1, 3, 25])
    def test_debajo_de_la_ultima_linea_no_queda_cuadricula(self, n):
        ws, usadas = self._acta(n)
        sobrantes = [
            f
            for f in range(12 + usadas, 272)
            if any(self._tiene_borde(ws.cell(f, c)) for c in range(3, 24))
        ]
        assert sobrantes == [], f"quedaron {len(sobrantes)} filas con bordes"

    def test_las_lineas_de_verdad_conservan_su_formato(self):
        """Limpiar lo que sobra no puede llevarse lo que sí va."""
        ws, usadas = self._acta(3)
        assert all(self._tiene_borde(ws.cell(f, 5)) for f in range(12, 12 + usadas))

    def test_el_pie_del_acta_no_se_movio(self):
        """Por eso NO se borran las filas: borrarlas correría el pie hacia
        arriba y rompería sus celdas combinadas."""
        ws, _ = self._acta(3)
        del_modelo = openpyxl.load_workbook(MODELO, keep_vba=True)["ACTA"]
        assert len(ws.merged_cells.ranges) == len(del_modelo.merged_cells.ranges)
        assert any(str(r) == "C274:W276" for r in ws.merged_cells.ranges)
