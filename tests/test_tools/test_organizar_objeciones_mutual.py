"""Tests del organizador de objeciones de MUTUAL SER
(tools/organizar_objeciones_mutual.py).

Lo que se cuida acá:

- Que el NOMBRE del servicio se rescate de la observación. La columna
  «SERVICIO» de MUTUAL trae el código, no el nombre; sin ese rescate el cruce
  contra el DGH se quedaría sin con qué desempatar.
- Las cuatro reglas fijas del archivo de OBJECIONES (ver CLAUDE.md):
  CTNCENCOS vacía, CROTIPOBJ por factura, SLNSERPRO sin inventar y el 100% de
  los renglones.
- Que un archivo con otras columnas se detenga con un mensaje claro en vez de
  entregar algo silenciosamente malo.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

# El script vive en tools/ (sin __init__.py): lo importamos por ruta.
_TOOLS = Path(__file__).resolve().parent.parent.parent / "tools"
sys.path.insert(0, str(_TOOLS))

import organizar_objeciones_mutual as org  # noqa: E402

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook, load_workbook  # noqa: E402


ENCABEZADOS = [
    "Número de factura",
    "SERVICIO",
    "Cantidad facturada",
    "Valor glosado",
    "Concepto de glosa",
    "Código de glosa",
    "Observacion",
]

# Los renglones reales del lote del 7 de septiembre, en miniatura.
FILAS = [
    [
        "HUS0000544271",
        "903883",
        1,
        "$\xa0304.500",
        "Consultas, interconsultas y atenciones (visitas) domiciliarias - TARIFAS",
        "TA0201",
        "La tecnología 903883 - GLUCOSA SEMIAUTOMATIZADA [GLUCOMETRIA] no se "
        "encuentra dentro del contrato número 20352.",
    ],
    [
        "HUS0000544271",
        "931001",
        1,
        "$\xa05.088",
        "Recargos no pactados - TARIFAS",
        "TA2901",
        "La tarifa facturada 31800 no coincide con la tarifa 26712.0 definida en "
        "el contrato para la tecnologia 931001 - TERAPIA FISICA INTEGRAL.",
    ],
]


def _xlsx(ruta: Path, filas=None, encabezados=None, hoja="CONSOLIDADO") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws.append(encabezados or ENCABEZADOS)
    for f in filas if filas is not None else FILAS:
        ws.append(f)
    wb.save(str(ruta))
    return ruta


def _xlsx_dgh(ruta: Path, filas) -> Path:
    """Export de servicios facturados del DGH, con los encabezados reales."""
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "SERVICIOS DGH",
            "DESCRIPCION INSTITUCIONAL",
            "SLNSERPRO_CUPS",
            "DESCRIPCION CUPS",
            "FACTURA",
            "CAT_SERVICIOS",
            "Vr_SERVICIO",
            "SALDO_FACT",
        ]
    )
    for f in filas:
        ws.append(f)
    wb.save(str(ruta))
    return ruta


def _dgh_del_lote(tmp_path: Path) -> Path:
    return _xlsx_dgh(
        tmp_path / "dgh.xlsx",
        [
            # El DGH nombra la glucometría con su código interno, y el CUPS es
            # el que manda MUTUAL: por ahí es que cruzan.
            ["19992", "GLUCOMETRIA", "903883", "GLUCOSA SEMIAUTOMATIZADA", "HUS0000544271", 29, 304500, 900000],
            ["931001", "TERAPIA FISICA INTEGRAL", "931001", "TERAPIA FISICA", "HUS0000544271", 1, 31800, 900000],
        ],
    )  # fmt: skip


def _leer_objeciones(ruta: Path) -> list[dict]:
    ws = load_workbook(str(ruta))["OBJECIONES"]
    columnas = [c.value for c in ws[1]]
    return [dict(zip(columnas, fila)) for fila in ws.iter_rows(min_row=2, values_only=True)]


# ---------------------------------------------------------------------------
# El nombre del servicio, que viene escondido en la observación
# ---------------------------------------------------------------------------


class TestNombreDelServicio:
    def test_lo_saca_del_texto_de_la_tecnologia(self):
        obs = (
            "La tecnología 903883 - GLUCOSA SEMIAUTOMATIZADA [GLUCOMETRIA] no se "
            "encuentra dentro del contrato número 20352."
        )
        assert org.nombre_del_servicio(obs, "903883") == "GLUCOSA SEMIAUTOMATIZADA [GLUCOMETRIA]"

    def test_lo_saca_cuando_el_codigo_va_al_final(self):
        obs = (
            "La tarifa facturada 31800 no coincide con la tarifa 26712.0 definida "
            "en el contrato para la tecnologia 931001 - TERAPIA FISICA INTEGRAL."
        )
        assert org.nombre_del_servicio(obs, "931001") == "TERAPIA FISICA INTEGRAL"

    def test_sirve_con_codigos_del_hospital(self):
        obs = (
            "La tecnología FMQ0817 - FILTRO HME ADULTO/PEDIATRICO REF 352/5996 no "
            "se encuentra dentro del contrato número 20352."
        )
        assert (
            org.nombre_del_servicio(obs, "FMQ0817") == "FILTRO HME ADULTO/PEDIATRICO REF 352/5996"
        )

    def test_no_se_confunde_con_los_guiones_del_concepto(self):
        """'Recargos no pactados - TARIFAS' también tiene un guion: el corte
        se hace a partir del código, no del primer guion que aparezca."""
        obs = "Recargos no pactados - TARIFAS. La tecnología 903883 - HEMOGRAMA no se encuentra."
        assert org.nombre_del_servicio(obs, "903883") == "HEMOGRAMA"

    def test_cuando_la_observacion_no_nombra_el_servicio_devuelve_vacio(self):
        """No se inventa un nombre: el cruce se apoyará en el código y el valor."""
        assert org.nombre_del_servicio("El servicio facturado 389002 no se encuentra habilitado", "389002") == ""  # fmt: skip
        assert org.nombre_del_servicio("texto suelto sin el código", "903883") == ""
        assert org.nombre_del_servicio("", "903883") == ""
        assert org.nombre_del_servicio(None, "903883") == ""


# ---------------------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------------------


class TestLectura:
    def test_lee_los_renglones_con_sus_datos(self, tmp_path):
        objeciones = org.leer_mutual(_xlsx(tmp_path / "m.xlsx"))
        assert len(objeciones) == 2
        primera = objeciones[0]
        assert primera["cxc"] == "HUS0000544271"
        assert primera["codigo"] == "TA0201"
        assert primera["cod_servicio"] == "903883"
        assert primera["servicio"] == "GLUCOSA SEMIAUTOMATIZADA [GLUCOMETRIA]"
        assert primera["cantidad"] == 1

    def test_el_valor_viene_con_signo_de_pesos(self, tmp_path):
        """MUTUAL manda '$ 304.500' como texto, con espacio duro."""
        objeciones = org.leer_mutual(_xlsx(tmp_path / "m.xlsx"))
        assert [o["valor"] for o in objeciones] == [304500, 5088]

    def test_los_encabezados_toleran_tildes_y_mayusculas(self, tmp_path):
        ruta = _xlsx(
            tmp_path / "m.xlsx",
            encabezados=[
                "NUMERO FACTURA",
                "codigo servicio",
                "Cantidad",
                "Valor glosa",
                "Concepto",
                "Cod glosa",
                "Observaciones",
            ],
        )
        assert len(org.leer_mutual(ruta)) == 2

    def test_el_formato_corto_del_8_de_septiembre(self, tmp_path):
        """MUTUAL cambia las columnas entre lotes.

        El del 8 de septiembre llegó con CINCO: la del código rebautizada
        «Tecnología» y sin cantidad ni concepto. La pantalla lo rechazaba con
        «no trae las columnas cod_servicio».
        """
        ruta = _xlsx(
            tmp_path / "m8.xlsx",
            encabezados=[
                "Número de factura",
                "Tecnología",
                "Valor glosado",
                "Código de glosa",
                "Observacion",
            ],
            filas=[
                [
                    "HUS0000544255",
                    "890494",
                    165400,
                    "TA0201",
                    "La tecnología 890494 - INTERCONSULTA HOSPITALARIA POR UROLOGIA "
                    "no se encuentra dentro del contrato número 20352.",
                ]
            ],
            hoja="Hoja1",
        )
        objeciones = org.leer_mutual(ruta)
        assert len(objeciones) == 1
        o = objeciones[0]
        assert o["cod_servicio"] == "890494"
        assert o["servicio"] == "INTERCONSULTA HOSPITALARIA POR UROLOGIA"
        assert o["valor"] == 165400
        assert o["codigo"] == "TA0201"
        # Las dos que no vinieron quedan vacías, no rompen nada.
        assert not o["cantidad"]
        assert o["concepto"] == ""

    def test_sin_cantidad_el_renglon_igual_se_arma(self, tmp_path):
        """El valor unitario se deja en 0 en vez de dividir por cero."""
        ruta = _xlsx(
            tmp_path / "m8.xlsx",
            encabezados=[
                "Número de factura",
                "Tecnología",
                "Valor glosado",
                "Código de glosa",
                "Observacion",
            ],
            filas=[["HUS0000544255", "890494", 165400, "TA0201", "sin nombre"]],
            hoja="Hoja1",
        )
        filas = org.construir_filas(org.leer_mutual(ruta), _fecha())
        assert len(filas) == 1
        assert filas[0]["CROVALOBJ"] == 165400

    def test_un_archivo_de_otra_entidad_se_detiene_con_mensaje_claro(self, tmp_path):
        ruta = _xlsx(
            tmp_path / "otro.xlsx",
            filas=[["a", "b", "c"]],
            encabezados=["COSA", "OTRA COSA", "UNA MAS"],
        )
        with pytest.raises(ValueError, match="no trae las columnas"):
            org.leer_mutual(ruta)

    def test_hoja_inexistente(self, tmp_path):
        with pytest.raises(ValueError, match="no existe"):
            org.leer_mutual(_xlsx(tmp_path / "m.xlsx"), hoja="NO_ESTA")

    def test_renglon_sin_factura_se_omite(self, tmp_path):
        filas = FILAS + [["", "999", 1, "$ 100", "x", "TA0101", "y"]]
        assert len(org.leer_mutual(_xlsx(tmp_path / "m.xlsx", filas=filas))) == 2


# ---------------------------------------------------------------------------
# Las cuatro reglas fijas
# ---------------------------------------------------------------------------


class TestCodigoDeGlosaLimpio:
    @pytest.mark.parametrize(
        "crudo, esperado",
        [("TA0201", "TA0201"), ("TA02 01", "TA0201"), ("ta0201", "TA0201"), ("", "")],
    )
    def test_deja_solo_el_codigo(self, crudo, esperado):
        assert org.codigo_glosa(crudo) == esperado


class TestTipoDeObjecionPorFactura:
    def test_los_tres_valores(self):
        assert org.crotipobj_factura({"TA", "FA"}) == 0
        assert org.crotipobj_factura({"CL"}) == 1
        assert org.crotipobj_factura({"CL", "TA"}) == 2

    def test_el_lote_real_es_administrativo(self, tmp_path):
        filas = org.construir_filas(org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha())
        assert {f["CROTIPOBJ"] for f in filas} == {0}

    def test_una_glosa_clinica_vuelve_mixta_toda_la_factura(self, tmp_path):
        filas_excel = FILAS + [
            ["HUS0000544271", "903883", 1, "$ 1.000", "Pertinencia", "CL0101", "no pertinente"]
        ]
        filas = org.construir_filas(
            org.leer_mutual(_xlsx(tmp_path / "m.xlsx", filas=filas_excel)), _fecha()
        )
        assert {f["CROTIPOBJ"] for f in filas} == {2}


class TestCentroDeCostoVacio:
    def test_sale_vacia_siempre(self, tmp_path):
        filas = org.construir_filas(org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha())
        assert {f["CTNCENCOS"] for f in filas} == {None}


class TestCruceContraElDgh:
    def test_llena_slnserpro_con_el_codigo_del_hospital(self, tmp_path):
        """MUTUAL manda el CUPS (903883); el DGH lo tiene como 19992."""
        from _cruce_dgh import leer_servicios_dgh

        servicios = leer_servicios_dgh(_dgh_del_lote(tmp_path))
        filas = org.construir_filas(
            org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha(), servicios
        )
        assert [f["SLNSERPRO"] for f in filas] == ["19992", "931001"]

    def test_sin_export_no_se_escribe_ningun_codigo(self, tmp_path):
        """Sin el DGH no hay con qué comprobar: la celda queda vacía."""
        filas = org.construir_filas(org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha())
        assert {f["SLNSERPRO"] for f in filas} == {None}

    def test_un_servicio_que_no_esta_en_esa_factura_queda_vacio(self, tmp_path):
        from _cruce_dgh import leer_servicios_dgh

        servicios = leer_servicios_dgh(
            _xlsx_dgh(
                tmp_path / "dgh.xlsx",
                [["111111", "OTRA COSA", "111111", "OTRA", "HUS0000999999", 1, 999, 999]],
            )
        )
        filas = org.construir_filas(
            org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha(), servicios
        )
        assert {f["SLNSERPRO"] for f in filas} == {None}

    def test_el_archivo_lleva_el_100_por_ciento_de_los_renglones(self, tmp_path):
        from _cruce_dgh import leer_servicios_dgh

        servicios = leer_servicios_dgh(_xlsx_dgh(tmp_path / "dgh.xlsx", []))
        filas = org.construir_filas(
            org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha(), servicios
        )
        assert len(filas) == 2
        assert [f["CROVALOBJ"] for f in filas] == [304500, 5088]

    def test_deja_traza_para_la_hoja_revisar(self, tmp_path):
        from _cruce_dgh import leer_servicios_dgh

        servicios = leer_servicios_dgh(_dgh_del_lote(tmp_path))
        trazas: list[dict] = []
        org.construir_filas(
            org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha(), servicios, trazas
        )
        assert len(trazas) == 2
        assert all(t["factura"] == "HUS0000544271" for t in trazas)


class TestElArchivoTerminadoCumpleLasReglas:
    def test_verificar_reglas_no_encuentra_fallas(self, tmp_path):
        from _cruce_dgh import leer_servicios_dgh, verificar_reglas

        dgh = _dgh_del_lote(tmp_path)
        servicios = leer_servicios_dgh(dgh)
        filas = org.construir_filas(
            org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha(), servicios
        )
        salida = tmp_path / "OBJECIONES_MUTUAL.xlsx"
        org.escribir_objeciones(filas, salida)
        escritas = _leer_objeciones(salida)
        assert (
            verificar_reglas(
                [
                    {
                        "factura": f["CRNCXC"],
                        "slnserpro": f["SLNSERPRO"],
                        "ctncencos": f["CTNCENCOS"],
                        "crotipobj": f["CROTIPOBJ"],
                        "codigo_glosa": f["CRNCONOBJ"],
                    }
                    for f in escritas
                ],
                servicios,
            )
            == []
        )


# ---------------------------------------------------------------------------
# Escritura y CLI
# ---------------------------------------------------------------------------


def _fecha():
    import datetime as _dt

    return _dt.datetime(2026, 9, 7)


class TestEscritura:
    def test_el_excel_sale_con_las_16_columnas(self, tmp_path):
        filas = org.construir_filas(org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha())
        salida = tmp_path / "o.xlsx"
        org.escribir_objeciones(filas, salida)
        ws = load_workbook(str(salida))["OBJECIONES"]
        assert [c.value for c in ws[1]] == list(org.COLUMNAS_OBJECIONES)
        assert ws.max_row == 3

    def test_un_archivo_por_factura(self, tmp_path):
        filas = org.construir_filas(org.leer_mutual(_xlsx(tmp_path / "m.xlsx")), _fecha())
        generados = org.escribir_por_factura(filas, tmp_path / "salida", "OBJECIONES_MUTUAL")
        assert [p.name for p in generados] == ["OBJECIONES_MUTUAL_HUS0000544271.xlsx"]


class TestCli:
    def test_de_punta_a_punta(self, tmp_path):
        entrada = _xlsx(tmp_path / "MUTUAL_7_SEPTIEMBRE.xlsx")
        dgh = _dgh_del_lote(tmp_path)
        salida = tmp_path / "OBJECIONES_MUTUAL_07092026.xlsx"
        reporte = tmp_path / "CRUCE_MUTUAL_07092026.xlsx"
        rc = org.main(
            [
                "--entrada",
                str(entrada),
                "--servicios-dgh",
                str(dgh),
                "--salida",
                str(salida),
                "--consolidado",
                "--reporte-cruce",
                str(reporte),
                "--fecha",
                "2026-09-07",
            ]
        )
        assert rc == 0
        filas = _leer_objeciones(salida)
        assert len(filas) == 2
        assert [f["SLNSERPRO"] for f in filas] == ["19992", "931001"]
        assert {f["CTNCENCOS"] for f in filas} == {None}
        assert {f["CROTIPOBJ"] for f in filas} == {0}
        assert set(load_workbook(str(reporte)).sheetnames) >= {"CRUCE", "REVISAR", "RESUMEN"}

    def test_entrada_inexistente(self, tmp_path):
        assert org.main(["--entrada", str(tmp_path / "no.xlsx"), "--salida", str(tmp_path)]) == 1

    def test_reporte_sin_export_es_error_claro(self, tmp_path):
        entrada = _xlsx(tmp_path / "m.xlsx")
        rc = org.main(
            [
                "--entrada",
                str(entrada),
                "--salida",
                str(tmp_path / "o.xlsx"),
                "--consolidado",
                "--reporte-cruce",
                str(tmp_path / "r.xlsx"),
            ]
        )
        assert rc == 1

    def test_export_inexistente(self, tmp_path):
        entrada = _xlsx(tmp_path / "m.xlsx")
        rc = org.main(
            [
                "--entrada",
                str(entrada),
                "--salida",
                str(tmp_path / "o.xlsx"),
                "--consolidado",
                "--servicios-dgh",
                str(tmp_path / "no.xlsx"),
            ]
        )
        assert rc == 1


# ---------------------------------------------------------------------------
# Doble glosa: el mismo servicio objetado bajo dos conceptos
# ---------------------------------------------------------------------------


def _obj(cxc, cod_servicio, cantidad, valor, codigo, concepto="X", observacion=""):
    """Una objeción ya leída, como la devuelve `leer_mutual`."""
    return {
        "fila_excel": 0,
        "cxc": cxc,
        "codigo": codigo,
        "cod_servicio": cod_servicio,
        "servicio": "",
        "cantidad": cantidad,
        "concepto": concepto,
        "observacion": observacion,
        "valor": valor,
    }


class TestDobleGlosa:
    """El caso del lote del 7-sep: sumar todas las filas daba $26.636.056
    cuando MUTUAL reportó $24.462.346."""

    def test_mismo_servicio_bajo_dos_conceptos_se_cuenta_una_vez(self):
        objeciones = [
            _obj("HUS0000544271", "FMQ0463", 1, 244400, "TA0201"),
            _obj("HUS0000544271", "FMQ0463", 1, 244400, "TA0601"),
        ]
        quedan = org.fusionar_dobles_glosas(objeciones)
        assert len(quedan) == 1
        assert quedan[0]["valor"] == 244400
        assert quedan[0]["fusionadas"] == [("TA0601", 244400)]

    def test_con_valores_distintos_manda_el_mayor(self):
        """389002: glosado por tarifa ($54.594) y por no estar habilitado
        ($181.900). MUTUAL sólo cuenta el mayor."""
        objeciones = [
            _obj("HUS0000544271", "389002", 1, 54594, "TA2901"),
            _obj("HUS0000544271", "389002", 1, 181900, "FA1305"),
        ]
        quedan = org.fusionar_dobles_glosas(objeciones)
        assert len(quedan) == 1
        assert quedan[0]["valor"] == 181900
        assert quedan[0]["codigo"] == "FA1305"
        assert quedan[0]["fusionadas"] == [("TA2901", 54594)]

    def test_no_junta_cantidades_distintas_del_mismo_servicio(self):
        """FMQ0178 se factura en dos renglones (x23 y x5): son dos glosas."""
        objeciones = [
            _obj("HUS0000544271", "FMQ0178", 23, 23000, "TA0201"),
            _obj("HUS0000544271", "FMQ0178", 23, 23000, "TA0601"),
            _obj("HUS0000544271", "FMQ0178", 5, 4000, "TA0201"),
            _obj("HUS0000544271", "FMQ0178", 5, 4000, "TA0601"),
        ]
        quedan = org.fusionar_dobles_glosas(objeciones)
        assert sorted(o["valor"] for o in quedan) == [4000, 23000]

    def test_no_junta_servicios_de_facturas_distintas(self):
        objeciones = [
            _obj("HUS0000000001", "FMQ0463", 1, 244400, "TA0201"),
            _obj("HUS0000000002", "FMQ0463", 1, 244400, "TA0601"),
        ]
        assert len(org.fusionar_dobles_glosas(objeciones)) == 2

    def test_el_mismo_codigo_repetido_NO_se_fusiona(self):
        """Dos renglones de verdad con el mismo código —como las 9 terapias
        respiratorias del Dispensario— no son doble glosa: juntarlos sería
        borrar una objeción real. Se dejan y se avisa."""
        objeciones = [
            _obj("HUS0000544271", "939403", 1, 851, "TA5701"),
            _obj("HUS0000544271", "939403", 1, 851, "TA5701"),
        ]
        avisos: list[str] = []
        quedan = org.fusionar_dobles_glosas(objeciones, avisar=avisos.append)
        assert len(quedan) == 2
        assert avisos and "NO se fusiona" in avisos[0]

    def test_un_servicio_solo_no_se_toca(self):
        objeciones = [_obj("HUS0000544271", "903883", 1, 304500, "TA0201")]
        quedan = org.fusionar_dobles_glosas(objeciones)
        assert len(quedan) == 1
        assert "fusionadas" not in quedan[0]

    def test_conserva_el_orden_original(self):
        objeciones = [
            _obj("HUS0000544271", "AAA", 1, 100, "TA0201"),
            _obj("HUS0000544271", "BBB", 1, 200, "TA0201"),
            _obj("HUS0000544271", "BBB", 1, 200, "TA0601"),
            _obj("HUS0000544271", "CCC", 1, 300, "TA0201"),
        ]
        quedan = org.fusionar_dobles_glosas(objeciones)
        assert [o["cod_servicio"] for o in quedan] == ["AAA", "BBB", "CCC"]


class TestLaDobleGlosaLlegaAlArchivo:
    FILAS_DOBLES = [
        [
            "HUS0000544271",
            "FMQ0463",
            1,
            "$\xa0244.400",
            "Consultas, interconsultas y atenciones (visitas) domiciliarias - TARIFAS",
            "TA0201",
            "La tecnología FMQ0463 - SISTEMA DE SUCCION CERRADA NO 12 no se "
            "encuentra dentro del contrato número 20352.",
        ],
        [
            "HUS0000544271",
            "FMQ0463",
            1,
            "$\xa0244.400",
            "Dispositivos médicos - TARIFAS",
            "TA0601",
            "La tecnología FMQ0463 - SISTEMA DE SUCCION CERRADA NO 12 no se "
            "encuentra dentro del contrato número 20352.",
        ],
    ]

    def _correr(self, tmp_path):
        entrada = _xlsx(tmp_path / "MUTUAL.xlsx", filas=self.FILAS_DOBLES)
        dgh = _xlsx_dgh(
            tmp_path / "dgh.xlsx",
            [
                ["QX9999", "SISTEMA DE SUCCION CERRADA NO 12", "FMQ0463", "SUCCION",
                 "HUS0000544271", 1, 244400, 900000],
            ],
        )  # fmt: skip
        salida = tmp_path / "OBJECIONES.xlsx"
        assert (
            org.main(
                [
                    "--entrada",
                    str(entrada),
                    "--servicios-dgh",
                    str(dgh),
                    "--salida",
                    str(salida),
                    "--consolidado",
                    "--fecha",
                    "2026-09-07",
                ]
            )
            == 0
        )
        return _leer_objeciones(salida)

    def test_el_archivo_trae_un_solo_renglon(self, tmp_path):
        filas = self._correr(tmp_path)
        assert len(filas) == 1
        assert filas[0]["CROVALOBJ"] == 244400

    def test_el_otro_codigo_queda_anotado_en_la_observacion(self, tmp_path):
        """No se pierde: el auditor tiene que poder ver bajo qué más lo objetaron."""
        filas = self._correr(tmp_path)
        assert "también glosado como TA0601 $244400" in filas[0]["CRDOBSERV"]

    def test_el_codigo_absorbido_cuenta_para_el_tipo_de_la_factura(self, tmp_path):
        """TA0201 y TA0601 son administrativos: la factura sigue en 0."""
        filas = self._correr(tmp_path)
        assert filas[0]["CROTIPOBJ"] == 0

    def test_una_glosa_clinica_absorbida_deja_la_factura_mixta(self):
        """Si el código que se absorbe es CL, la factura es mixta: la glosa
        clínica existe aunque no se cuente su plata dos veces."""
        objeciones = [
            _obj("HUS0000544271", "AAA", 1, 500, "TA0201"),
            _obj("HUS0000544271", "AAA", 1, 300, "CL0101"),
        ]
        quedan = org.fusionar_dobles_glosas(objeciones)
        filas = org.construir_filas(quedan, _fecha())
        assert filas[0]["CROTIPOBJ"] == 2
