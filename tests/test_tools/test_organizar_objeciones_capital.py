"""Tests del organizador de objeciones de CAPITAL SALUD
(tools/organizar_objeciones_capital.py).

Lo propio de CAPITAL, y por eso lo que más se prueba: no manda el código del
servicio (sólo el nombre, con guiones en vez de espacios) y tampoco manda el
código de glosa en columna propia — ese va pegado al principio de
`DescripcionGlosa`.

Además, las cuatro reglas fijas del archivo de OBJECIONES (ver CLAUDE.md).
"""

from __future__ import annotations

import datetime as _dt
import sys
from pathlib import Path

import pytest

# El script vive en tools/ (sin __init__.py): lo importamos por ruta.
_TOOLS = Path(__file__).resolve().parent.parent.parent / "tools"
sys.path.insert(0, str(_TOOLS))

import organizar_objeciones_capital as org  # noqa: E402

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook, load_workbook  # noqa: E402


FECHA = _dt.datetime(2026, 9, 9)

ENCABEZADOS = ["FACTURA", "servicio", "Cantidad", "DescripcionGlosa", "OBSERVACION", "VALOR GLOSA"]

GLOSA_SO = "SO4201 - Existe ausencia total, parcial o inconsistencia de la lista de precios"
GLOSA_TA = "TA0801 - Los cargos por apoyo diagnóstico presentan diferencias"
GLOSA_CL = "CL0101 - El servicio no es pertinente"


def _xlsx(ruta: Path, filas=None, encabezados=None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(encabezados or ENCABEZADOS)
    for f in filas if filas is not None else _FILAS:
        ws.append(f)
    wb.save(str(ruta))
    return ruta


_FILAS = [
    ["HUS0000532847", "KETAMINA", 1, GLOSA_SO, "no adjuntan factura", 30200],
    [
        "HUS0000532847",
        "INTERNACION-ADULTOS-COMPLEJIDAD-ALTA",
        3,
        GLOSA_TA,
        "mayor valor cobrado",
        148800,
    ],
]


def _xlsx_dgh(ruta: Path, filas) -> Path:
    """Export de servicios facturados del DGH, con los encabezados reales."""
    wb = Workbook()
    ws = wb.active
    ws.append(
        ["FACTURA", "SERVICIOS DGH", "DESCRIPCION INSTITUCIONAL", "CAT SERVICIOS", "VR_SERVICIO"]
    )
    for f in filas:
        ws.append(f)
    wb.save(str(ruta))
    return ruta


def _dgh(tmp_path):
    from _cruce_dgh import leer_servicios_dgh

    ruta = _xlsx_dgh(
        tmp_path / "dgh.xlsx",
        [
            ["HUS0000532847", "20048691-1", "KETAMINA CLORHIDRATO I.V. AMP X 500MG/10ML", 1, 30200],
            [
                "HUS0000532847",
                "129A02",
                "INTERNACION ADULTOS COMPLEJIDAD ALTA HABITACION MULTIPLE",
                3,
                148800,
            ],
        ],
    )
    return leer_servicios_dgh(ruta)


# ─── El código de glosa, que viene pegado a su descripción ───────────────────


class TestCodigoYConcepto:
    @pytest.mark.parametrize(
        "crudo, codigo",
        [
            (GLOSA_SO, "SO4201"),
            (GLOSA_TA, "TA0801"),
            ("SO42 01 - texto con el código partido", "SO4201"),
            ("so4201 - en minúscula", "SO4201"),
            ("SO4201 sin guion separador", "SO4201"),
        ],
    )
    def test_separa_el_codigo(self, crudo, codigo):
        assert org.codigo_y_concepto(crudo)[0] == codigo

    def test_devuelve_el_texto_sin_el_codigo(self):
        assert org.codigo_y_concepto(GLOSA_SO)[1].startswith("Existe ausencia total")

    def test_sin_codigo_reconocible_no_inventa_uno(self):
        assert org.codigo_y_concepto("Glosa sin código al principio") == (
            "",
            "Glosa sin código al principio",
        )

    def test_celda_vacia(self):
        assert org.codigo_y_concepto(None) == ("", "")


# ─── El nombre del servicio, que viene con guiones ───────────────────────────


class TestNombreServicio:
    def test_los_guiones_se_vuelven_espacios(self):
        assert (
            org.nombre_servicio("INTERNACION-ADULTOS-COMPLEJIDAD-ALTA")
            == "INTERNACION ADULTOS COMPLEJIDAD ALTA"
        )

    def test_conserva_los_parentesis(self):
        assert org.nombre_servicio("HEMOGRAMA-IV-(HEMOGLOBINA-HEMATOCRITO)") == (
            "HEMOGRAMA IV (HEMOGLOBINA HEMATOCRITO)"
        )

    def test_un_nombre_sin_guiones_no_cambia(self):
        assert org.nombre_servicio("KETAMINA") == "KETAMINA"

    def test_celda_vacia(self):
        assert org.nombre_servicio(None) == ""

    def test_asi_es_como_el_cruce_puede_comparar(self):
        """Con guiones el parecido se hunde; con espacios es exacto."""
        from _cruce_dgh import norm_desc, parecido_desc

        capital = "INTERNACION-ADULTOS-COMPLEJIDAD-ALTA"
        dgh = norm_desc("INTERNACION ADULTOS COMPLEJIDAD ALTA")
        assert parecido_desc(capital, dgh) < 0.6
        assert parecido_desc(norm_desc(org.nombre_servicio(capital)), dgh) == 1.0


# ─── Lectura del Excel ───────────────────────────────────────────────────────


class TestLectura:
    def test_lee_las_objeciones(self, tmp_path):
        objeciones = org.leer_capital(_xlsx(tmp_path / "c.xlsx"))
        assert len(objeciones) == 2
        assert objeciones[0]["cxc"] == "HUS0000532847"
        assert objeciones[0]["codigo"] == "SO4201"
        assert objeciones[0]["servicio"] == "KETAMINA"
        assert objeciones[0]["valor"] == 30200
        assert objeciones[1]["servicio"] == "INTERNACION ADULTOS COMPLEJIDAD ALTA"

    def test_si_falta_una_columna_indispensable_se_detiene(self, tmp_path):
        ruta = _xlsx(
            tmp_path / "c.xlsx",
            filas=[["HUS0000532847", "KETAMINA", 1, "x"]],
            encabezados=["FACTURA", "servicio", "Cantidad", "OTRA COSA"],
        )
        with pytest.raises(ValueError, match="no trae las columnas"):
            org.leer_capital(ruta)

    def test_el_error_dice_cuál_falta(self, tmp_path):
        ruta = _xlsx(
            tmp_path / "c.xlsx",
            filas=[["HUS0000532847", "KETAMINA"]],
            encabezados=["FACTURA", "servicio"],
        )
        with pytest.raises(ValueError, match="glosa"):
            org.leer_capital(ruta)

    def test_las_filas_vacias_se_saltan(self, tmp_path):
        ruta = _xlsx(tmp_path / "c.xlsx", filas=[*_FILAS, [None] * 6, ["", "", "", "", "", ""]])
        assert len(org.leer_capital(ruta)) == 2

    def test_archivo_sin_datos(self, tmp_path):
        assert org.leer_capital(_xlsx(tmp_path / "c.xlsx", filas=[])) == []


# ─── Las reglas fijas del archivo de OBJECIONES ──────────────────────────────


class TestTipoDeObjecionPorFactura:
    def test_solo_administrativas_es_cero(self):
        assert org.crotipobj_factura({"SO", "TA", "FA"}) == 0

    def test_solo_clinicas_es_uno(self):
        assert org.crotipobj_factura({"CL"}) == 1

    def test_mezcla_es_dos(self):
        assert org.crotipobj_factura({"CL", "SO"}) == 2

    def test_toda_la_factura_lleva_el_mismo_tipo(self, tmp_path):
        ruta = _xlsx(
            tmp_path / "c.xlsx",
            filas=[
                ["HUS0000532847", "KETAMINA", 1, GLOSA_SO, "x", 30200],
                ["HUS0000532847", "CONSULTA", 1, GLOSA_CL, "x", 1000],
                ["HUS0000999999", "OTRO", 1, GLOSA_SO, "x", 500],
            ],
        )
        filas = org.construir_filas(org.leer_capital(ruta), FECHA)
        tipos = {}
        for f in filas:
            tipos.setdefault(f["CRNCXC"], set()).add(f["CROTIPOBJ"])
        assert tipos == {"HUS0000532847": {2}, "HUS0000999999": {0}}


class TestCentroDeCostoVacio:
    def test_sale_vacia_siempre(self, tmp_path):
        filas = org.construir_filas(
            org.leer_capital(_xlsx(tmp_path / "c.xlsx")), FECHA, _dgh(tmp_path)
        )
        assert {f["CTNCENCOS"] for f in filas} == {None}


class TestCruceContraElDgh:
    def test_llena_el_codigo_que_capital_no_manda(self, tmp_path):
        filas = org.construir_filas(
            org.leer_capital(_xlsx(tmp_path / "c.xlsx")), FECHA, _dgh(tmp_path)
        )
        assert [f["SLNSERPRO"] for f in filas] == ["20048691-1", "129A02"]

    def test_sin_export_la_columna_queda_vacia(self, tmp_path):
        filas = org.construir_filas(org.leer_capital(_xlsx(tmp_path / "c.xlsx")), FECHA)
        assert {f["SLNSERPRO"] for f in filas} == {None}

    def test_lo_que_no_cruza_no_se_borra(self, tmp_path):
        from _cruce_dgh import leer_servicios_dgh

        vacio = leer_servicios_dgh(_xlsx_dgh(tmp_path / "dgh0.xlsx", []))
        filas = org.construir_filas(org.leer_capital(_xlsx(tmp_path / "c.xlsx")), FECHA, vacio)
        assert len(filas) == 2
        assert all(f["SLNSERPRO"] is None for f in filas)
        assert [f["CROVALOBJ"] for f in filas] == [30200, 148800]

    def test_deja_traza_para_la_hoja_revisar(self, tmp_path):
        trazas: list[dict] = []
        org.construir_filas(
            org.leer_capital(_xlsx(tmp_path / "c.xlsx")), FECHA, _dgh(tmp_path), trazas
        )
        assert len(trazas) == 2
        assert all(t["confianza"] in ("ALTA", "MEDIA", "BAJA", "SIN CRUCE") for t in trazas)

    def test_los_renglones_repetidos_se_conservan(self, tmp_path):
        """CAPITAL repite el servicio tantas veces como renglones tenga la
        factura; el DGH los tiene igual. Juntarlos borraría objeciones."""
        ruta = _xlsx(
            tmp_path / "c.xlsx",
            filas=[["HUS0000532847", "KETAMINA", 1, GLOSA_SO, "x", 30200]] * 3,
        )
        filas = org.construir_filas(org.leer_capital(ruta), FECHA, _dgh(tmp_path))
        assert len(filas) == 3
        assert sum(f["CROVALOBJ"] for f in filas) == 90600


class TestElArchivoTerminadoCumpleLasReglas:
    def test_verificar_reglas_no_encuentra_fallas(self, tmp_path):
        from _cruce_dgh import verificar_reglas

        servicios = _dgh(tmp_path)
        filas = org.construir_filas(org.leer_capital(_xlsx(tmp_path / "c.xlsx")), FECHA, servicios)
        fallas = verificar_reglas(
            [
                {
                    "factura": f["CRNCXC"],
                    "slnserpro": f["SLNSERPRO"],
                    "ctncencos": f["CTNCENCOS"],
                    "crotipobj": f["CROTIPOBJ"],
                    "codigo_glosa": f["CRNCONOBJ"],
                }
                for f in filas
            ],
            servicios,
        )
        assert fallas == []


# ─── Escritura y CLI ─────────────────────────────────────────────────────────


class TestEscritura:
    def test_el_excel_sale_con_las_16_columnas(self, tmp_path):
        filas = org.construir_filas(org.leer_capital(_xlsx(tmp_path / "c.xlsx")), FECHA)
        salida = tmp_path / "o.xlsx"
        org.escribir_objeciones(filas, salida)
        ws = load_workbook(str(salida))["OBJECIONES"]
        assert [c.value for c in ws[1]] == list(org.COLUMNAS_OBJECIONES)
        assert ws.max_row == 3

    def test_crdobserv_lleva_codigo_servicio_y_valor(self, tmp_path):
        filas = org.construir_filas(org.leer_capital(_xlsx(tmp_path / "c.xlsx")), FECHA)
        assert filas[0]["CRDOBSERV"].startswith("SO4201 KETAMINA")
        assert filas[0]["CRDOBSERV"].endswith("$30200")


class TestCli:
    def _correr(self, argv):
        return org.main(argv)

    def test_de_punta_a_punta(self, tmp_path):
        entrada = _xlsx(tmp_path / "CAPITAL.xlsx")
        dgh = _xlsx_dgh(
            tmp_path / "dgh.xlsx",
            [
                [
                    "HUS0000532847",
                    "20048691-1",
                    "KETAMINA CLORHIDRATO I.V. AMP X 500MG/10ML",
                    1,
                    30200,
                ],
                [
                    "HUS0000532847",
                    "129A02",
                    "INTERNACION ADULTOS COMPLEJIDAD ALTA HABITACION MULTIPLE",
                    3,
                    148800,
                ],
            ],
        )
        salida = tmp_path / "OBJECIONES.xlsx"
        reporte = tmp_path / "CRUCE.xlsx"
        rc = self._correr(
            [
                "--entrada",
                str(entrada),
                "--servicios-dgh",
                str(dgh),
                "--salida",
                str(salida),
                "--consolidado",
                "--reporte-cruce",
                str(reporte),
                "--fecha",
                "2026-09-09",
            ]
        )
        assert rc == 0
        assert salida.is_file()
        assert set(load_workbook(str(reporte)).sheetnames) >= {"CRUCE", "REVISAR", "RESUMEN"}

    def test_reporte_sin_export_es_error_claro(self, tmp_path):
        rc = self._correr(
            [
                "--entrada",
                str(_xlsx(tmp_path / "c.xlsx")),
                "--salida",
                str(tmp_path / "o.xlsx"),
                "--consolidado",
                "--reporte-cruce",
                str(tmp_path / "r.xlsx"),
            ]
        )
        assert rc == 1

    def test_entrada_inexistente(self, tmp_path):
        assert (
            self._correr(
                ["--entrada", str(tmp_path / "no.xlsx"), "--salida", str(tmp_path / "o.xlsx")]
            )
            == 1
        )

    def test_fecha_al_reves_se_rechaza(self, tmp_path):
        with pytest.raises(SystemExit):
            self._correr(
                [
                    "--entrada",
                    str(_xlsx(tmp_path / "c.xlsx")),
                    "--salida",
                    str(tmp_path / "o.xlsx"),
                    "--consolidado",
                    "--fecha",
                    "09-09-2026",
                ]
            )
