"""Tests del bot que saca el PDF y el Word de respuesta por factura.

Los dos formatos son los que el auditor ya venía armando a mano:
`RTA_ADRES_<FACTURA>.pdf` y `Reporte_Factura_<FACTURA>_<GESTOR>.docx`.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

_TOOLS = Path(__file__).resolve().parent.parent.parent / "tools"
sys.path.insert(0, str(_TOOLS))

openpyxl = pytest.importorskip("openpyxl")
pytest.importorskip("docx")
pytest.importorskip("reportlab")

import respuestas_adres_por_factura as rr  # noqa: E402

from docx import Document  # noqa: E402

ENCABEZADO = {
    rr.COL_RADICACION: "Número Radicación",
    rr.COL_FACTURA: "Número Factura",
    rr.COL_PAQUETE: "Número Paquete",
    rr.COL_VALOR_GLOSADO: "Valor Glosado",
    rr.COL_DOC_VICTIMA: "Tip- Num Doc Victima",
    rr.COL_DESC_ELEMENTO: "Descripción Elemento",
    rr.COL_DESC_GLOSA: "Descripción Glosa",
    rr.COL_OBSERVACION: "OBSERVACION",
    rr.COL_RTA_COMPLETA: "RTA GLOSA COMPLETA",
    rr.COL_VALOR_ACEPTADO: "VALOR ACEPTADO",
    rr.COL_GESTOR: "GESTOR",
}


def fila(
    factura="HUS352890",
    *,
    radicado="14344788",
    paquete="31068",
    glosado=0,
    documento="CC-1005338825",
    elemento="Consulta de urgencias",
    glosa="3202- La consulta no esta justificada",
    observacion="SE OBJETA",
    rta="3202-SE OBJETA -Consulta de urgenciasPOR VALOR DE  $85.800",
    aceptado=0,
    gestor="OSCAR",
):
    return {
        rr.COL_RADICACION: radicado,
        rr.COL_FACTURA: factura,
        rr.COL_PAQUETE: paquete,
        rr.COL_VALOR_GLOSADO: glosado,
        rr.COL_DOC_VICTIMA: documento,
        rr.COL_DESC_ELEMENTO: elemento,
        rr.COL_DESC_GLOSA: glosa,
        rr.COL_OBSERVACION: observacion,
        rr.COL_RTA_COMPLETA: rta,
        rr.COL_VALOR_ACEPTADO: aceptado,
        rr.COL_GESTOR: gestor,
    }


def escribir_macro(ruta: Path, filas) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=21, value="-   CANTIDAD ACEPTADA  . POR VALOR $91.617.467")
    for col, texto in ENCABEZADO.items():
        ws.cell(row=2, column=col, value=texto)
    for i, f in enumerate(filas, start=3):
        for col, valor in f.items():
            ws.cell(row=i, column=col, value=valor)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


def parrafos(ruta: Path) -> list[str]:
    return [p.text for p in Document(str(ruta)).paragraphs]


class TestLeerLaMacro:
    def test_agrupa_por_factura_con_sus_datos(self, tmp_path):
        macro = escribir_macro(
            tmp_path / "m.xlsx",
            [fila("HUS1", glosado=100), fila("HUS1", glosado=200), fila("HUS2", glosado=50)],
        )
        datos = rr.leer_macro(macro)
        assert set(datos) == {"1", "2"}
        fac = datos["1"]
        assert len(fac.glosas) == 2
        assert fac.radicacion == "14344788"
        assert fac.documento_paciente == "CC-1005338825"
        assert fac.gestor == "OSCAR"

    def test_las_filas_sin_causal_son_glosa_total(self, tmp_path):
        macro = escribir_macro(
            tmp_path / "m.xlsx", [fila(glosa="3202- algo"), fila(glosa=""), fila(glosa=None)]
        )
        fac = rr.leer_macro(macro)["352890"]
        assert [g.glosa_total for g in fac.glosas] == [False, True, True]

    def test_filtra_por_gestor_sin_importar_mayusculas(self, tmp_path):
        macro = escribir_macro(
            tmp_path / "m.xlsx",
            [fila("HUS1", gestor="CAROLINA"), fila("HUS2", gestor="OSCAR")],
        )
        assert list(rr.leer_macro(macro, gestor="carolina")) == ["1"]

    def test_filtra_por_paquete(self, tmp_path):
        macro = escribir_macro(
            tmp_path / "m.xlsx",
            [fila("HUS1", paquete="31068"), fila("HUS2", paquete="31078")],
        )
        assert list(rr.leer_macro(macro, paquete="31078")) == ["2"]


class TestElOrdenYLosNombres:
    def test_lo_aceptado_va_de_primeras(self, tmp_path):
        """La regla del auditor: todo lo que se acepta sale de primero."""
        macro = escribir_macro(
            tmp_path / "m.xlsx",
            [
                fila(observacion="SE OBJETA", rta="A objetada"),
                fila(observacion="SE SUBSANA", rta="B subsanada"),
                fila(observacion="SE ACEPTA", rta="C aceptada", aceptado=100),
            ],
        )
        fac = rr.leer_macro(macro)["352890"]
        assert [g.rta_completa for g in fac.responder()] == [
            "C aceptada",
            "A objetada",
            "B subsanada",
        ]

    def test_los_nombres_de_los_archivos(self):
        assert rr.nombre_pdf("HUS352890") == "RTA_ADRES_HUS352890.pdf"
        assert rr.nombre_word("HUS298253", "CAROLINA") == "Reporte_Factura_HUS298253_CAROLINA.docx"
        assert rr.nombre_word("HUS1", "") == "Reporte_Factura_HUS1.docx"


class TestElWord:
    def _word(self, tmp_path, filas, **kw):
        macro = escribir_macro(tmp_path / "m.xlsx", filas)
        fac = rr.leer_macro(macro)["352890"]
        destino = tmp_path / "salida.docx"
        rr.generar_word(fac, destino, **kw)
        return destino

    def test_encabeza_con_lo_aceptado(self, tmp_path):
        ruta = self._word(
            tmp_path,
            [fila(observacion="SE ACEPTA", rta="C aceptada", aceptado=1820)],
            consecutivo="123",
        )
        assert parrafos(ruta)[0] == (
            "123 SE IDENTIFICAN HALLAZGOS Y SE REALIZAN LOS AJUSTES CORRESPONDIENTES "
            "POR GLOSA ACEPTADA PARCIAL POR VALOR DE $1.820"
        )

    def test_sin_consecutivo_el_encabezado_no_lleva_numero(self, tmp_path):
        ruta = self._word(tmp_path, [fila(observacion="SE ACEPTA", rta="C", aceptado=1820)])
        assert parrafos(ruta)[0].startswith("SE IDENTIFICAN HALLAZGOS")

    def test_sin_nada_aceptado_no_hay_encabezado(self, tmp_path):
        ruta = self._word(tmp_path, [fila(observacion="SE OBJETA", rta="solo objetada")])
        assert parrafos(ruta)[0] == "solo objetada"

    def test_una_respuesta_por_parrafo_tal_como_la_escribio_el_auditor(self, tmp_path):
        ruta = self._word(
            tmp_path,
            [fila(rta="PRIMERA respuesta"), fila(rta="SEGUNDA respuesta")],
        )
        assert parrafos(ruta)[:2] == ["PRIMERA respuesta", "SEGUNDA respuesta"]

    def test_la_nota_de_extemporaneidad_solo_si_se_pide(self, tmp_path):
        sin = self._word(tmp_path, [fila(rta="X")])
        assert not any("EXTEMPORANEA" in p for p in parrafos(sin))
        con = self._word(tmp_path, [fila(rta="X")], extemporanea=True)
        assert parrafos(con)[-1] == rr.NOTA_EXTEMPORANEA

    def test_las_glosas_totales_no_entran_pero_se_avisan(self, tmp_path):
        ruta = self._word(
            tmp_path,
            [fila(rta="con causal"), fila(glosa="", rta="sin causal", glosado=9000)],
        )
        textos = parrafos(ruta)
        assert "sin causal" not in textos
        assert any("no se relacionan 1 renglón(es) de GLOSA TOTAL por $9.000" in p for p in textos)

    def test_con_la_opcion_las_glosas_totales_si_entran(self, tmp_path):
        ruta = self._word(
            tmp_path,
            [fila(rta="con causal"), fila(glosa="", rta="sin causal")],
            incluir_totales=True,
        )
        textos = parrafos(ruta)
        assert "sin causal" in textos
        assert not any("GLOSA TOTAL" in p for p in textos)

    def test_avisa_las_glosas_sin_decidir(self, tmp_path):
        ruta = self._word(tmp_path, [fila(rta="X"), fila(observacion="", rta="Y")])
        assert any("quedan 1 glosa(s) sin decidir" in p for p in parrafos(ruta))

    def test_una_fila_sin_respuesta_escrita_no_deja_un_parrafo_vacio(self, tmp_path):
        ruta = self._word(tmp_path, [fila(rta="X"), fila(rta="")])
        assert [p for p in parrafos(ruta) if p.strip()] == ["X"]


class TestElPDF:
    def test_usa_el_texto_del_auditor_tal_cual(self, tmp_path):
        from app.services.evidencia_adres_pdf import rta_glosa_completa

        assert rta_glosa_completa({"rta_completa": "LO QUE ESCRIBIO EL AUDITOR"}) == (
            "LO QUE ESCRIBIO EL AUDITOR"
        )

    def test_si_no_viene_escrita_la_arma_como_antes(self):
        from app.services.evidencia_adres_pdf import rta_glosa_completa

        armada = rta_glosa_completa(
            {"causal_codigo": "3202", "decision": "SE OBJETA", "descripcion": "Consulta"}
        )
        assert armada.startswith("3202-SE OBJETA-Consulta")

    def test_una_respuesta_larguisima_no_tumba_el_pdf(self, tmp_path):
        """Hay respuestas de más de 2.500 caracteres: una sola celda más alta
        que la hoja. El PDF tiene que armarse igual."""
        from app.services.evidencia_adres_pdf import generar_pdf_evidencia

        datos = {
            "factura": "HUS1",
            "radicacion": "1",
            "documento_paciente": "CC-1",
            "glosas": [
                {"valor_glosado": 100, "descripcion": "X", "rta_completa": "PALABRA " * 500}
            ],
            "resumen": {"valor_glosado": 100, "valor_aceptado": 0},
        }
        assert generar_pdf_evidencia(datos)[:4] == b"%PDF"

    def test_los_datos_del_encabezado_llegan_al_pdf(self, tmp_path):
        macro = escribir_macro(tmp_path / "m.xlsx", [fila(glosado=85800)])
        fac = rr.leer_macro(macro)["352890"]
        datos = rr.datos_para_pdf(fac)
        assert datos["factura"] == "HUS352890"
        assert datos["radicacion"] == "14344788"
        assert datos["documento_paciente"] == "CC-1005338825"
        assert datos["resumen"]["valor_glosado"] == 85800

    def test_el_resumen_cuenta_las_glosas_totales_aparte(self, tmp_path):
        macro = escribir_macro(
            tmp_path / "m.xlsx",
            [fila(glosado=100), fila(glosa="", glosado=900)],
        )
        resumen = rr.datos_para_pdf(rr.leer_macro(macro)["352890"])["resumen"]
        assert resumen["valor_glosado"] == 100
        assert resumen["glosas_totales_ocultas"] == 1
        assert resumen["valor_glosas_totales"] == 900


class TestCorridaCompleta:
    def _correr(self, tmp_path, filas, extra=()):
        macro = escribir_macro(tmp_path / "m.xlsx", filas)
        salida = tmp_path / "out"
        assert rr.main(["--macro", str(macro), "--salida", str(salida), *extra]) == 0
        return salida

    def test_deja_un_pdf_y_un_word_por_factura(self, tmp_path):
        salida = self._correr(
            tmp_path, [fila("HUS1", gestor="CAROLINA"), fila("HUS2", gestor="OSCAR")]
        )
        assert (salida / "RTA_ADRES_HUS1.pdf").exists()
        assert (salida / "Reporte_Factura_HUS1_CAROLINA.docx").exists()
        assert (salida / "RTA_ADRES_HUS2.pdf").exists()
        assert (salida / "Reporte_Factura_HUS2_OSCAR.docx").exists()

    def test_carpeta_por_factura(self, tmp_path):
        salida = self._correr(tmp_path, [fila("HUS1")], extra=["--carpeta-por-factura"])
        assert (salida / "HUS1" / "RTA_ADRES_HUS1.pdf").exists()

    def test_solo_un_gestor(self, tmp_path):
        salida = self._correr(
            tmp_path,
            [fila("HUS1", gestor="CAROLINA"), fila("HUS2", gestor="OSCAR")],
            extra=["--gestor", "CAROLINA"],
        )
        assert (salida / "RTA_ADRES_HUS1.pdf").exists()
        assert not (salida / "RTA_ADRES_HUS2.pdf").exists()

    def test_el_reporte_csv(self, tmp_path):
        macro = escribir_macro(
            tmp_path / "m.xlsx",
            [fila("HUS1", glosado=100, observacion="SE ACEPTA", aceptado=40)],
        )
        reporte = tmp_path / "r.csv"
        assert (
            rr.main(
                [
                    "--macro",
                    str(macro),
                    "--salida",
                    str(tmp_path / "out"),
                    "--reporte-csv",
                    str(reporte),
                ]
            )
            == 0
        )
        lineas = reporte.read_text(encoding="utf-8-sig").splitlines()
        assert lineas[0].startswith("FACTURA;GESTOR;PDF;WORD")
        campos = lineas[1].split(";")
        assert campos[0] == "HUS1" and campos[7] == "100.00" and campos[8] == "40.00"

    def test_el_resumen_en_pantalla(self, tmp_path, capsys):
        self._correr(tmp_path, [fila("HUS1", glosado=100, observacion="SE ACEPTA", aceptado=40)])
        texto = capsys.readouterr().out
        assert "Facturas          : 1" in texto
        assert "documentos       : 1 PDF y 1 Word" in texto
        assert "aceptado         : $40" in texto

    def test_una_macro_sin_facturas_avisa(self, tmp_path, caplog):
        macro = escribir_macro(tmp_path / "m.xlsx", [])
        assert rr.main(["--macro", str(macro), "--salida", str(tmp_path / "out")]) == 1
        assert "no trae ninguna factura" in caplog.text
