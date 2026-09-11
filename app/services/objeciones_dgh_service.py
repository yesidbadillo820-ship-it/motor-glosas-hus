"""Objeciones para DGH — armar el archivo de cargue desde la aplicación.

Es el mismo trabajo que el auditor hacía pidiendo el cruce por chat: sube el
Excel de glosas que mandó la entidad y el export de servicios facturados del
DGH, y la pantalla devuelve los dos archivos de siempre —el que se sube y el
respaldo con la hoja REVISAR— más el resumen del cruce para mirarlo antes de
descargar nada.

**No reimplementa nada.** Los bots de `tools/` son la única fuente de la
verdad: este servicio los importa, les pasa los dos archivos y recoge lo que
producen. Si una regla cambia, cambia en el bot y la pantalla la hereda.

Entidades que sabe leer hoy:

    FAMISANAR    4 columnas; el servicio va escondido en el texto de la glosa
    DISPENSARIO  5 columnas; trae el nombre del servicio en columna propia
    SAVIA SALUD  8 columnas; trae código y nombre del servicio
    SALUD TOTAL  6 columnas; sólo el nombre del servicio, sin código
    SANITAS      7 columnas; ojo, la 2ª se llama «NUMERO DE FACTURA» pero trae
                 el código de glosa
    VCO          consolidado del acta del portal VCO (COOSALUD, FIDUPREVISORA,
                 SAVIA…): 10 columnas, con el acta en la primera
    EMSSANAR     no manda Excel: son los PDF de objeción de ripslink, uno por
                 factura. Se pueden subir varios de una vez.
    CAPITAL      6 columnas. No manda código de servicio (sólo el nombre, con
                 guiones) ni columna de código de glosa: ese va dentro de
                 «DescripcionGlosa».
    ADRES        Excel de glosas del ADRES. Tiene motor propio (homologación
                 SOAT↔CUPS, topes de valor); acepta el homologador Gold
                 Standard como segundo archivo.
    MUTUAL SER   consolidado de 5 o 7 columnas, según el lote. La columna del
                 servicio («SERVICIO» o «Tecnología») trae el código, no el
                 nombre: el nombre va dentro de la observación.

Las reglas fijas del formato (CTNCENCOS vacía, CROTIPOBJ por factura,
SLNSERPRO sin códigos inventados, el 100% de los renglones) están en CLAUDE.md
y las aplica y comprueba el propio bot: acá solo se recoge el veredicto.
"""

from __future__ import annotations

import importlib
import sys
import tempfile
from dataclasses import dataclass, field
from collections.abc import Sequence
from datetime import date, datetime
from pathlib import Path
from typing import Any

_TOOLS = Path(__file__).resolve().parent.parent.parent / "tools"

# Tamaño máximo por archivo. El export del DGH más grande visto ronda los 2 MB
# (2.264 servicios); 25 MB deja margen sobrado sin permitir un volcado entero.
MAX_BYTES = 25 * 1024 * 1024

CONFIANZAS = ("ALTA", "MEDIA", "BAJA", "SIN CRUCE")

# Tope del DGH: no recibe más de 300 facturas en un archivo de cargue.
MAX_FACTURAS_DGH = 300


class ErrorObjeciones(RuntimeError):
    """Algo que el auditor puede entender y corregir."""


def _cargar(modulo: str):
    """Importa un bot de `tools/` sin dejar `sys.path` sucio."""
    ruta = str(_TOOLS)
    agregado = ruta not in sys.path
    if agregado:
        sys.path.insert(0, ruta)
    try:
        return importlib.import_module(modulo)
    finally:
        if agregado and ruta in sys.path:
            sys.path.remove(ruta)


# ─── Qué entidad mandó el archivo ────────────────────────────────────────────


@dataclass(frozen=True)
class Entidad:
    """Una entidad que el motor sabe procesar."""

    id: str
    nombre: str
    # Nombre corto para los archivos, el mismo que usa el bot por consola
    # (OBJECIONES_SALUDTOTAL_…, no OBJECIONES_SALUD_…).
    corto: str
    modulo: str
    # Cómo se reconoce su archivo: encabezados que deben estar (normalizados).
    senas: tuple[str, ...]
    columnas: int
    ayuda: str
    # "xlsx" (un Excel) o "pdf" (uno o varios PDF, como EMSSANAR).
    formato: str = "xlsx"

    def como_dict(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "nombre": self.nombre,
            "corto": self.corto,
            "columnas": self.columnas,
            "ayuda": self.ayuda,
            "formato": self.formato,
        }


ENTIDADES: tuple[Entidad, ...] = (
    Entidad(
        id="famisanar",
        nombre="FAMISANAR",
        corto="FAMISANAR",
        modulo="organizar_objeciones_famisanar",
        senas=("NRO_FACTURA", "CODIGO_DEVOLUCION", "VALOR DEVOLUCION", "OBSERVACION"),
        columnas=4,
        ayuda="Export DEVYGLOSAS: factura, código de devolución, valor y observación.",
    ),
    Entidad(
        id="dispensario",
        nombre="Dispensario Médico",
        corto="DISPENSARIO",
        modulo="organizar_objeciones_dispensario",
        senas=("FACTURA", "VALOR GLOSA INICIAL", "SERVICIO OBJETADO", "CODIGO GLOSA INICIAL"),
        columnas=5,
        ayuda="Excel de glosa inicial, con el servicio objetado en su propia columna.",
    ),
    Entidad(
        id="savia",
        nombre="SAVIA SALUD",
        corto="SAVIA",
        modulo="organizar_objeciones_savia",
        senas=("NUMERO_FACTURA", "COD_SERVICIO", "VALOR_GLOSA", "MOTIVO_ESP_GLOSA_VALOR_A"),
        columnas=8,
        ayuda="Export de 8 columnas, con el código y el nombre del servicio en columnas propias.",
    ),
    Entidad(
        id="saludtotal",
        nombre="SALUD TOTAL",
        corto="SALUDTOTAL",
        modulo="organizar_objeciones_saludtotal",
        senas=("NUMEROFAC_", "NOMBRESERVICIO", "VALORGLOSATOTALXSERV", "CODMOTVGLOSAESPC"),
        columnas=6,
        ayuda="Export NotificacionGLS. No manda código de servicio: se ubica por nombre y valor.",
    ),
    Entidad(
        id="sanitas",
        nombre="SANITAS",
        corto="SANITAS",
        modulo="organizar_objeciones_sanitas",
        senas=("NUMERO DE FACTURA", "VALOR REAL GLOSA", "CODIGO PROCEDIMIENTO"),
        columnas=7,
        ayuda="Hoja «Glosa». Ojo: la 2ª columna dice «NUMERO DE FACTURA» pero trae el código de glosa.",
    ),
    Entidad(
        id="vco",
        nombre="VCO (acta del portal)",
        corto="VCO",
        modulo="organizar_objeciones_vco",
        senas=("NUMERO FACTURA", "VALOR GLOSA", "CODIGO GLOSA ESPECIFICA", "DESCRIPCION SERVICIO"),
        columnas=10,
        ayuda="Consolidado del acta del portal VCO (COOSALUD, FIDUPREVISORA, SAVIA…).",
    ),
    Entidad(
        id="emssanar",
        nombre="EMSSANAR",
        corto="EMSSANAR",
        modulo="organizar_objeciones_emssanar",
        # No se reconoce por encabezados: llega en PDF y se detecta por eso.
        senas=(),
        columnas=0,
        ayuda="PDF de objeción de ripslink («Objeción a Factura N° HUS…»), uno por factura.",
        formato="pdf",
    ),
    Entidad(
        id="adres",
        nombre="ADRES",
        corto="ADRES",
        modulo="organizar_objeciones_adres",
        senas=("COD ELEMENTO", "VALOR GLOSADO", "VALOR ACEPTADO", "CODIGO NUMERICO"),
        columnas=0,
        ayuda=(
            "Excel de glosas del ADRES. Se puede subir además el Homologador "
            "Gold Standard CUPS↔SOAT como segundo archivo."
        ),
    ),
    Entidad(
        id="mutual",
        nombre="MUTUAL SER",
        corto="MUTUAL",
        modulo="organizar_objeciones_mutual",
        # MUTUAL cambia las columnas entre lotes (7 el 7-sep, 5 el 8-sep, con
        # «SERVICIO» rebautizada «Tecnología»). Se listan las de los dos
        # formatos: las estables reconocen ambos y las que varían sólo suman.
        senas=(
            "NUMERO DE FACTURA",
            "VALOR GLOSADO",
            "CODIGO DE GLOSA",
            "TECNOLOGIA",
            "CANTIDAD FACTURADA",
            "CONCEPTO DE GLOSA",
        ),
        columnas=0,  # varía entre lotes
        ayuda=(
            "Consolidado de MUTUAL SER (5 o 7 columnas, según el lote). La columna "
            "«SERVICIO» o «Tecnología» trae el código, no el nombre; el nombre se "
            "saca de la observación."
        ),
    ),
    Entidad(
        id="capital",
        nombre="CAPITAL SALUD",
        corto="CAPITAL_SALUD",
        modulo="organizar_objeciones_capital",
        # «DESCRIPCIONGLOSA» (sin espacio) es la seña que no comparte con nadie:
        # ahí adentro va el código de glosa, que no viene en columna propia.
        senas=("FACTURA", "DESCRIPCIONGLOSA", "VALOR GLOSA", "OBSERVACION"),
        columnas=6,
        ayuda=(
            "Export de CAPITAL SALUD. No manda código de servicio (sólo el nombre, "
            "con guiones) ni columna de código de glosa: ese va dentro de "
            "«DescripcionGlosa»."
        ),
    ),
)


def catalogo_entidades() -> list[dict[str, Any]]:
    """Las entidades que la pantalla puede ofrecer."""
    return [e.como_dict() for e in ENTIDADES]


def _encabezados(datos: bytes) -> list[str]:
    """Los encabezados de la primera hoja con contenido, normalizados."""
    from openpyxl import load_workbook

    cruce = _cargar("_cruce_dgh")
    with tempfile.TemporaryDirectory(prefix="objeciones-") as carpeta:
        ruta = Path(carpeta) / "entrada.xlsx"
        ruta.write_bytes(datos)
        try:
            wb = load_workbook(filename=str(ruta), data_only=True, read_only=True)
        except Exception as exc:  # archivo corrupto o que no es Excel
            raise ErrorObjeciones(
                "No pude abrir el archivo: ¿seguro que es un Excel (.xlsx)?"
            ) from exc
        try:
            for ws in wb.worksheets:
                fila = next(ws.iter_rows(values_only=True), None)
                if fila and any(fila):
                    return [cruce.norm_header(c) for c in fila if c is not None]
            return []
        finally:
            wb.close()


def es_pdf(datos: bytes) -> bool:
    """Un PDF empieza por %PDF: es lo único que no llega en Excel."""
    return datos[:4] == b"%PDF"


def detectar_entidad(datos: bytes) -> Entidad:
    """Adivina de quién es el archivo por sus encabezados.

    Gana la entidad que tenga más señas presentes; si ninguna llega a dos, se
    devuelve un error que dice qué se leyó, en vez de procesar a ciegas con el
    lector equivocado (que es como salen los archivos malos sin que se note).

    Un PDF es de EMSSANAR: es la única entidad que no manda Excel.
    """
    if es_pdf(datos):
        return entidad_por_id("emssanar")

    cabecera = set(_encabezados(datos))
    if not cabecera:
        raise ErrorObjeciones("El archivo de la entidad está vacío.")

    puntajes = [(sum(1 for s in e.senas if s in cabecera), e) for e in ENTIDADES]
    puntajes.sort(key=lambda p: p[0], reverse=True)
    mejor, entidad = puntajes[0]
    if mejor < 2:
        leidos = ", ".join(sorted(cabecera)[:8])
        raise ErrorObjeciones(
            "No reconozco de qué entidad es este archivo. Encabezados leídos: "
            f"{leidos}. Elegí la entidad a mano si el formato cambió."
        )
    return entidad


def entidad_por_id(id_entidad: str) -> Entidad:
    for e in ENTIDADES:
        if e.id == id_entidad:
            return e
    raise ErrorObjeciones(f"No conozco la entidad '{id_entidad}'.")


# ─── Utilidades ──────────────────────────────────────────────────────────────


def _fecha(valor: str | date | datetime | None) -> datetime:
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime(valor.year, valor.month, valor.day)
    texto = (valor or "").strip()
    if not texto:
        hoy = date.today()
        return datetime(hoy.year, hoy.month, hoy.day)
    try:
        return datetime.strptime(texto, "%Y-%m-%d")
    except ValueError as exc:
        raise ErrorObjeciones(f"Fecha inválida: '{texto}'. Se espera AAAA-MM-DD.") from exc


# ─── El resultado ────────────────────────────────────────────────────────────


@dataclass
class Resultado:
    """Lo que la pantalla necesita mostrar y lo que el auditor descarga."""

    entidad: str
    entidad_id: str
    fecha: str
    facturas: int
    objeciones: int
    valor_total: int
    confianza: dict[str, int] = field(default_factory=dict)
    revisar: list[dict[str, Any]] = field(default_factory=list)
    por_factura: list[dict[str, Any]] = field(default_factory=list)
    reglas_ok: bool = True
    fallas_reglas: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    nombre_objeciones: str = ""
    nombre_cruce: str = ""
    objeciones_xlsx: bytes = b""
    cruce_xlsx: bytes = b""

    @property
    def ubicadas(self) -> int:
        return self.confianza.get("ALTA", 0) + self.confianza.get("MEDIA", 0)

    @property
    def pendientes(self) -> int:
        """Lo que hay que mirar a mano: cruce débil, sin cruce o con alerta."""
        return len(self.revisar)

    def resumen(self) -> dict[str, Any]:
        """El resultado SIN los archivos, para responder en JSON."""
        return {
            "entidad": self.entidad,
            "entidad_id": self.entidad_id,
            "fecha": self.fecha,
            "facturas": self.facturas,
            "objeciones": self.objeciones,
            "valor_total": self.valor_total,
            "confianza": self.confianza,
            "ubicadas": self.ubicadas,
            "pendientes": self.pendientes,
            "revisar": self.revisar,
            "por_factura": self.por_factura,
            "reglas_ok": self.reglas_ok,
            "fallas_reglas": self.fallas_reglas,
            "avisos": self.avisos,
            "nombre_objeciones": self.nombre_objeciones,
            "nombre_cruce": self.nombre_cruce,
        }


# ─── El armado, entidad por entidad ──────────────────────────────────────────
#
# Todos reciben la LISTA de archivos que subió el auditor: casi todas las
# entidades mandan uno solo, pero EMSSANAR manda un PDF por factura.


def _uno(rutas: list[Path]) -> Path:
    """El único archivo esperado; si vienen varios se avisa en vez de ignorarlos."""
    if len(rutas) > 1:
        raise ValueError(
            f"Esta entidad manda un solo Excel por lote. Subiste {len(rutas)} archivos: dejá uno."
        )
    return rutas[0]


def _filas_famisanar(
    bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas
) -> list[dict]:
    ruta = _uno(rutas)
    return bot.construir_registros(
        ruta,
        fecha=fecha,
        consecutivo=1,
        codigo_sufijo="01",
        mapa_codigos=None,
        servicios_dgh=servicios,
        trazas=trazas,
    )


def _filas_dispensario(
    bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas
) -> list[dict]:
    ruta = _uno(rutas)
    objeciones = bot.leer_excel_glosa_inicial(ruta)
    filas = bot.filas_desde_excel(objeciones, fecha, servicios, trazas)
    # Ese bot devuelve listas posicionales; acá se usan dicts como los demás.
    return [dict(zip(bot.ENCABEZADOS, f, strict=True)) for f in filas]


def _filas_savia(
    bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas
) -> list[dict]:
    ruta = _uno(rutas)
    return bot.construir_registros(
        ruta,
        fecha=fecha,
        consecutivo=1,
        codigo_sufijo="01",
        mapa_codigos=None,
        servicios_dgh=servicios,
        trazas=trazas,
    )


def _filas_saludtotal(
    bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas
) -> list[dict]:
    ruta = _uno(rutas)
    return bot.construir_registros(
        ruta, fecha=fecha, consecutivo=1, servicios_dgh=servicios, trazas=trazas
    )


def _filas_mutual(
    bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas
) -> list[dict]:
    # MUTUAL lista el mismo servicio bajo varios conceptos de glosa pero lo
    # cuenta una sola vez: sin fusionar, el archivo reclamaría de más.
    objeciones = bot.fusionar_dobles_glosas(bot.leer_mutual(_uno(rutas)))
    return bot.construir_filas(objeciones, fecha, servicios, trazas)


def _filas_sanitas(
    bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas
) -> list[dict]:
    ruta = _uno(rutas)
    return bot.construir_filas(bot.leer_sanitas(ruta), fecha, servicios, trazas)


# GENUSUARIO4 con el que la pantalla firma todos los cargues, el mismo de los
# demás bots (el de VCO por consola escribe "CARTERA").
GENUSUARIO4_PANTALLA = "999"


@dataclass
class _ConfigVco:
    """Lo que el bot de VCO espera de su `argparse` cuando lo llama la pantalla."""

    fecha_documento: date
    fecha_objecion: date
    consecutivo_inicial: int = 1
    referencia: str = ""
    usuario: str = GENUSUARIO4_PANTALLA
    centro_costos: str = ""  # regla fija: CTNCENCOS va vacía
    tipo_objecion: str = ""  # vacío = se decide por factura (0/1/2)
    sin_prefijo: bool = False
    detalle_servicio: bool = False


def _filas_vco(
    bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas
) -> list[dict]:
    ruta = _uno(rutas)
    formato, filas, idx = bot.leer_entrada(ruta, None)
    if formato != "consolidado":
        raise ValueError(
            "Ese archivo ya viene en formato OBJECIONES (16 columnas). Acá se "
            "sube el consolidado del acta del portal VCO, el de NUMERO FACTURA "
            "y VALOR GLOSA."
        )
    cfg = _ConfigVco(fecha_documento=fecha.date(), fecha_objecion=fecha.date())
    return bot.consolidado_a_registros(filas, idx, cfg, servicios_dgh=servicios, trazas=trazas)


def _filas_emssanar(
    bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas
) -> list[dict]:
    """EMSSANAR manda un PDF por factura; se procesan todos de una vez."""
    filas: list[dict] = []
    vistas: set[str] = set()
    for consec, ruta in enumerate(sorted(rutas), start=1):
        try:
            res = bot.procesar_pdf(ruta)
        except Exception as exc:
            raise ValueError(
                f"No pude leer el PDF «{ruta.name}»: ¿es una objeción de ripslink "
                f"(«Objeción a Factura N° HUS…»)? ({exc})"
            ) from exc
        factura = res["encabezado"].get("factura") or bot.normalizar_factura(
            ruta.stem.split("_")[-1]
        )
        if factura in vistas:  # el mismo PDF subido dos veces
            continue
        vistas.add(factura)
        for renglon in res["renglones"]:
            filas.append(
                bot.renglon_a_fila(
                    renglon,
                    consec=consec,
                    factura=factura,
                    fecha=res["encabezado"].get("fecha_objecion") or fecha,
                    usuario=GENUSUARIO4_PANTALLA,
                    tipobj=0,  # provisional: la regla lo decide por factura
                    servicios_dgh=servicios,
                    trazas=trazas,
                )
            )
    return bot.aplicar_crotipobj_por_factura(filas)


# El ADRES tiene su propio motor de homologación (SOAT↔CUPS, topes de valor,
# lotes de 300 facturas), así que no usa el cruce de `_cruce_dgh`. Para que la
# pantalla muestre el mismo resumen, sus métodos se traducen a confianzas.
_CONFIANZA_ADRES = {
    "codigo directo": "ALTA",
    "homologado SOAT→CUPS": "ALTA",
    "descripcion igual": "ALTA",
    "descripcion empieza igual": "MEDIA",
    "valor + palabras en comun": "MEDIA",
    "descripcion parecida": "MEDIA",
}


def _filas_adres(bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas):
    """El Excel de glosas del ADRES, más el homologador CUPS↔SOAT si lo suben.

    El segundo archivo (opcional) es el Homologador Gold Standard: sin él, los
    códigos SOAT del ADRES sólo cruzan por nombre y valor.
    """
    glosas = rutas[0]
    homologador = rutas[1] if len(rutas) > 1 else None
    if len(rutas) > 2:
        raise ValueError(
            "Para el ADRES se suben uno o dos archivos: el Excel de glosas y, "
            f"si lo tenés, el homologador CUPS↔SOAT. Subiste {len(rutas)}."
        )

    filas_adres = bot.leer_adres(glosas)
    if not filas_adres:
        raise ValueError(
            "No encontré ninguna glosa en ese archivo. ¿Es el Excel de glosas "
            "del ADRES, el que arma el auditor con la clasificación y el valor "
            "aceptado de cada renglón?"
        )
    soat_a_cups = bot.leer_homologador(homologador) if homologador else {}
    conversion = bot.construir_registros(
        filas_adres,
        bot.leer_dgh(ruta_dgh),
        soat_a_cups,
        fecha=fecha,
    )

    for registro, resolucion in zip(conversion.registros, conversion.resoluciones, strict=True):
        factura = str(registro["CRNCXC"])
        # El grupo no sale del código (el ADRES usa 3106, 3209…): lo dice la
        # clasificación, que el bot ya resumió por factura.
        registro["_grupos"] = sorted(conversion.grupos.get(factura, set()))
        trazas.append(
            {
                "factura": factura,
                "codigo_objecion": registro["CRNCONOBJ"] or "",
                "valor": int(registro["CROVALOBJ"] or 0),
                "cod_entidad": "",
                "desc_entidad": "",
                "unitario_entidad": 0,
                "observacion": registro["CRDOBSERV"] or "",
                "servicio_dgh": resolucion.candidato_desc or "",
                "cod_dgh": registro["SLNSERPRO"] or "",
                "unitario_dgh": int(resolucion.tope_servicio or 0),
                "centro_costo": resolucion.centro_costo or "",
                "confianza": _CONFIANZA_ADRES.get(resolucion.metodo, "SIN CRUCE")
                if registro["SLNSERPRO"]
                else "SIN CRUCE",
                "motivos": resolucion.metodo or "",
                "puntaje": 0,
                "aviso": "" if registro["SLNSERPRO"] else _cargar("_cruce_dgh").AVISO_SIN_CRUCE,
            }
        )
    return conversion.registros


def _filas_capital(
    bot, rutas: list[Path], ruta_dgh: Path, fecha: datetime, servicios, trazas
) -> list[dict]:
    return bot.construir_filas(bot.leer_capital(_uno(rutas)), fecha, servicios, trazas)


_ARMADORES = {
    "famisanar": _filas_famisanar,
    "dispensario": _filas_dispensario,
    "savia": _filas_savia,
    "saludtotal": _filas_saludtotal,
    "sanitas": _filas_sanitas,
    "vco": _filas_vco,
    "emssanar": _filas_emssanar,
    "adres": _filas_adres,
    "mutual": _filas_mutual,
    "capital": _filas_capital,
}


def procesar(
    archivo_entidad: bytes | Sequence[bytes],
    archivo_dgh: bytes,
    entidad_id: str | None = None,
    fecha: str | date | datetime | None = None,
) -> Resultado:
    """Arma los dos archivos a partir de lo que subió el auditor.

    `archivo_entidad` es el Excel de glosas de la entidad; EMSSANAR manda PDF
    en vez de Excel, y uno por factura, así que también se acepta una lista.
    """
    crudos = (
        [archivo_entidad]
        if isinstance(archivo_entidad, (bytes, bytearray))
        else list(archivo_entidad)
    )
    archivos = [bytes(a) for a in crudos if a]
    if not archivos:
        raise ErrorObjeciones("Falta el archivo de glosas de la entidad.")
    if not archivo_dgh:
        raise ErrorObjeciones(
            "Falta el export de servicios facturados del DGH. Sin él no se puede "
            "saber a qué servicio va cada objeción y el archivo saldría sin códigos."
        )
    for datos, cual in [(a, "de la entidad") for a in archivos] + [(archivo_dgh, "del DGH")]:
        if len(datos) > MAX_BYTES:
            raise ErrorObjeciones(f"El archivo {cual} pesa más de {MAX_BYTES // (1024 * 1024)} MB.")

    entidad = entidad_por_id(entidad_id) if entidad_id else detectar_entidad(archivos[0])
    momento = _fecha(fecha)
    cruce = _cargar("_cruce_dgh")
    bot = _cargar(entidad.modulo)
    avisos: list[str] = []

    # Todo pasa dentro de una carpeta temporal que se borra sola: los archivos
    # del auditor no quedan tirados en el servidor.
    with tempfile.TemporaryDirectory(prefix="objeciones-") as carpeta:
        base = Path(carpeta)
        ruta_dgh = base / "servicios_dgh.xlsx"
        salida_obj = base / "objeciones.xlsx"
        salida_cruce = base / "cruce.xlsx"
        ruta_dgh.write_bytes(archivo_dgh)
        rutas_entidad = []
        for n, datos in enumerate(archivos, start=1):
            # El nombre importa: el bot de EMSSANAR saca la factura del PDF,
            # pero si el encabezado no la trae cae al nombre del archivo.
            ruta = base / f"glosas_{n:03d}{'.pdf' if es_pdf(datos) else '.xlsx'}"
            ruta.write_bytes(datos)
            rutas_entidad.append(ruta)

        try:
            servicios = cruce.leer_servicios_dgh(ruta_dgh, avisar=avisos.append)
        except ValueError as exc:
            raise ErrorObjeciones(str(exc)) from exc

        trazas: list[dict] = []
        try:
            filas = _ARMADORES[entidad.id](bot, rutas_entidad, ruta_dgh, momento, servicios, trazas)
        except ValueError as exc:
            raise ErrorObjeciones(str(exc)) from exc
        if not filas:
            raise ErrorObjeciones("El archivo de la entidad no trae ninguna objeción.")

        # Las reglas fijas se comprueban sobre lo que se va a entregar.
        fallas = cruce.verificar_reglas(
            [
                {
                    "factura": f["CRNCXC"],
                    "slnserpro": f["SLNSERPRO"],
                    "ctncencos": f["CTNCENCOS"],
                    "crotipobj": f["CROTIPOBJ"],
                    "codigo_glosa": f["CRNCONOBJ"],
                    # El ADRES no dice el grupo en el código: lo trae aparte.
                    "grupos": f.get("_grupos"),
                }
                for f in filas
            ],
            servicios,
        )

        _escribir_objeciones(bot, entidad.id, filas, salida_obj)
        cruce.escribir_reporte_cruce(trazas, salida_cruce, entidad=entidad.nombre)

        sufijo = momento.strftime("%d%m%Y")
        corto = entidad.corto
        facturas = sorted({f["CRNCXC"] for f in filas})
        if len(facturas) > MAX_FACTURAS_DGH:
            avisos.append(
                f"El lote trae {len(facturas)} facturas y el DGH no recibe más de "
                f"{MAX_FACTURAS_DGH} por archivo: hay que partirlo antes de subirlo."
            )
        return Resultado(
            entidad=entidad.nombre,
            entidad_id=entidad.id,
            fecha=momento.strftime("%Y-%m-%d"),
            facturas=len(facturas),
            objeciones=len(filas),
            valor_total=sum(int(f["CROVALOBJ"] or 0) for f in filas),
            confianza={k: sum(1 for t in trazas if t["confianza"] == k) for k in CONFIANZAS},
            revisar=_filas_revisar(trazas),
            por_factura=_por_factura(filas, trazas),
            reglas_ok=not fallas,
            fallas_reglas=fallas,
            avisos=avisos,
            nombre_objeciones=f"OBJECIONES_{corto}_{sufijo}.xlsx",
            nombre_cruce=f"CRUCE_{corto}_{sufijo}.xlsx",
            objeciones_xlsx=salida_obj.read_bytes(),
            cruce_xlsx=salida_cruce.read_bytes(),
        )


def _escribir_objeciones(bot, entidad_id: str, filas: list[dict], salida: Path) -> None:
    """Cada bot escribe su Excel; el formato de celdas es el suyo."""
    if entidad_id == "saludtotal":
        bot.escribir_consolidado(filas, salida)
    elif entidad_id == "dispensario":
        bot.escribir_excel([[f[c] for c in bot.ENCABEZADOS] for f in filas], salida)
    elif entidad_id in ("sanitas", "mutual", "capital"):
        bot.escribir_objeciones(filas, salida)
    elif entidad_id == "vco":
        bot.escribir_cargue([[f[c] for c in bot.COLUMNAS_CARGUE] for f in filas], salida)
    elif entidad_id == "emssanar":
        bot.escribir_excel(filas, salida)
    elif entidad_id == "adres":
        bot.escribir_objeciones(filas, salida)
    else:
        bot.escribir_consolidado(filas, salida)


def _filas_revisar(trazas: list[dict]) -> list[dict[str, Any]]:
    """Lo que el auditor tiene que mirar, listo para la tabla de la pantalla."""
    return [
        {
            "factura": t["factura"],
            "codigo_glosa": t["codigo_objecion"],
            "valor": t["valor"],
            "servicio_entidad": t["desc_entidad"],
            "codigo_entidad": t["cod_entidad"],
            "servicio_dgh": t["servicio_dgh"],
            "codigo_dgh": t["cod_dgh"],
            "confianza": t["confianza"],
            "motivos": t["motivos"],
            "aviso": t["aviso"],
        }
        for t in trazas
        if t["aviso"] or t["confianza"] in ("BAJA", "SIN CRUCE")
    ]


def _por_factura(filas: list[dict], trazas: list[dict]) -> list[dict[str, Any]]:
    """Una línea por factura: cuántas objeciones, cuánto y qué falta."""
    resumen: dict[str, dict[str, Any]] = {}
    for f in filas:
        d = resumen.setdefault(
            f["CRNCXC"],
            {"factura": f["CRNCXC"], "objeciones": 0, "valor": 0, "sin_codigo": 0, "tipo": 0},
        )
        d["objeciones"] += 1
        d["valor"] += int(f["CROVALOBJ"] or 0)
        d["sin_codigo"] += 1 if not f["SLNSERPRO"] else 0
        d["tipo"] = f["CROTIPOBJ"]
    pendientes: dict[str, int] = {}
    for t in trazas:
        if t["aviso"] or t["confianza"] in ("BAJA", "SIN CRUCE"):
            pendientes[t["factura"]] = pendientes.get(t["factura"], 0) + 1
    for factura, d in resumen.items():
        d["revisar"] = pendientes.get(factura, 0)
    return sorted(resumen.values(), key=lambda d: d["factura"])
